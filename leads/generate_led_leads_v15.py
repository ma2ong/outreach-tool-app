"""LED Display Leads — Americas + Korea — v15
Extends v14 (420 records) with 6 new entries = 426 total.
Sources: Instagram opencli + Facebook opencli
New: Argentina +2, Mexico +1, Korea +1, Colombia +2
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v14 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12", "v13", "v14"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (421 – 426) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Argentina new (2) ====================
    {
        "no":421,"country":"Argentina","region":"Americas",
        "company_en":"TNX Argentina","company_local":"",
        "city":"Argentina",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Professional audio, lighting, LED panels sales and engineering services; all types of projects; Instagram @tnx_arg (6,484 followers, Verified, 245 posts)",
        "facebook":"","instagram":"tnx_arg","linkedin":"",
    },
    {
        "no":422,"country":"Argentina","region":"Americas",
        "company_en":"LED VJ Argentina","company_local":"",
        "city":"Buenos Aires",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+54 11 5745-4880",
        "website":"",
        "business":"LED VJ audiovisual systems service; indoor/outdoor LED panels for events; social, corporate, cultural, festival; Buenos Aires; Instagram @led.vj (1,098 followers); phone confirmed from bio",
        "facebook":"","instagram":"led.vj","linkedin":"",
    },
    # ==================== Mexico new (1) ====================
    {
        "no":423,"country":"Mexico","region":"Americas",
        "company_en":"Pantallas LED MTY","company_local":"",
        "city":"Monterrey, Nuevo León",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental; AV solutions for corporate and social events; Monterrey Mexico; Instagram @pantallasled.mty (400 followers, 39 posts)",
        "facebook":"","instagram":"pantallasled.mty","linkedin":"",
    },
    # ==================== Korea new (1) ====================
    {
        "no":424,"country":"Korea","region":"Asia",
        "company_en":"Sign Pro","company_local":"네오사인",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"1644-2546",
        "website":"",
        "business":"Sign & LED total solution specialist; 15 years experience; public institution signage; Incheon International Airport project; disability-owned enterprise (수의계약 가능); Instagram @sign_pro_ (1,147 followers, 137 posts); phone confirmed from bio",
        "facebook":"","instagram":"sign_pro_","linkedin":"",
    },
    # ==================== Colombia new (2) ====================
    {
        "no":425,"country":"Colombia","region":"Americas",
        "company_en":"Azkar Eventos","company_local":"",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED giant screen & AV event production; corporate and major events; Bogotá Colombia; FIFA World Cup 2026 event viewing screens; Facebook @azkareventos (2,290 followers)",
        "facebook":"azkareventos","instagram":"","linkedin":"",
    },
    {
        "no":426,"country":"Colombia","region":"Americas",
        "company_en":"Proyectores y Pantallas LED","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel and projector import & sales; IT company; Colombia; Facebook @proyectoresypantallasledsas (3,189 followers)",
        "facebook":"proyectoresypantallasledsas","instagram":"","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v15 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-07  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v15"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","사인코드(010-5180-0775, Verified)3000+装案。네오사인Sign Pro(1644-2546)仁川机场项目。빛가람시스템(bgr@bitgaram.biz)釜山。현대퓨처넷大型集成商。LSI정보기술(042-627-8400)大田。이지미디어(010-8780-1003)教堂。"),
    ("USA","Brightlink(sales@brightlinkav.com, 855-449-4733)全北美。LED Factory USA(Houston)零售供应。FastInstall(425-766-3232)Seattle安装。Eagle LED USA(info@eagleled.us)分销商。"),
    ("Brazil","AMS Locação(André Moreira +55 18 99749-5003)。Azul Light(+55 44 3031-5492)22K粉。MW LED(mwled.com.br)。Infinity(2,873粉)、Locação Pelotas(2,132粉)活跃。"),
    ("Canada","Inno-Leader有深圳背景。AVT/Illusion Universe/Digital Edge适合经销商。Brightlink也覆盖加拿大。"),
    ("Chile","CGS Chile(+56 2 2418-1225)重点。arriendopantallasledchile已联系。LED Chile Pantallas新增。"),
    ("Argentina","TNX(6,484粉, Verified)买卖LED+音响+灯光全套。LED VJ(+54 11 5745-4880 Buenos Aires)活动租赁。Tucumán(+54 386-551-0502)在库。"),
    ("Colombia","Azkar Eventos(Bogotá, 世界杯活动大屏)、Proyectores y Pantallas LED(3,189粉)新增。360GROUP(24K粉)、Eventos Kaos(6K粉, Medellín)。Visionled(Fredy Ruiz +57 310-492-3133)重点。LED Master(+57 300-204-4085)有电话。"),
    ("Peru","Lima核心，WhatsApp是有效渠道。LED Maker Peru(ledmaker.com.pe)已在库。"),
    ("Mexico","Pantallas LED MTY(Monterrey)新增。PUNTO LED(Zapopan)制造分销。MV Producciones GDL(Guadalajara)租赁。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v15.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
