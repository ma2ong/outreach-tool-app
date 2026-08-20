"""LED Display Leads — Americas + Korea — v12
Extends v11 (396 records) with 7 new entries = 403 total.
Sources: Instagram opencli + Facebook opencli + Jina website fetch
New: Korea +3, Brazil +2, Mexico +2
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v11 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (397 – 403) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Korea new (3) ====================
    {
        "no":397,"country":"Korea","region":"Asia",
        "company_en":"EZMedia Korea","company_local":"이지미디어",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 10-8780-1003",
        "website":"",
        "business":"Church video/audio/special lighting design & installation specialist; LED signboard + projector; Instagram @ezmedia_official (80 followers); phone confirmed from bio",
        "facebook":"","instagram":"ezmedia_official","linkedin":"",
    },
    {
        "no":398,"country":"Korea","region":"Asia",
        "company_en":"LSI Information Technology","company_local":"(주)LSI정보기술",
        "city":"Daejeon",
        "contact_name":"","title":"",
        "email":"lsi8400@lsit.kr","phone_whatsapp":"042-627-8400",
        "website":"",
        "business":"LED signboard display specialist; outdoor LED, indoor LED, LED screen; Daejeon region focus; Instagram @lsit8400 (672 followers); phone + email confirmed from bio",
        "facebook":"","instagram":"lsit8400","linkedin":"",
    },
    {
        "no":399,"country":"Korea","region":"Asia",
        "company_en":"MOESS","company_local":"모이써",
        "city":"Daegu",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Daegu AV equipment: projector, screen, TV, signage, DID kiosk, audio — sales & installation; KakaoTalk Plus: MOESS; Instagram @moess_jnroad (56 followers, 104 posts)",
        "facebook":"","instagram":"moess_jnroad","linkedin":"",
    },
    # ==================== Brazil new (2) ====================
    {
        "no":400,"country":"Brazil","region":"Americas",
        "company_en":"MW LED","company_local":"",
        "city":"Brazil (nationwide)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://mwled.com.br",
        "business":"LED panel rental & sales for events nationwide; fairs, shows, corporate events; Instagram @mw.led (10,779 followers); website mwled.com.br",
        "facebook":"","instagram":"mw.led","linkedin":"",
    },
    {
        "no":401,"country":"Brazil","region":"Americas",
        "company_en":"RC LED Produções Visuais","company_local":"",
        "city":"Brazil (nationwide)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel rental company; shows, fairs, rodeos, corporate events, nationwide Brazil; Facebook @comercial2rcled.com.br (749 followers)",
        "facebook":"comercial2rcled.com.br","instagram":"","linkedin":"",
    },
    # ==================== Mexico new (2) ====================
    {
        "no":402,"country":"Mexico","region":"Americas",
        "company_en":"PUNTO LED","company_local":"",
        "city":"Zapopan, Jalisco",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED large screen manufacture + distribution; billboards, banners, price boards, commercial signage; Zapopan Jalisco; Facebook @puntoledpantallas (410 followers)",
        "facebook":"puntoledpantallas","instagram":"","linkedin":"",
    },
    {
        "no":403,"country":"Mexico","region":"Americas",
        "company_en":"MV Producciones GDL","company_local":"",
        "city":"Guadalajara, Jalisco",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental + CCTV + audio + lighting; events; Guadalajara; Instagram @pantallasledguadalajara (57 followers)",
        "facebook":"","instagram":"pantallasledguadalajara","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v12 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-07  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v12"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","LSI정보기술大田(042-627-8400, lsi8400@lsit.kr)有直接联系方式。이지미디어(010-8780-1003)专攻教堂市场。모이써(KakaoTalk)大邱渠道。白象LED釜山(1833-7989)。汉邦IT首尔(02-467-7447)。"),
    ("USA","FastInstall(425-766-3232)做视频墙安装，Eagle LED USA(info@eagleled.us)是分销商。"),
    ("Brazil","MW LED(mwled.com.br, 10K+粉丝)是大型活动租赁商。Azul Light+BSL Light已有联系。RC LED Produções做全国租赁。"),
    ("Canada","Inno-Leader有深圳背景，沟通顺畅。AVT/Illusion Universe体量较大，适合做经销商。"),
    ("Chile","CGS Chile(+56 2 2418-1225 / WA +56 9 9532-2717)已联系。"),
    ("Argentina","Tucumán(+54 386-551-0502)、Si-Led Buenos Aires等3条新增。"),
    ("Colombia","Visionled(Fredy Ruiz +57 310-492-3133)、Alquiler Pantallas LED Colombia(Bogotá)是重点。"),
    ("Peru","Lima核心，LED Maker Peru已在库。可跟进alquiler_de_pantallas_led_peru。"),
    ("Mexico","PUNTO LED(Zapopan Jalisco)制造分销大屏。MV Producciones GDL做Guadalajara租赁。Monterrey/CDMX继续开发。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v12.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
