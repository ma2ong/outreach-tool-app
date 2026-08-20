"""LED Display Leads — Americas + Korea — v20
Extends v19 (509 records) with 16 new entries = 525 total.
Sources: WebSearch + WebFetch (Jina/company websites)
New: USA +9, Mexico +7
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v19 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12", "v13", "v14", "v15", "v16", "v17", "v18", "v19"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (510 – 525) ────────────────────────────────────────────────────
new_entries = [
    # ==================== USA new (9) ====================
    {
        "no":510,"country":"USA","region":"Americas",
        "company_en":"Z3 LED Solutions","company_local":"",
        "city":"Pompano Beach, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"(954) 971-8842",
        "website":"z3ledsolutions.com",
        "business":"LED video wall rental and event production; South Florida (Dade/Broward/Palm Beach); P7 indoor/outdoor panels; website confirmed z3ledsolutions.com; Facebook @z3ledsolutions; Yelp listed (954) 971-8842; address: 4100 N Powerline Rd Ste U1, Pompano Beach FL 33066",
        "facebook":"z3ledsolutions","instagram":"","linkedin":"",
    },
    {
        "no":511,"country":"USA","region":"Americas",
        "company_en":"RentaLED USA","company_local":"",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental for events; full-service AV production for congresses and conventions; Instagram @rentaledusa (1,241 followers)",
        "facebook":"","instagram":"rentaledusa","linkedin":"",
    },
    {
        "no":512,"country":"USA","region":"Americas",
        "company_en":"TriVision Studios","company_local":"",
        "city":"Chantilly, VA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 703-608-9680",
        "website":"trivisionstudios.com",
        "business":"LED video wall rental and AV integration; Washington DC/Northern VA/NYC/Baltimore/Richmond; indoor LED fine pitch P1.9mm; 30+ years experience; corporate conferences, government forums, live performances; phone confirmed from website; Instagram @trivisioninc; Facebook @trivisionstudios",
        "facebook":"trivisionstudios","instagram":"trivisioninc","linkedin":"",
    },
    {
        "no":513,"country":"USA","region":"Americas",
        "company_en":"Vantage Production Group","company_local":"",
        "city":"Lemont, IL",
        "contact_name":"","title":"",
        "email":"info@vantagepg.com",
        "phone_whatsapp":"815-469-0000",
        "website":"vantagepg.com",
        "business":"LED video wall rental and full-service corporate AV; Chicago and Midwest; 20+ years; serves Downtown Chicago, North Shore, Milwaukee; phone/email confirmed from website; Instagram @Soundworksproductions; Facebook @vantageproductiongroup",
        "facebook":"vantageproductiongroup","instagram":"Soundworksproductions","linkedin":"",
    },
    {
        "no":514,"country":"USA","region":"Americas",
        "company_en":"AVR Expos","company_local":"",
        "city":"USA (Nationwide)",
        "contact_name":"","title":"",
        "email":"info@avrexpos.com",
        "phone_whatsapp":"1-800-693-6668",
        "website":"avrexpos.com",
        "business":"LED video wall rental nationwide for trade shows, conferences, live events; beMatrix 2.5mm technology; offices in Atlanta, Boston, Chicago, Dallas, LA, NYC, Washington DC; phone/email confirmed from website; Facebook @AVRexposLLC",
        "facebook":"AVRexposLLC","instagram":"","linkedin":"",
    },
    {
        "no":515,"country":"USA","region":"Americas",
        "company_en":"Insane Impact","company_local":"",
        "city":"Des Moines, IA",
        "contact_name":"","title":"",
        "email":"info@insaneimpact.com",
        "phone_whatsapp":"(515) 349-7708",
        "website":"insaneimpact.com",
        "business":"Nationwide LED screen rental; 500+ LED trailers across all 50 states; 2500+ events; mobile jumbotrons + permanent installation; phone/email confirmed from website; Instagram @insaneimpact (1,626 followers); Facebook @InsaneImpact",
        "facebook":"InsaneImpact","instagram":"insaneimpact","linkedin":"",
    },
    {
        "no":516,"country":"USA","region":"Americas",
        "company_en":"NABS Creative","company_local":"",
        "city":"Utah, USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"(435) 268-2221",
        "website":"nabscreative.com",
        "business":"LED screen rental nationwide (mainland USA); mobile LED screen trailers + outdoor modular systems + indoor modular; sizes 17x10 to 30x17; motorsport, equestrian, concerts, church graduations; phone confirmed from website",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":517,"country":"USA","region":"Americas",
        "company_en":"CMG Visuals","company_local":"",
        "city":"Dallas, TX",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"214-714-2153",
        "website":"cmgvisuals.com",
        "business":"LED video wall sales, installation, rental, and service; Dallas-Fort Worth and beyond; serves houses of worship, TV/film, tradeshows, broadcast studios, retail, corporate; phone confirmed from website; Instagram @cmgvisualsled; Facebook profile confirmed",
        "facebook":"","instagram":"cmgvisualsled","linkedin":"",
    },
    {
        "no":518,"country":"USA","region":"Americas",
        "company_en":"LED Wall Systems (EncoreX)","company_local":"",
        "city":"Union City, NJ",
        "contact_name":"","title":"",
        "email":"sales@ledwallsystems.com",
        "phone_whatsapp":"(646) 229-2995",
        "website":"ledwallsystems.com",
        "business":"LED video wall rental and sale; NYC and globally; partners: Leyard, Unilumin, Absen, ROE, Novastar; volume stages, broadcast studios, DOOH; phone/email confirmed from website; address: 109 45th St, Union City NJ 07087; Facebook listed",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== Mexico new (7) ====================
    {
        "no":519,"country":"Mexico","region":"Americas",
        "company_en":"ADOC Producciones","company_local":"",
        "city":"Ciudad de México",
        "contact_name":"","title":"",
        "email":"contacto@adocproducciones.com.mx",
        "phone_whatsapp":"+52 55 6880 6063",
        "website":"adocproducciones.com.mx",
        "business":"Renta de pantallas LED y produccion audiovisual para eventos corporativos; CDMX, Monterrey, Guadalajara; phone/email/WhatsApp confirmed from website (+52 55 6880 6063); also: MTY +52 81 8084 8841 / GDL +52 33 3228 0736; Instagram @adoc_producciones; Facebook @ADOCPRODUCCIONES",
        "facebook":"ADOCPRODUCCIONES","instagram":"adoc_producciones","linkedin":"",
    },
    {
        "no":520,"country":"Mexico","region":"Americas",
        "company_en":"Renta Pantalla LED MTY","company_local":"",
        "city":"Monterrey, NL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+52 81 1616 5994",
        "website":"rentadepantallasledmty.com",
        "business":"Renta de pantallas LED para eventos corporativos en Monterrey; phone/WhatsApp confirmed from website (8116165994 mobile); Instagram @pantallasled.mty; Facebook @rentadepantallasledmty",
        "facebook":"rentadepantallasledmty","instagram":"pantallasled.mty","linkedin":"",
    },
    {
        "no":521,"country":"Mexico","region":"Americas",
        "company_en":"Proyector MTY","company_local":"",
        "city":"Monterrey, NL",
        "contact_name":"","title":"",
        "email":"proyectormty@gmail.com",
        "phone_whatsapp":"+52 81 1805 1719",
        "website":"proyectormty.wixsite.com/rentadeproyectores",
        "business":"Renta de proyectores y pantallas LED en Monterrey y area metropolitana; phone/email confirmed from Google search; Facebook @proyectormty",
        "facebook":"proyectormty","instagram":"","linkedin":"",
    },
    {
        "no":522,"country":"Mexico","region":"Americas",
        "company_en":"Audiovisuales Kanek","company_local":"",
        "city":"Guadalajara, Jalisco",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+52 33 3440 1490",
        "website":"audiovisualeskanek.com",
        "business":"Renta de pantallas LED gigantes 3.91mm y 2.6mm para eventos; Guadalajara, Puerto Vallarta, Tequila, Ciudad Guzman; phone/WhatsApp confirmed from website (+52 33 3440 1490); also office (33) 3665 8078; Instagram @audiovisuales_kanek; Facebook @KanekAudiovisuales",
        "facebook":"KanekAudiovisuales","instagram":"audiovisuales_kanek","linkedin":"",
    },
    {
        "no":523,"country":"Mexico","region":"Americas",
        "company_en":"WLK Producciones","company_local":"",
        "city":"Ciudad de México",
        "contact_name":"","title":"",
        "email":"wlkaudiovisual@gmail.com",
        "phone_whatsapp":"+52 55 4615 2641",
        "website":"wlkproducciones.com.mx",
        "business":"Renta de pantallas LED y audio profesional en Mexico; 2500+ m2 de pantallas LED; conciertos, festivales, estadios, ferias, eventos corporativos; cobertura nacional; phone/email confirmed from website; also +52 55 3046 8005 / +52 55 8029 7034; Instagram @wlkproducciones; Facebook @wlkaudiovisual",
        "facebook":"wlkaudiovisual","instagram":"wlkproducciones","linkedin":"",
    },
    {
        "no":524,"country":"Mexico","region":"Americas",
        "company_en":"Kzas Audiovisual","company_local":"",
        "city":"Tijuana, Baja California",
        "contact_name":"","title":"",
        "email":"kzaspro@gmail.com",
        "phone_whatsapp":"+52 664 508 5791",
        "website":"kzasaudiovisual.com",
        "business":"Renta de pantallas LED interior y exterior; 10+ anos de experiencia; Tijuana, Rosarito, Mexicali, Ensenada, Valle de Guadalupe; phone/WhatsApp/email confirmed from website (+526645085791); Instagram @kzasaudiovisual; Facebook @KzasAudiovisual",
        "facebook":"KzasAudiovisual","instagram":"kzasaudiovisual","linkedin":"",
    },
    {
        "no":525,"country":"Mexico","region":"Americas",
        "company_en":"Grupo PROMOLED","company_local":"",
        "city":"Guadalajara, Jalisco",
        "contact_name":"","title":"",
        "email":"info@grupopromoled.com",
        "phone_whatsapp":"+52 33 3560 7886",
        "website":"grupopromoled.com",
        "business":"Venta y renta de pantallas LED interior y exterior; 1000+ m2 de pantallas; distribucion nacional; Juan Pablo 483 Col. Chapalita, Guadalajara Jalisco; phone/email confirmed from Foursquare + search; also (33) 1727-9527; Instagram @grupopromoled; Facebook @promoledpantallas (5,808 likes)",
        "facebook":"promoledpantallas","instagram":"grupopromoled","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v20 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-20  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v20"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","케이시스(1600-6187, Novastar合作)重点推。위디에스(02-6953-4419)透明屏/视频墙。디스플레이비투비(032-363-4770)三星商业伙伴。MUXWAVE(02-716-8850)透明LED专门。사인코드(010-5180-0775)3000+装案。STN Media(02-2088-1662)数字标牌+媒体幕墙。Ksys Digital Signage(449粉)LED/视频墙/标牌组合。DS Lighting(1131粉)激光+LED舞台。WS International LED安装专门。"),
    ("USA","LED Wall Systems/EncoreX(Union City NJ, sales@ledwallsystems.com, 646-229-2995)重点—NYC+全球，合作Unilumin/Absen/Novastar。Vantage PG(Chicago, info@vantagepg.com, 815-469-0000)芝加哥20年。AVR Expos(全国, info@avrexpos.com, 800-693-6668)贸易展会专门。CMG Visuals(Dallas TX, 214-714-2153)TX教堂+广播。Insane Impact(Des Moines, info@insaneimpact.com)500+拖车全国。TriVision Studios(Chantilly VA, 703-608-9680)DC政府+企业。Z3 LED Solutions(Pompano Beach FL, 954-971-8842)南佛州活动租赁。Brightlink(sales@brightlinkav.com)全北美。LED Market USA(info@ledmarketusa.com)已收录重点。"),
    ("Brazil","InLED Brasil(contato@inledbrasil.com, 61-99205-7388)重点。CIA do LED(51-992314663)。LED Eart(85-9.9432.4544)。The LED(50K粉)巴西最大。Mídia InLED(15K粉)巴西/美/欧。ALMO RAEL(15624粉, 20年市场)新增重点。LED ABC(3090粉, ABC SP)新增。PLL Painéis LED(542粉)住宅+商业。Ledfy(1135粉)OOH户外数字。"),
    ("Canada","Genoptic Smart Displays(1221粉)模块化LED标牌。Toronto LED Video Wall Rental(645粉)多伦多活动租赁新增。Inno-Leader深圳背景。Brightlink也覆盖加拿大。"),
    ("Chile","CGS Chile(+56 2 2418-1225)重点。NordeLED(1130粉, Viña del Mar)。Pantallas LED en Santiago(FB 1661粉)租赁。AgsLed(FB, Concepción)。"),
    ("Argentina","Prina(16435粉, Verified)重点。GC Pantallas LED(4263粉, Tucumán到全国)新增重点。Golden Rocket(615粉)LED/Totem/Tunnel租赁新增。TNX(6484粉, Verified)买卖全套。LED VJ(+54 11 5745-4880)活动租赁。"),
    ("Colombia","ZULED Soluciones(10K粉, Medellin, +57 310 385 0334)重点。China LED SAS(5K粉, +57 313 831 6469)直接进口商重点。Target Colombia(1668粉, Cali, +57 318 707 1440)。Vision LED Colombia(3239粉, +57 310 492 3133)全国。Nueva Era Producciones(16K粉, Medellin)大账号。OPE Producciones(1154粉, Bogotá, +57 310 814 1832)。Visionled(Fredy Ruiz +57 310-492-3133)重点。LED Pixel Colombia(1713粉)全国安装。Fabulux+FabuluxLED(Bogotá)。360GROUP(24K粉)。"),
    ("Peru","Espacio Digital SAC(Lima, +51 905 431 819, ventas@espaciodigital360.com)重点—有邮箱。INNOVALED Peru(Lima, +51 986 534 469)videowall+数字标牌重点。Corporacion Cervantes(Lima, +51 997 547 440, 4810粉)大账号。SMLED Peru(Lima, 1486粉)。JSA Rental(Lima, +51 966 123 897)。ELES Producciones(Lima, +51 997 473 896)。Big LED Peru(Lima, +51 990 594 725)。Pantallas LED A&V(Lima, +51 997 059 555)。Ledex Digital Outdoor(2008粉, Lima)户外LED广告重点。PubliLED SAC(919 474 350, Arequipa+Tacna)有电话重点。"),
    ("Mexico","ADOC Producciones(CDMX, contacto@adocproducciones.com.mx, +52 55 6880 6063)重点—CDMX/MTY/GDL全国。WLK Producciones(CDMX, wlkaudiovisual@gmail.com, +52 55 4615 2641)2500m2 LED重点。Grupo PROMOLED(Guadalajara, info@grupopromoled.com, +52 33 3560 7886)1000m2 GDL。Audiovisuales Kanek(Guadalajara, +52 33 3440 1490)活动租赁。Kzas Audiovisual(Tijuana, kzaspro@gmail.com, +52 664 508 5791)Baja CA全覆盖。Renta Pantalla LED MTY(Monterrey, +52 81 1616 5994)。RGB Media(1433粉, 1000+项目)重点。Pantallas LED Rino Monterrey(855粉, 3.9pitch)。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v20.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
