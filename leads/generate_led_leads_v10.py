"""LED Display Leads — Americas + Korea — v10
Extends v9 (379 records) with 10 new entries = 389 total.
Sources: Instagram opencli + Facebook opencli + Jina website fetch
New: Brazil +2, USA +1, Colombia +3, Argentina +3, Mexico +1
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 / v6 / v7 / v8 / v9 new_entries ───────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (380 – 389) ────────────────────────────────────────────────────
new_entries = [
    # ==================== USA new (1) ====================
    {
        "no":380,"country":"USA","region":"Americas",
        "company_en":"Eagle LED USA","company_local":"",
        "city":"Orlando, FL / Lynnwood, WA",
        "contact_name":"","title":"",
        "email":"info@eagleled.us","phone_whatsapp":"+1 407-693-4509",
        "website":"https://eagleled.us",
        "business":"LED display supplier; indoor/outdoor LED walls, scoreboards, digital signage; US-based team; multiple sales contacts (italo@eagleled.us +1 305-964-8152); 5,711 IG followers",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== Brazil new (2) ====================
    {
        "no":381,"country":"Brazil","region":"Americas",
        "company_en":"Azul Light","company_local":"",
        "city":"Maringá, PR (branches: Goiânia GO, Balneário Camboriú SC)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+55 44 3031-5492",
        "website":"https://azullight.com.br",
        "business":"LED display specialist; indoor/outdoor LED panels for commerce, churches, shows, outdoor advertising; 4 branches across Paraná, Goiás, Santa Catarina; e-commerce store; 22,314 IG followers (Verified); wa.me confirmed",
        "facebook":"","instagram":"azullightbrasil","linkedin":"",
    },
    {
        "no":382,"country":"Brazil","region":"Americas",
        "company_en":"GT Painel de Led","company_local":"",
        "city":"Aparecida de Goiânia, GO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel & AV full-service rental; 15 years experience; LED+video+audio+light events; Goiânia/Goiás region; Facebook: gtledpainel (957 followers); Instagram: gtledpainel (5,874 followers, Verified)",
        "facebook":"gtledpainel","instagram":"gtledpainel","linkedin":"",
    },
    # ==================== Colombia new (3) ====================
    {
        "no":383,"country":"Colombia","region":"Americas",
        "company_en":"Visionled Colombia","company_local":"",
        "city":"Medellín",
        "contact_name":"Fredy Ruiz","title":"",
        "email":"","phone_whatsapp":"+57 310-492-3133",
        "website":"",
        "business":"LED display integrator & rental; media facades, video walls, events; Medellín; confirmed phone from Facebook post; Facebook @magoarlekin (5,308 followers)",
        "facebook":"magoarlekin","instagram":"","linkedin":"",
    },
    {
        "no":384,"country":"Colombia","region":"Americas",
        "company_en":"Pantallas LED de Colombia","company_local":"",
        "city":"Colombia (nationwide)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display sales & installation; indoor/outdoor P5/P3/P2 panels + video processors; nationwide delivery; Instagram @pantallas_led_de_colombia (1,091 followers)",
        "facebook":"","instagram":"pantallas_led_de_colombia","linkedin":"",
    },
    {
        "no":385,"country":"Colombia","region":"Americas",
        "company_en":"Alquiler Pantallas LED Colombia","company_local":"",
        "city":"Bogotá (Cra 70F # 73a - 45)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental nationwide; large format screens; events, advertising; 24-hour service; Bogotá; Facebook @Alquilerpantallasledcolombia (1,653 followers)",
        "facebook":"Alquilerpantallasledcolombia","instagram":"","linkedin":"",
    },
    # ==================== Argentina new (3) ====================
    {
        "no":386,"country":"Argentina","region":"Americas",
        "company_en":"Si-Led Argentina","company_local":"",
        "city":"Buenos Aires",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel rental; DJ/events/streaming setups; Buenos Aires; Facebook @alquilerdepantalla (749 followers)",
        "facebook":"alquilerdepantalla","instagram":"","linkedin":"",
    },
    {
        "no":387,"country":"Argentina","region":"Americas",
        "company_en":"Alquiler Pantallas LED Tucumán","company_local":"",
        "city":"Tucumán",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+54 386-551-0502",
        "website":"",
        "business":"LED screen rental; events; Tucumán province; Instagram @alquiler_pantallasled_tucuman (514 followers); phone confirmed from bio",
        "facebook":"","instagram":"alquiler_pantallasled_tucuman","linkedin":"",
    },
    {
        "no":388,"country":"Argentina","region":"Americas",
        "company_en":"Alquiler Pantallas Led Argentina","company_local":"",
        "city":"Argentina (nationwide)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental nationwide; advertising LED, large format event screens, totems; Instagram @alquilerpantallasled (481 followers)",
        "facebook":"","instagram":"alquilerpantallasled","linkedin":"",
    },
    # ==================== Mexico new (1) ====================
    {
        "no":389,"country":"Mexico","region":"Americas",
        "company_en":"Pantallas LED Monterrey","company_local":"",
        "city":"Monterrey, Nuevo León",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental; Monterrey area; Instagram @pantallasledmty (361 followers)",
        "facebook":"","instagram":"pantallasledmty","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v10 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-06  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v10"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","韩国现有110条，仍是最大缺口。AVTEAM有Kakao官方账号，可直接Kakao联系。Lzone专注大邱/庆北地区，CS Sheet有柔性LED产品线。Eone OMS(eoneoms.com)专攻数字标牌，070-4211-0510。"),
    ("USA","邮件+WhatsApp为主。Eagle LED USA(eagleled.us)有多个销售联系人，italo@eagleled.us / info@eagleled.us。重点推大屏广告、体育场馆、教堂市场。"),
    ("Brazil","Azul Light(azullight.com.br, +55 44 3031-5492 WA)是目前最大的巴西客户候选，4家分店，22K粉丝。BSL Light已在库。GT Painel de Led活跃于Goiânia，有15年经验。"),
    ("Canada","Inno-Leader有深圳背景，沟通顺畅。AVT/Illusion Universe体量较大，适合做经销商。"),
    ("Chile","CGS Chile(+56 2 2418-1225 / WA +56 9 9532-2717)、arriendopantallasledchile(+56 974753938)是本轮新增优质联系人。"),
    ("Argentina","Tucumán(+54 386-551-0502)、Si-Led Buenos Aires等3条新增。经济不稳，优先推性价比型号，用WhatsApp开发。"),
    ("Colombia","Visionled Colombia(Fredy Ruiz +57 310-492-3133, Medellín)、Alquiler Pantallas LED Colombia(Bogotá)是本轮重点。用WhatsApp+西班牙语开发信。"),
    ("Peru","Lima核心，WhatsApp是唯一有效渠道。从展会/活动rental切入。LED Maker Peru(ledmaker.com.pe)已在库。"),
    ("Mexico","Pantallas LED Monterrey新增。CDMX+Guadalajara+Monterrey三大市场。Fénix Evolution专注活动租赁，推rental产品线。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v10.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
