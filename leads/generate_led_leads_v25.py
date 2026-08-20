"""
LED Display Leads v25 — Argentina (Buenos Aires x2, Rosario) + Peru (Lima)
3 Argentina + 1 Peru, nos 547-550
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 25)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 547,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "LATINLED Argentina",
        "company_local": "LATINLED",
        "city": "Haedo, Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "info@latinled.com.ar",
        "phone_whatsapp": "+54 9 11 3651-9054",
        "website": "https://latinled.com.ar",
        "business": "LED screen technology — modular display systems for event halls, churches, nightclubs, sports, outdoor advertising; 15+ years experience; 12-installment financing; Buenos Aires province; Instagram @latinled",
        "facebook": "",
        "instagram": "latinled",
        "linkedin": "",
    },
    {
        "no": 548,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Imagenes del Sur",
        "company_local": "Imágenes del Sur",
        "city": "Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+54 9 11 51850754",
        "website": "https://www.imagenesdelsur.com",
        "business": "LED screen rental for events — modular 1m x 0.5m panels, scalable to any size; audiovisual services including LED and professional sound; Buenos Aires",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 549,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "MCS y Design Rosario",
        "company_local": "MCS y Rental / MCS y Design",
        "city": "Rosario, Santa Fe",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "https://mcsyrentalrosario.com.ar",
        "business": "AV equipment rental — high-definition LED screens as lead service, 4K displays for exhibitions and digital signage, plus sound/projection/lighting; Rosario Argentina; Instagram @mcsydesign",
        "facebook": "",
        "instagram": "mcsydesign",
        "linkedin": "",
    },
    {
        "no": 550,
        "country": "Peru",
        "region": "Americas",
        "company_en": "Audilux Peru",
        "company_local": "Audilux",
        "city": "Miraflores, Lima",
        "contact_name": "",
        "title": "",
        "email": "eventos@audilux.pe",
        "phone_whatsapp": "+51 972060955",
        "website": "https://audilux.pe",
        "business": "LED screen rental for events in Lima and Callao; corporate events, conferences, exhibitions; Miraflores Lima Peru",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v25"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v25.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
