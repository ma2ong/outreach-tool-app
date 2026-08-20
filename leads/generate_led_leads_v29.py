"""
LED Display Leads v29 — USA (6 new entries, nos 562-567)
Aria AV, Rent For Event, LED Rentals LA, Rentex, Big Bang Companies, 4Wall Entertainment
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 29)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 562,
        "country": "USA",
        "region": "Americas",
        "company_en": "Aria AV Rentals",
        "company_local": "Aria AV",
        "city": "Nationwide (Anaheim, Atlanta, Chicago, Dallas, Denver, Las Vegas, Nashville +)",
        "contact_name": "",
        "title": "",
        "email": "Rentals@AriaAV.com",
        "phone_whatsapp": "",
        "website": "https://www.ariaav.com",
        "business": "Nationwide AV and LED display rental — video walls, monitors, LED panels for trade shows, corporate events and conferences; serves 30+ US cities including Atlanta, Chicago, Dallas, Denver, Las Vegas, Nashville; ariaav.com",
        "facebook": "",
        "instagram": "ariatechnologyrentals",
        "linkedin": "",
    },
    {
        "no": 563,
        "country": "USA",
        "region": "Americas",
        "company_en": "Rent For Event",
        "company_local": "Rent For Event",
        "city": "Los Angeles, CA (serves nationwide)",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "https://rentforevent.com",
        "business": "LED screen and video wall rental for indoor and outdoor events; serves NY, LA, Houston, Phoenix, Chicago, NJ, Dallas, Austin, Tampa; also LED display purchase and installation; rentforevent.com",
        "facebook": "rentforeventla",
        "instagram": "rentforeventus",
        "linkedin": "",
    },
    {
        "no": 564,
        "country": "USA",
        "region": "Americas",
        "company_en": "LED Rentals LA",
        "company_local": "LED Rentals LA",
        "city": "Los Angeles, CA",
        "contact_name": "",
        "title": "",
        "email": "info@ledrentalsla.com",
        "phone_whatsapp": "",
        "website": "https://ledrentalsla.com",
        "business": "LED video wall rentals in Los Angeles and Southern California; indoor and outdoor LED display solutions for corporate meetings, private events; customized solutions with 24/7 support; phone 747-262-4770",
        "facebook": "",
        "instagram": "led_rentals_la",
        "linkedin": "",
    },
    {
        "no": 565,
        "country": "USA",
        "region": "Americas",
        "company_en": "Rentex AV Rentals",
        "company_local": "Rentex",
        "city": "Boston, MA (warehouses: Las Vegas, Orlando, Dallas, Chicago, NYC, Nashville, Phoenix)",
        "contact_name": "",
        "title": "",
        "email": "info@rentex.com",
        "phone_whatsapp": "",
        "website": "https://www.rentex.com",
        "business": "B2B cross-rental partner for AV production firms — LED tiles, LED iPoster screens, Novastar and Brompton LED processors; 40+ years in business; warehouses in 7 US cities; does NOT rent to end users, only to AV production companies; rentex.com",
        "facebook": "",
        "instagram": "rentexavrentals",
        "linkedin": "",
    },
    {
        "no": 566,
        "country": "USA",
        "region": "Americas",
        "company_en": "Big Bang Companies",
        "company_local": "Big Bang Companies",
        "city": "Rochester, MN (serves all 50 states)",
        "contact_name": "Brandon",
        "title": "",
        "email": "brandon@thebigbang.us",
        "phone_whatsapp": "",
        "website": "https://www.bigbangcompanies.com",
        "business": "LED video wall rental available locally and nationwide — modular indoor/outdoor LED panels in various sizes (7x7, 10x7, 16x7 ft); stacking, truss or fly mounting; provides truss and hardware; serves all 50 states from Rochester MN base; bigbangcompanies.com",
        "facebook": "bigbangmn",
        "instagram": "bigbangmn",
        "linkedin": "",
    },
    {
        "no": 567,
        "country": "USA",
        "region": "Americas",
        "company_en": "4Wall Entertainment",
        "company_local": "4Wall, Inc.",
        "city": "Nationwide (Atlanta, Boston, Detroit, Las Vegas, LA, Miami, Nashville, NYC, Orlando, DC +)",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "https://www.4wall.com",
        "business": "Major event production company — LED video walls, lighting, audio for live events, concerts, festivals, trade shows, worship, theatre and TV/film; 12+ US locations; also offers LED display sales and systems design; 4wall.com; Instagram @4wall",
        "facebook": "4Wall.Entertainment",
        "instagram": "4wall",
        "linkedin": "4wallentertainment",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v29"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v29.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
