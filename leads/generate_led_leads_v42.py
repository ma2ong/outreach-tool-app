"""
LED Display Leads Generator v42 — USA batch (6 companies, nos 651-656)
Cities: Orlando FL, Cerritos CA, Wixom MI, Louisville KY, Omaha NE, Jacksonville FL
"""
import json, ast, re, os
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font

base = Path(__file__).parent

# --- Load all previous leads (chain) ---
_v4_src = open(base / "generate_led_leads_v4.py", encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
for _vname in ("v5","v6","v7","v8","v9","v10","v11","v12","v13","v14","v15","v16",
               "v17","v18","v19","v20","v21","v22","v23","v24","v25","v26","v27",
               "v28","v29","v30","v31","v32","v33","v34","v35","v36","v37","v38",
               "v39","v40","v41"):
    _path = base / f"generate_led_leads_{_vname}.py"
    if not _path.exists():
        continue
    _src = open(_path, encoding="utf-8").read()
    _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

new_entries = [
    {
        "no": 651,
        "country": "USA",
        "region": "Americas",
        "company_en": "GSE Audiovisual Inc",
        "company_local": "",
        "city": "Orlando, FL",
        "contact_name": "",
        "title": "",
        "email": "sales@gseav.com",
        "phone_whatsapp": "888-573-6847",
        "website": "gseav.com",
        "business": "Full-service national AV supplier: LED video walls (Absen, Unilumin, Aluvision, INFiLED, Barco), lighting & staging, video production. Offices in Orlando/Atlanta/Boston/Chicago/Denver/Hawaii/Las Vegas/Miami/NYC/NJ/New Orleans/San Antonio/San Diego/SF/DC/Vancouver. Website gseav.com",
        "facebook": "gseav",
        "instagram": "",
        "linkedin": "gse-audiovisual-inc",
    },
    {
        "no": 652,
        "country": "USA",
        "region": "Americas",
        "company_en": "Ultimate Outdoor Entertainment",
        "company_local": "",
        "city": "Cerritos, CA",
        "contact_name": "",
        "title": "",
        "email": "events@uoe.com",
        "phone_whatsapp": "888-669-1298",
        "website": "uoe.com",
        "business": "LED video wall rentals, mobile LED screen trailers, outdoor projector screens — serves nationwide including TX, CA, NV, AZ, CO, OK, AR, AL, TN, FL, NC, VA, PA, NJ, MD, DC. Instagram @ultimateoutdoorentertainment. Website uoe.com",
        "facebook": "ultimateoutdoorentertainment",
        "instagram": "ultimateoutdoorentertainment",
        "linkedin": "",
    },
    {
        "no": 653,
        "country": "USA",
        "region": "Americas",
        "company_en": "Mercury Sound & Lighting",
        "company_local": "",
        "city": "Wixom, MI",
        "contact_name": "",
        "title": "",
        "email": "info@mercurysl.com",
        "phone_whatsapp": "734-507-1177",
        "website": "mercurysl.com",
        "business": "Metro Detroit AV company: LED video wall rentals + installations, live event production, staging, audio, lighting. 20+ years experience. Also offices in Orlando FL and Coppell TX. Instagram @mercurysl. Website mercurysl.com",
        "facebook": "mercurysl",
        "instagram": "mercurysl",
        "linkedin": "",
    },
    {
        "no": 654,
        "country": "USA",
        "region": "Americas",
        "company_en": "C&H Audio Visual Services",
        "company_local": "",
        "city": "Louisville, KY",
        "contact_name": "",
        "title": "",
        "email": "rentals@chavs.net",
        "phone_whatsapp": "502-637-4595",
        "website": "chavs.net",
        "business": "AV rental and event production in Louisville KY (also serves Cincinnati and Indianapolis): LED walls, projectors, screens, sound, lighting, stage design. Instagram @chaudiovisual, FB: CHAudioVisuals. Website chavs.net",
        "facebook": "CHAudioVisuals",
        "instagram": "chaudiovisual",
        "linkedin": "",
    },
    {
        "no": 655,
        "country": "USA",
        "region": "Americas",
        "company_en": "Nova Productions LLC",
        "company_local": "",
        "city": "Omaha, NE",
        "contact_name": "",
        "title": "",
        "email": "info@novaproductionsomaha.com",
        "phone_whatsapp": "402-208-8536",
        "website": "novaomaha.com",
        "business": "Omaha event production: LED video wall rentals, portable stages, audio, lighting. Serves Midwest including Iowa, South Dakota, Kansas, Missouri. Instagram @nova_productions_llc, FB: novaproductionsomaha. Website novaomaha.com",
        "facebook": "novaproductionsomaha",
        "instagram": "nova_productions_llc",
        "linkedin": "",
    },
    {
        "no": 656,
        "country": "USA",
        "region": "Americas",
        "company_en": "PRI Productions",
        "company_local": "",
        "city": "Jacksonville, FL",
        "contact_name": "",
        "title": "",
        "email": "info@priproductions.com",
        "phone_whatsapp": "904-398-8179",
        "website": "priproductions.com",
        "business": "Jacksonville FL full-service event production: LED video wall rentals (indoor/outdoor), concert production, AV technology, trade shows. Serves entire Southeast: FL, GA, NC. Instagram @priproductions, FB: PRIProductions. Website priproductions.com",
        "facebook": "PRIProductions",
        "instagram": "priproductions",
        "linkedin": "",
    },
]

leads.extend(new_entries)

# --- Print summary ---
print(f"Total leads: {len(leads)}")
usa = [l for l in leads if l.get("country") == "USA"]
print(f"USA leads: {len(usa)}")
print(f"New entries: {len(new_entries)} (nos {new_entries[0]['no']}–{new_entries[-1]['no']})")

# --- Export to Excel ---
COUNTRY_COLORS = {
    "USA": "DDEEFF",
    "Korea": "FFE8CC",
    "Brazil": "E8F5E9",
    "Canada": "FFF9C4",
    "Colombia": "FCE4EC",
    "Mexico": "F3E5F5",
    "Chile": "E0F7FA",
    "Argentina": "FBE9E7",
    "Peru": "F1F8E9",
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads v42"

headers = ["No", "Country", "Region", "Company EN", "Company Local", "City",
           "Contact Name", "Title", "Email", "Phone/WhatsApp",
           "Website", "Business", "Facebook", "Instagram", "LinkedIn"]
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="334455")
    cell.font = Font(bold=True, color="FFFFFF")

for lead in leads:
    row = [
        lead.get("no"), lead.get("country"), lead.get("region"),
        lead.get("company_en"), lead.get("company_local"), lead.get("city"),
        lead.get("contact_name"), lead.get("title"), lead.get("email"),
        lead.get("phone_whatsapp"), lead.get("website"), lead.get("business"),
        lead.get("facebook"), lead.get("instagram"), lead.get("linkedin"),
    ]
    ws.append(row)
    color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=color)

ws.column_dimensions["A"].width = 6
ws.column_dimensions["D"].width = 30
ws.column_dimensions["F"].width = 25
ws.column_dimensions["I"].width = 32
ws.column_dimensions["J"].width = 20
ws.column_dimensions["K"].width = 28
ws.column_dimensions["L"].width = 60

out = base / "LED_Display_Leads_v42.xlsx"
wb.save(out)
print(f"Saved: {out}")
