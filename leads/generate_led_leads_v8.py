"""LED Display Leads — Americas + Korea — v8
Extends v7 (358 records) with 14 new entries = 372 total.
Sources: Instagram/Facebook opencli + WebSearch + Jina/WebFetch
New: Korea +9, Brazil +2, Mexico +1, Chile +1, Peru +1
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
_leads_src = re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1)
leads = ast.literal_eval(_leads_src)

# ── Load v5 / v6 / v7 new_entries ─────────────────────────────────────────────
for _vname in ("v5", "v6", "v7"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (359 – 372) ────────────────────────────────────────────────────
new_entries = [
    # ==================== KOREA new (9) ====================
    {
        "no":359,"country":"Korea","region":"Asia",
        "company_en":"Korea AV Solutions","company_local":"코리아 AV 솔루션",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2-776-3333",
        "website":"",
        "business":"LED 전광판 + 미디어 파사드 전문기업; media facade, large outdoor LED, AV integration — Instagram: @av_korea",
        "facebook":"","instagram":"av_korea","linkedin":"",
    },
    {
        "no":360,"country":"Korea","region":"Asia",
        "company_en":"Asone Tech","company_local":"주식회사 애즈원",
        "city":"Seoul (Geumcheon-gu)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2-2088-0920",
        "website":"http://www.asonetech.co.kr",
        "business":"LED display total solution; 9-year Novastar official distributor; 1,900 installs; public-sector top bidder; KakaoTalk: pf.kakao.com/_xaqZKs",
        "facebook":"","instagram":"asonetech","linkedin":"",
    },
    {
        "no":361,"country":"Korea","region":"Asia",
        "company_en":"DisplayHub","company_local":"디스플레이허브",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2-546-3288",
        "website":"",
        "business":"Large LED display specialist; all-in-one service (install to maintenance) since 2010 — Instagram: @displayhub_ltd",
        "facebook":"","instagram":"displayhub_ltd","linkedin":"",
    },
    {
        "no":362,"country":"Korea","region":"Asia",
        "company_en":"OnBit LED","company_local":"온빛전자",
        "city":"Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 31-975-5689",
        "website":"",
        "business":"LED 전광판 전문기업; outdoor/indoor LED displays, nationwide installation — Instagram: @onbitled",
        "facebook":"","instagram":"onbitled","linkedin":"",
    },
    {
        "no":363,"country":"Korea","region":"Asia",
        "company_en":"Gadjet Networks","company_local":"(주)가제트네트웍스",
        "city":"Deogyang-gu, Gyeonggi-do",
        "contact_name":"정청수","title":"대표",
        "email":"zoro0309@naver.com","phone_whatsapp":"+82 2-2012-2142",
        "website":"https://gadjetnet.com",
        "business":"LED display + video wall + DID; media art production; total media solution; rental & permanent — Instagram: @led_gadjet_networks",
        "facebook":"","instagram":"led_gadjet_networks","linkedin":"",
    },
    {
        "no":364,"country":"Korea","region":"Asia",
        "company_en":"BasicTech","company_local":"(주)베이직테크",
        "city":"Paju, Gyeonggi-do",
        "contact_name":"윤영산","title":"대표",
        "email":"","phone_whatsapp":"+82 70-7004-2200",
        "website":"http://www.basictech.co.kr",
        "business":"LED display / media facade / media art / XR-VP studio; development, manufacturing, rental; visual solution specialist",
        "facebook":"","instagram":"basictech_official","linkedin":"",
    },
    {
        "no":365,"country":"Korea","region":"Asia",
        "company_en":"Display RCH","company_local":"디스플레이RCH",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"http://www.ricaled.com",
        "business":"LED display manufacturer; outdoor modules, media poles, full-color panels — Instagram: @rch_led",
        "facebook":"","instagram":"rch_led","linkedin":"",
    },
    {
        "no":366,"country":"Korea","region":"Asia",
        "company_en":"Seonggwang Electric Lighting","company_local":"성광전기조명",
        "city":"Daegu (Buk-gu)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 53-428-3864",
        "website":"https://seonggwang.iotled.co.kr",
        "business":"LED display & lighting specialist; outdoor signage, indoor panels, commercial installs; Daegu/Gyeongbuk region",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":367,"country":"Korea","region":"Asia",
        "company_en":"Vision Catcher","company_local":"비전캐처",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 1644-0271",
        "website":"",
        "business":"LED display + AV + lighting rental & integration; concerts, events, corporate; sound/light/LED/video total service — Instagram: @visioncatcher_rental",
        "facebook":"","instagram":"visioncatcher_rental","linkedin":"",
    },
    # ==================== BRAZIL new (2) ====================
    {
        "no":368,"country":"Brazil","region":"Americas",
        "company_en":"LED 360","company_local":"Led 360",
        "city":"São Paulo, SP",
        "contact_name":"Fábio","title":"",
        "email":"fabio@led360.com.br","phone_whatsapp":"+55 11 95858-9230",
        "website":"https://led360.com.br",
        "business":"LED panel specialist; rental & sale; indoor/outdoor P2/P3/P4; São Paulo; WhatsApp: +55 11 95858-9230 (confirmed on website)",
        "facebook":"","instagram":"painelled360","linkedin":"",
    },
    {
        "no":369,"country":"Brazil","region":"Americas",
        "company_en":"LED Expert","company_local":"LED Expert",
        "city":"Montes Claros, MG",
        "contact_name":"","title":"",
        "email":"ledexpertmoc@ledexpert.com.br","phone_whatsapp":"+55 38 3222-0500",
        "website":"https://ledexpert.com.br",
        "business":"LED panel manufacturer/integrator; 15+ years IoT/Industry 5.0; ISO/UL/CE/TUV certified; 5-year warranty; branches in BH, SP, RJ; national support",
        "facebook":"","instagram":"ledexpertoficial","linkedin":"",
    },
    # ==================== MEXICO new (1) ====================
    {
        "no":370,"country":"Mexico","region":"Americas",
        "company_en":"Fénix Evolution","company_local":"Producciones Fénix Evolution",
        "city":"León, Guanajuato",
        "contact_name":"","title":"",
        "email":"ventas@fenix-evolution.com.mx","phone_whatsapp":"+52 477-125-5013",
        "website":"https://fenix-evolution.com.mx",
        "business":"LED screen rental/sales; indoor/outdoor; events, expos, congresses; León + Irapuato base; nationwide service; WhatsApp León: +52 477-125-5013",
        "facebook":"produccionesfenixevolution","instagram":"fenixevolution_led","linkedin":"",
    },
    # ==================== CHILE new (1) ====================
    {
        "no":371,"country":"Chile","region":"Americas",
        "company_en":"Liu King LED Screen Chile","company_local":"Liu King",
        "city":"Chile",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Wholesale LED panel distributor; P2/P3/P4/P5 indoor/outdoor; pre-sale & stock for immediate delivery; nationwide Chile — Instagram: @liukingledscreenchile",
        "facebook":"","instagram":"liukingledscreenchile","linkedin":"",
    },
    # ==================== PERU new (1) ====================
    {
        "no":372,"country":"Peru","region":"Americas",
        "company_en":"LED Maker Peru","company_local":"Led Maker - Pantallas Led Perú",
        "city":"Lima",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display importer & distributor; sale & installation; indoor/outdoor panels; Lima, Peru — Facebook: @LEDmakerPeru (1,294 followers)",
        "facebook":"LEDmakerPeru","instagram":"","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

# ── Sheet 1: All Markets ───────────────────────────────────────────────────────
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn",
    "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font      = header_font
    cell.fill      = header_fill
    cell.alignment = header_align
    cell.border    = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""),
        lead.get("region",""),
        lead.get("country",""),
        lead.get("city",""),
        lead.get("company_en",""),
        lead.get("company_local",""),
        lead.get("contact_name",""),
        lead.get("title",""),
        lead.get("email",""),
        lead.get("phone_whatsapp",""),
        wa_flag,
        kk_flag,
        lead.get("website",""),
        lead.get("business",""),
        lead.get("facebook",""),
        lead.get("instagram",""),
        lead.get("linkedin",""),
        "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill      = row_fill
        cell.border    = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font      = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14

ws2["A1"] = "LED Display Leads — v8 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")

ws2["A3"] = "Country"
ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font      = Font(bold=True, color="FFFFFF")
    cell.fill      = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border    = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea", "USA", "Brazil", "Canada", "Chile", "Argentina", "Colombia", "Peru", "Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    ws2.cell(row=r, column=2, value=counts.get(country, 0)).border = thin_border
    ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    r += 1

ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
ws2.cell(row=r, column=2, value=len(leads)).font = Font(bold=True)
ws2.cell(row=r, column=2).border = thin_border
ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")

ws2["A" + str(r+2)] = f"Last updated: 2026-05-06  |  Total: {len(leads)} records"
ws2["A" + str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60

ws3["A1"] = "市场开发建议 — v8"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")

advice = [
    ("Korea", "优先Kakao联系，邮件/电话备用。公共事业/政府项目多，重点推系统集成方案。애즈원/케이시스等有Kakao官方账号，可直接消息开发。"),
    ("USA", "优先WhatsApp/邮件。重点推大屏广告、活动租赁和体育场馆方向。强调交期+美标认证。XcolorLED/TechLed适合做经销商。"),
    ("Brazil", "多数公司留有WhatsApp，直接发WA最有效。LED Expert有全国网点(BH/SP/RJ/MG)，适合做区域经销商。ledwave.com.br也是大公司。"),
    ("Canada", "邮件为主，电话跟进。重点推零售/企业室内屏和政府招标。Inno-Leader有Shenzhen背景，沟通成本低。"),
    ("Chile", "邮件+WhatsApp。刘王(Liu King)是批发型分销商，适合报大批量价格。圣地亚哥活动/广告市场优先。"),
    ("Argentina", "通货膨胀环境，买家更关注性价比和分期。优先推标准现货型号，避免定制。"),
    ("Colombia", "Bogotá市场最大，Medellín其次。活动/租赁方向需求活跃。西班牙语开发信效果更好。WhatsApp是主流。"),
    ("Peru", "Lima为核心市场，零售广告+事件屏需求为主。WhatsApp是主流沟通工具。LED Maker从Instagram/FB引流。"),
    ("Mexico", "CDMX市场最大，Guadalajara/Monterrey/León其次。Fénix Evolution专注活动租赁，适合推租赁产品线。DMX Technologies有20年，适合长期合作。"),
]

ws3["A3"] = "国家"
ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font   = Font(bold=True, color="FFFFFF")
    cell.fill   = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border

for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v8.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
