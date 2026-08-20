"""
LED Display Leads v41 — USA (8 new entries, nos 643-650)
New cities/regions: DC/Mid-Atlantic (x3), SF Bay Area (x3), New Orleans (x1), San Francisco (x1)
Channels: Email (7), IG DM (3: avactions_dc, amospro_av, corplighting), FB (3)
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 41)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 643,
        "country": "USA",
        "region": "Americas",
        "company_en": "Rent LED Video Walls",
        "company_local": "",
        "city": "Annapolis, MD",
        "contact_name": "",
        "title": "",
        "email": "info@rentledvideowall.com",
        "phone_whatsapp": "+1 443-333-9972",
        "website": "https://rentledvideowall.com",
        "business": "Full-service LED video wall rental for concerts, conferences, corporate events and permanent installs; serving Maryland, Washington DC and Virginia; Annapolis-based. Email confirmed from contact page.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 644,
        "country": "USA",
        "region": "Americas",
        "company_en": "Media Support Services Inc",
        "company_local": "",
        "city": "Baltimore, MD",
        "contact_name": "",
        "title": "",
        "email": "info@mssav.com",
        "phone_whatsapp": "+1 410-669-1100",
        "website": "https://mssav.com",
        "business": "AV rental company serving DC/MD/VA/DE/NJ/NY/PA since 1996; LED screens, projectors, staging, corporate events; two offices: Baltimore (410-669-1100) and Washington DC (202-990-8557). Email confirmed from website.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 645,
        "country": "USA",
        "region": "Americas",
        "company_en": "AV Actions Inc",
        "company_local": "",
        "city": "Alexandria, VA",
        "contact_name": "",
        "title": "",
        "email": "info@avactions.com",
        "phone_whatsapp": "+1 703-751-1010",
        "website": "https://avactions.com",
        "business": "Full-service AV rental and LED wall production; offices in Alexandria VA, Washington DC, Rockville MD, and Orlando FL; corporate events, trade shows. Instagram @avactions_dc. Email confirmed from contact page.",
        "facebook": "",
        "instagram": "avactions_dc",
        "linkedin": "",
    },
    {
        "no": 646,
        "country": "USA",
        "region": "Americas",
        "company_en": "Stage Lights and Sound",
        "company_local": "",
        "city": "Richmond, CA",
        "contact_name": "",
        "title": "",
        "email": "Sales@StageLightsandSound.com",
        "phone_whatsapp": "+1 415-652-0080",
        "website": "https://stagelightsandsound.com",
        "business": "Full-service event production company; LED video wall rental (Absen LED), staging, audio, lighting; serving Northern California and SF Bay Area; Richmond CA base. Email confirmed from website.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 647,
        "country": "USA",
        "region": "Americas",
        "company_en": "Amos Productions",
        "company_local": "",
        "city": "Livermore, CA",
        "contact_name": "",
        "title": "",
        "email": "info@amospro.com",
        "phone_whatsapp": "+1 925-449-3847",
        "website": "https://amospro.com",
        "business": "Full-service AV production; LED video walls, live sound, lighting, DJ, photo booth, video production; serving SF Bay Area, Silicon Valley, and Napa Valley; Livermore CA base. Instagram @amospro_av; Facebook amosproductions. Email confirmed from contact page.",
        "facebook": "amosproductions",
        "instagram": "amospro_av",
        "linkedin": "",
    },
    {
        "no": 648,
        "country": "USA",
        "region": "Americas",
        "company_en": "Megahertz AV",
        "company_local": "",
        "city": "Santa Clara, CA",
        "contact_name": "",
        "title": "",
        "email": "info@mhzav.com",
        "phone_whatsapp": "+1 650-641-7135",
        "website": "https://mhzav.com",
        "business": "AV equipment rental company; LED video wall rental in San Francisco; offices in Santa Clara and San Francisco; corporate events, conferences, trade shows. Facebook: mhzent. Email confirmed from contact page.",
        "facebook": "mhzent",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 649,
        "country": "USA",
        "region": "Americas",
        "company_en": "Corporate Lighting and Audio",
        "company_local": "",
        "city": "New Orleans, LA",
        "contact_name": "",
        "title": "",
        "email": "sales@corplighting.com",
        "phone_whatsapp": "+1 504-837-3947",
        "website": "https://corplighting.com",
        "business": "National production company based in New Orleans; LED video wall rental, audio, lighting, stage design for conventions, concerts, corporate events; 5207 River Rd. Instagram @corplighting; Facebook corplighting. Email confirmed from contact page.",
        "facebook": "corplighting",
        "instagram": "corplighting",
        "linkedin": "",
    },
    {
        "no": 650,
        "country": "USA",
        "region": "Americas",
        "company_en": "Bay Area Lighting and Sound",
        "company_local": "",
        "city": "San Francisco, CA",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+1 866-767-7623",
        "website": "https://bayarealightingandsound.com",
        "business": "Full AV production and equipment rental; LED video walls (indoor and outdoor), stage, sound, lighting, truss; 30 years experience in SF Bay Area and San Jose area. Phone confirmed from website.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]


if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    all_leads = load_all_leads()
    all_leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads"
    headers = ["no","country","region","company_en","company_local","city","contact_name","title",
               "email","phone_whatsapp","website","business","facebook","instagram","linkedin"]
    ws.append(headers)

    for lead in all_leads:
        row = [lead.get(h, "") for h in headers]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for cell in ws[ws.max_row]:
            cell.fill = fill

    out = os.path.join(BASE, "LED_Display_Leads_v41.xlsx")
    wb.save(out)
    print(f"Saved {out} with {len(all_leads)} leads")
    usa = [l for l in all_leads if l.get("country") == "USA"]
    print(f"USA total: {len(usa)}")
