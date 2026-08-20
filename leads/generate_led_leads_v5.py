"""LED Display Leads — Americas + Korea — v5
Extends v4 (209 records) with 91 new entries = 300 total.
New: Colombia +10, USA +30, Canada +31, Brazil +3, Chile +2, Korea +15
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads without re-running the workbook code ────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
_leads_src = re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1)
leads = ast.literal_eval(_leads_src)

# ── New entries (210 – 300) ────────────────────────────────────────────────────
new_entries = [
    # ==================== COLOMBIA new (10) ====================
    {
        "no":210,"country":"Colombia","region":"South America",
        "company_en":"Pantallas LED Colombia","company_local":"Pantallas LED Colombia",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"info@pantallasledcolombia.com","phone_whatsapp":"+57 321 4854822",
        "website":"https://pantallasledcolombia.com",
        "business":"LED screen sales & rental; indoor/outdoor advertising",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":211,"country":"Colombia","region":"South America",
        "company_en":"Gescom S.A.S.","company_local":"Gescom S.A.S.",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"info@gescom.com.co","phone_whatsapp":"+57 318 856 4608",
        "website":"https://gescom.com.co",
        "business":"LED display solutions & digital signage",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":212,"country":"Colombia","region":"South America",
        "company_en":"OOH Digital Networks","company_local":"OOH Digital Networks",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"info@oohrd.com","phone_whatsapp":"+57 601 6954409",
        "website":"https://oohrd.com",
        "business":"OOH digital LED advertising networks",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":213,"country":"Colombia","region":"South America",
        "company_en":"PubliMedia Colombia","company_local":"PubliMedia",
        "city":"Bogotá",
        "contact_name":"J. González","title":"",
        "email":"jgonzalez@publimedia.com.co","phone_whatsapp":"",
        "website":"https://publimedia.com.co",
        "business":"LED media & outdoor advertising",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":214,"country":"Colombia","region":"South America",
        "company_en":"Marketmedios","company_local":"Marketmedios",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"info@marketmedios.com.co","phone_whatsapp":"+571 635 0650",
        "website":"https://marketmedios.com.co",
        "business":"LED outdoor billboard & digital media",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":215,"country":"Colombia","region":"South America",
        "company_en":"ExpoRed Colombia","company_local":"ExpoRed",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"hola@expo.red","phone_whatsapp":"+57 300 2224957",
        "website":"https://expo.red",
        "business":"LED display & event technology",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":216,"country":"Colombia","region":"South America",
        "company_en":"Kender SAS","company_local":"Kender SAS",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"contacto@kender.com.co","phone_whatsapp":"+57 314 3584459",
        "website":"https://kender.com.co",
        "business":"LED display & technology solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":217,"country":"Colombia","region":"South America",
        "company_en":"Tekus Colombia","company_local":"Tekus",
        "city":"Medellín",
        "contact_name":"","title":"",
        "email":"hello@tekus.co","phone_whatsapp":"+57 323 252 9964",
        "website":"https://tekus.co",
        "business":"LED display & tech integration",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":218,"country":"Colombia","region":"South America",
        "company_en":"VCR Ltda.","company_local":"VCR Ltda.",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"info@vcr.com.co","phone_whatsapp":"+601 744 6629",
        "website":"https://vcr.com.co",
        "business":"LED display & visual communication solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":219,"country":"Colombia","region":"South America",
        "company_en":"LED MASTER Colombia","company_local":"The LED Master",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://theledmaster.com",
        "business":"LED display sales & installation",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== USA new (30) ====================
    {
        "no":220,"country":"USA","region":"North America",
        "company_en":"Planar Systems","company_local":"Planar Systems",
        "city":"Hillsboro, OR",
        "contact_name":"","title":"",
        "email":"info@planar.com","phone_whatsapp":"+1 503-748-5500",
        "website":"https://www.planar.com",
        "business":"Fine-pitch LED & LCD video wall displays; corporate/retail/broadcast",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":221,"country":"USA","region":"North America",
        "company_en":"Daktronics","company_local":"Daktronics",
        "city":"Brookings, SD",
        "contact_name":"","title":"",
        "email":"contact@daktronics.com","phone_whatsapp":"+1 800-325-8766",
        "website":"https://www.daktronics.com",
        "business":"LED scoreboards/billboards/displays; sports/outdoor/transportation",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":222,"country":"USA","region":"North America",
        "company_en":"Adscope Media","company_local":"Adscope Media",
        "city":"Marine City, MI",
        "contact_name":"","title":"",
        "email":"info@adscopemedia.com","phone_whatsapp":"+1 586-337-6500",
        "website":"https://adscopemedia.com",
        "business":"LED display sales & outdoor advertising media",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":223,"country":"USA","region":"North America",
        "company_en":"BASS Ltd","company_local":"BASS Ltd",
        "city":"Broussard, LA",
        "contact_name":"","title":"",
        "email":"info@bassltd.com","phone_whatsapp":"+1 337-981-1189",
        "website":"https://bassltd.com",
        "business":"LED display & AV solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":224,"country":"USA","region":"North America",
        "company_en":"Sigmax Media","company_local":"Sigmax Corp.",
        "city":"Santa Fe Springs, CA",
        "contact_name":"","title":"",
        "email":"sales@sigmaxcorp.com","phone_whatsapp":"",
        "website":"https://sigmaxcorp.com",
        "business":"LED display & digital signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":225,"country":"USA","region":"North America",
        "company_en":"Wilson Electronic Displays","company_local":"WED LED",
        "city":"Dayton, OH",
        "contact_name":"","title":"",
        "email":"info@wedled.com","phone_whatsapp":"+1 937-558-2416",
        "website":"https://wedled.com",
        "business":"LED display manufacturer & distributor",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":226,"country":"USA","region":"North America",
        "company_en":"Vu Volumes","company_local":"Vu Volumes",
        "city":"Tampa, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 888-575-1510",
        "website":"https://vuvolumes.com",
        "business":"LED display rental & event production",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":227,"country":"USA","region":"North America",
        "company_en":"E.C. Pro Video Systems","company_local":"EC Pro Store",
        "city":"New York, NY",
        "contact_name":"","title":"",
        "email":"sales@ecprostore.com","phone_whatsapp":"+1 212-333-5570",
        "website":"https://ecprostore.com",
        "business":"LED display & professional video systems",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":228,"country":"USA","region":"North America",
        "company_en":"LED Nation USA","company_local":"LED Nation USA",
        "city":"Doral, FL",
        "contact_name":"","title":"",
        "email":"info@lednationusa.com","phone_whatsapp":"+1 888-590-1720",
        "website":"https://lednationusa.com",
        "business":"LED display wholesale & distribution (Florida/LatAm hub)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":229,"country":"USA","region":"North America",
        "company_en":"Silicon Core","company_local":"Silicon-Core",
        "city":"Milpitas, CA",
        "contact_name":"","title":"",
        "email":"sales@silicon-core.com","phone_whatsapp":"",
        "website":"https://silicon-core.com",
        "business":"Fine-pitch indoor LED display manufacturer",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":230,"country":"USA","region":"North America",
        "company_en":"Jumbotron USA","company_local":"Jumbotron.com",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"info@jumbotron.com","phone_whatsapp":"+1 800-940-0922",
        "website":"https://jumbotron.com",
        "business":"Large LED display & video wall solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":231,"country":"USA","region":"North America",
        "company_en":"Vision Tech LED","company_local":"Vision Tech LED Displays",
        "city":"Springfield, MO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-927-1778",
        "website":"https://visiontechleddisplays.com",
        "business":"LED display manufacturer & reseller",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":232,"country":"USA","region":"North America",
        "company_en":"Trans-Lux Corporation","company_local":"Trans-Lux",
        "city":"Hazelwood, MO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-243-5544",
        "website":"https://trans-lux.com",
        "business":"LED display & ticker sign manufacturer (80+ years)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":233,"country":"USA","region":"North America",
        "company_en":"Hyoco Distribution","company_local":"Hyoco Distribution",
        "city":"Irvine, CA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 888-860-2249",
        "website":"https://hyocodistribution.com",
        "business":"LED display distributor (Chinese brand importer to US)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":234,"country":"USA","region":"North America",
        "company_en":"USA LED Sign","company_local":"USA LED Sign",
        "city":"Commerce, CA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 817-385-1028",
        "website":"https://usaledsign.com",
        "business":"LED sign & display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":235,"country":"USA","region":"North America",
        "company_en":"Vibrant Displays","company_local":"Vibrant Displays",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 855-527-4448",
        "website":"https://vibrantdisplays.com",
        "business":"LED display solutions & digital signage",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":236,"country":"USA","region":"North America",
        "company_en":"APG Displays","company_local":"APG Displays",
        "city":"Orlando, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-829-0593",
        "website":"https://apgdisplays.com",
        "business":"LED display manufacturer & integrator",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":237,"country":"USA","region":"North America",
        "company_en":"Xtreme LED Screens","company_local":"Xtreme LED Screens",
        "city":"Austin, TX",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 866-774-5196",
        "website":"https://xtremeledscreens.com",
        "business":"LED screen rental & sales; events & fixed install",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":238,"country":"USA","region":"North America",
        "company_en":"Numeritex Displays","company_local":"Numeritex",
        "city":"Murray, KY",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 270-753-3740",
        "website":"https://numeritex.com",
        "business":"LED display & digital signage manufacturer",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":239,"country":"USA","region":"North America",
        "company_en":"Pixel Wall Inc.","company_local":"PIXW",
        "city":"Nashville, TN",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 703-594-1288",
        "website":"https://pixw.us",
        "business":"LED video wall solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":240,"country":"USA","region":"North America",
        "company_en":"Sign-Tech","company_local":"Sign-Tech",
        "city":"Paragould, AR",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 870-239-3344",
        "website":"https://sign-tech.com",
        "business":"LED sign & display manufacturer",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":241,"country":"USA","region":"North America",
        "company_en":"LED Market USA","company_local":"LED Market USA",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://ledmarketusa.com",
        "business":"LED display market & distribution",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":242,"country":"USA","region":"North America",
        "company_en":"Refresh LED","company_local":"Refresh LED",
        "city":"PA, USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://refreshled.com",
        "business":"LED display solutions & service",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":243,"country":"USA","region":"North America",
        "company_en":"Mobile View Screens","company_local":"Mobile View Screens",
        "city":"Larkspur, CO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 303-770-3416",
        "website":"",
        "business":"Mobile LED screen rental for events",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":244,"country":"USA","region":"North America",
        "company_en":"Sunshine Electronic Display","company_local":"Sunshine Electronic Display",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-821-9013",
        "website":"https://sunshine.us.com",
        "business":"LED display manufacturer & distributor",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":245,"country":"USA","region":"North America",
        "company_en":"Boostr Displays","company_local":"Boostr Displays",
        "city":"AL, USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 205-523-4799",
        "website":"",
        "business":"LED display sales & installation",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":246,"country":"USA","region":"North America",
        "company_en":"TrailexLED","company_local":"TrailexLED",
        "city":"Canfield, OH",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Trailer-mounted LED screen rental",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":247,"country":"USA","region":"North America",
        "company_en":"Little Mountain Display","company_local":"Little Mountain Display",
        "city":"Tulsa, OK",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 918-583-7450",
        "website":"",
        "business":"LED display sales & installation (Tulsa region)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":248,"country":"USA","region":"North America",
        "company_en":"Insane Impact","company_local":"Insane Impact",
        "city":"Clive, IA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 855-923-4883",
        "website":"",
        "business":"LED display rental & event production",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":249,"country":"USA","region":"North America",
        "company_en":"Worship Productions","company_local":"Worship Productions",
        "city":"Brea, CA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display & AV for houses of worship",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== CANADA new (31) ====================
    {
        "no":250,"country":"Canada","region":"North America",
        "company_en":"FunFlicks Canada","company_local":"FunFlicks",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"info@funflicks.com","phone_whatsapp":"+1 877-263-0480",
        "website":"https://funflicks.com",
        "business":"LED screen rental & outdoor movie events",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":251,"country":"Canada","region":"North America",
        "company_en":"Leader LED","company_local":"Leader LED",
        "city":"Markham, ON",
        "contact_name":"","title":"",
        "email":"info@leaderled.ca","phone_whatsapp":"+1 647-395-8882",
        "website":"https://leaderled.ca",
        "business":"LED display sales & installation",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":252,"country":"Canada","region":"North America",
        "company_en":"Daktronics Canada","company_local":"Daktronics Canada",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-800-3258",
        "website":"https://www.daktronics.com",
        "business":"LED scoreboards, billboards, displays; sports/outdoor/transit",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":253,"country":"Canada","region":"North America",
        "company_en":"Netvisual","company_local":"Netvisual",
        "city":"Oakville, ON",
        "contact_name":"","title":"",
        "email":"info@netvisual.ca","phone_whatsapp":"+1 416-850-5059",
        "website":"https://netvisual.ca",
        "business":"LED display & digital signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":254,"country":"Canada","region":"North America",
        "company_en":"UTG Digital Media","company_local":"UTG Digital Media",
        "city":"Ottawa, ON",
        "contact_name":"","title":"",
        "email":"info@utgdm.com","phone_whatsapp":"+1 613-695-5550",
        "website":"https://utgdm.com",
        "business":"LED display & digital signage network management",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":255,"country":"Canada","region":"North America",
        "company_en":"LEDLE MEDIA","company_local":"LED Sign Toronto",
        "city":"Toronto, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 416-970-1983",
        "website":"https://ledsigntoronto.com",
        "business":"LED sign & display solutions (Toronto)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":256,"country":"Canada","region":"North America",
        "company_en":"Media Resources International","company_local":"Media Resources International",
        "city":"Oakville, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-667-4554",
        "website":"https://mediaresources.com",
        "business":"LED display & professional AV media solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":257,"country":"Canada","region":"North America",
        "company_en":"NUMMAX","company_local":"NUMMAX",
        "city":"Quebec City, QC",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 877-255-3471",
        "website":"https://nummax.com/en",
        "business":"LED display & digital signage manufacturer (Quebec)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":258,"country":"Canada","region":"North America",
        "company_en":"Blanchett Neon","company_local":"Blanchett Neon",
        "city":"Edmonton, AB",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 780-453-2441",
        "website":"https://blanchettneon.com",
        "business":"LED sign & display manufacturing (Alberta)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":259,"country":"Canada","region":"North America",
        "company_en":"Blue Leo Technologies","company_local":"Blue Leo Technologies",
        "city":"Kitchener, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 647-879-0944",
        "website":"https://blueleotech.com",
        "business":"LED display & digital signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":260,"country":"Canada","region":"North America",
        "company_en":"Genoptic Smart Displays","company_local":"Genoptic / LED Sign Supply",
        "city":"Calgary, AB",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 403-726-9260",
        "website":"https://ledsignsupply.com",
        "business":"LED sign supply & smart display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":261,"country":"Canada","region":"North America",
        "company_en":"Libertevision","company_local":"Libertévision",
        "city":"Sherbrooke, QC",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 855-437-0022",
        "website":"https://libertevision.com",
        "business":"LED display & digital signage (Quebec)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":262,"country":"Canada","region":"North America",
        "company_en":"eSigns Canada","company_local":"eSigns Canada",
        "city":"Bradford, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 905-775-9112",
        "website":"https://esignscanada.ca",
        "business":"LED sign & display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":263,"country":"Canada","region":"North America",
        "company_en":"Digital Edge LED Screens","company_local":"DELED",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 321-220-1713",
        "website":"https://deled.ca",
        "business":"LED screen solutions & digital signage",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":264,"country":"Canada","region":"North America",
        "company_en":"DYNAMIX","company_local":"DYNAMIX",
        "city":"Richmond Hill, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 519-681-5000",
        "website":"https://dynamix.ca",
        "business":"LED display & AV integration (Ontario)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":265,"country":"Canada","region":"North America",
        "company_en":"ColossoVision","company_local":"ColossoVision",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 855-855-1030",
        "website":"https://colossovision.ca",
        "business":"LED display & large-format visual solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":266,"country":"Canada","region":"North America",
        "company_en":"Pure AV","company_local":"Pure AV",
        "city":"Mississauga, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 800-929-7089",
        "website":"https://pureav.ca",
        "business":"LED display & professional AV solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":267,"country":"Canada","region":"North America",
        "company_en":"Bannerz Canada","company_local":"Bannerz Canada",
        "city":"Edmonton, AB",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 780-989-1190",
        "website":"https://bannerz.ca",
        "business":"LED sign & banner display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":268,"country":"Canada","region":"North America",
        "company_en":"Pro-LED Sign","company_local":"LED Sign 4 Life",
        "city":"Mississauga, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 888-880-8524",
        "website":"https://ledsign4life.ca",
        "business":"LED sign & display solutions (Ontario)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":269,"country":"Canada","region":"North America",
        "company_en":"RGB Sign & Print","company_local":"RGB Sign",
        "city":"Scarborough, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 416-498-5555",
        "website":"https://led-sign.ca",
        "business":"LED sign & printing solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":270,"country":"Canada","region":"North America",
        "company_en":"Telematics Canada","company_local":"Telematics Canada",
        "city":"Vancouver, BC",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 604-331-8795",
        "website":"https://telematics.ca",
        "business":"LED display & telematics solutions (BC)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":271,"country":"Canada","region":"North America",
        "company_en":"GTR Industries","company_local":"GTR Industries",
        "city":"Toronto, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 416-665-4209",
        "website":"https://gtrdirect.ca",
        "business":"LED display & digital signage",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":272,"country":"Canada","region":"North America",
        "company_en":"Digital Edge Media","company_local":"Digital Edge Media",
        "city":"Sherwood Park, AB",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 866-727-4543",
        "website":"https://digitaledgemedia.ca",
        "business":"LED display & digital media solutions (Alberta)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":273,"country":"Canada","region":"North America",
        "company_en":"Everbright Media","company_local":"Everbright LED",
        "city":"BC, Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 778-320-5255",
        "website":"https://everbrightled.ca",
        "business":"LED display & media solutions (BC)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":274,"country":"Canada","region":"North America",
        "company_en":"Adtronics","company_local":"Adtronics",
        "city":"Delta, BC",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 604-940-8696",
        "website":"https://adtronics.net",
        "business":"LED display & digital advertising solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":275,"country":"Canada","region":"North America",
        "company_en":"LIGHTVU","company_local":"LIGHTVU",
        "city":"Sherwood Park, AB",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 855-514-4888",
        "website":"https://lightvu.com",
        "business":"LED display & lighting visual solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":276,"country":"Canada","region":"North America",
        "company_en":"GreenTak Canada","company_local":"GreenTak",
        "city":"Mississauga, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 905-212-9776",
        "website":"https://greentak.com",
        "business":"LED display & green technology solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":277,"country":"Canada","region":"North America",
        "company_en":"LED Solutions Canada","company_local":"LED Solutions",
        "city":"Hamilton, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 866-768-1533",
        "website":"https://ledsolutions.ca",
        "business":"LED display solutions (Ontario)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":278,"country":"Canada","region":"North America",
        "company_en":"United Digitals","company_local":"United Digitals",
        "city":"Toronto, ON",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 416-663-4500",
        "website":"https://uniteddigitals.com",
        "business":"LED display & digital signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":279,"country":"Canada","region":"North America",
        "company_en":"ClearLED Displays","company_local":"ClearLED",
        "city":"Vancouver, BC",
        "contact_name":"","title":"",
        "email":"sales@clearleddisplays.com","phone_whatsapp":"+1 778-373-5701",
        "website":"https://clearleddisplays.com",
        "business":"Transparent & fixed LED display solutions (BC)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":280,"country":"Canada","region":"North America",
        "company_en":"Canadian LED","company_local":"Canadian LED",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 825-258-5393",
        "website":"https://canadian-led.com",
        "business":"LED display & lighting solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== BRAZIL new (3) ====================
    {
        "no":281,"country":"Brazil","region":"South America",
        "company_en":"WDC Networks","company_local":"WDC Networks",
        "city":"São Paulo, SP",
        "contact_name":"","title":"",
        "email":"contato@wdcnet.com.br","phone_whatsapp":"+55 11 3035-3777",
        "website":"https://wdcnet.com.br",
        "business":"LED display & networking solutions distribution",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":282,"country":"Brazil","region":"South America",
        "company_en":"Elétrica Tensão","company_local":"Elétrica Tensão",
        "city":"São Paulo, SP",
        "contact_name":"","title":"",
        "email":"vendas@eletricatensao.com.br",
        "phone_whatsapp":"+55 11 2548-0947 / WA: +55 11 2548-1215",
        "website":"https://eletricatensao.com.br",
        "business":"LED display & electrical solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":283,"country":"Brazil","region":"South America",
        "company_en":"Absen Brasil","company_local":"Absen Brasil",
        "city":"São Paulo, SP",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://absen.com.br",
        "business":"Fine-pitch LED display (Absen global brand, Brazil branch)",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== CHILE new (2) ====================
    {
        "no":284,"country":"Chile","region":"South America",
        "company_en":"Bokado Eventos","company_local":"Bokado Eventos",
        "city":"Coquimbo",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"0056 512 406339",
        "website":"https://bokadoeventos.cl",
        "business":"LED screen rental & event production (northern Chile)",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":285,"country":"Chile","region":"South America",
        "company_en":"YES Chile","company_local":"YES Chile",
        "city":"Santiago",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"00562 2321 2000",
        "website":"https://yes.cl",
        "business":"LED display & technology solutions (Santiago)",
        "facebook":"","instagram":"","linkedin":"",
    },
    # ==================== KOREA new (15) ====================
    {
        "no":286,"country":"Korea","region":"Asia",
        "company_en":"KoreaSign","company_local":"코리아싸인",
        "city":"Goyang, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"info@koreasign.co.kr","phone_whatsapp":"+82 31 924-0477",
        "website":"http://www.koreasign.co.kr",
        "business":"LED display, transparent LED film, traffic & environment signage, LED media art",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":287,"country":"Korea","region":"Asia",
        "company_en":"APS Inc Korea","company_local":"APS",
        "city":"Hwaseong, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 31 776-1800",
        "website":"https://apsinc.co.kr",
        "business":"LED display solutions & AV integration",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":288,"country":"Korea","region":"Asia",
        "company_en":"Glow M","company_local":"글로우엠",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"info@glowm.co.kr","phone_whatsapp":"+82 2 546-0804",
        "website":"https://glowm.co.kr",
        "business":"LED display & media facade solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":289,"country":"Korea","region":"Asia",
        "company_en":"LDS (LED Display Solution)","company_local":"LDS",
        "city":"Bucheon, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"info@ldss.co.kr","phone_whatsapp":"+82 32 678-7792",
        "website":"https://ldss.co.kr",
        "business":"Indoor/outdoor LED display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":290,"country":"Korea","region":"Asia",
        "company_en":"LUMENS LED","company_local":"루멘스",
        "city":"Yongin, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://lumensleds.com",
        "business":"LED display & lighting manufacturer",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":291,"country":"Korea","region":"Asia",
        "company_en":"Display Hub Ltd","company_local":"디스플레이허브",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://displayhub.co.kr",
        "business":"Professional large-screen display distributor; DynaScan 360° LED, Diamond Vision OLED",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":292,"country":"Korea","region":"Asia",
        "company_en":"Display and Life (DAL)","company_local":"디스플레이앤라이프",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"http://www.displayandlife.com",
        "business":"LCD/OLED/LED digital display design & manufacturing; retail/entertainment/sports",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":293,"country":"Korea","region":"Asia",
        "company_en":"Orion Display","company_local":"오리온디스플레이",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":294,"country":"Korea","region":"Asia",
        "company_en":"Cheonghae L&S","company_local":"청해엘엔에스",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display & signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":295,"country":"Korea","region":"Asia",
        "company_en":"Jubong Information Systems","company_local":"주봉정보시스템",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display & information systems",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":296,"country":"Korea","region":"Asia",
        "company_en":"S-Cube Lab","company_local":"에스큐브랩",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display R&D & solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":297,"country":"Korea","region":"Asia",
        "company_en":"HDigi Korea","company_local":"에이치디지코리아",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display digital solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":298,"country":"Korea","region":"Asia",
        "company_en":"ILMedia","company_local":"일미디어",
        "city":"Namyangju, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://ilmedia.co.kr",
        "business":"LED media display & signage solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":299,"country":"Korea","region":"Asia",
        "company_en":"LED Sign Korea","company_local":"LED사인코리아",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"https://ledsignkorea.com",
        "business":"LED sign & display solutions",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":300,"country":"Korea","region":"Asia",
        "company_en":"Ailed","company_local":"에이일레드",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2 525-0473",
        "website":"https://ailed.co.kr",
        "business":"LED display sales & installation",
        "facebook":"","instagram":"","linkedin":"",
    },
]

leads.extend(new_entries)

# Renumber sequentially (v4 nos already 1-209, new are 210-300)
for i, lead in enumerate(leads, 1):
    lead["no"] = i

# ─── Workbook ─────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

COLOR_HEADER_BG   = "1F4E79"
COLOR_HEADER_FONT = "FFFFFF"

COUNTRY_COLORS = {
    "Brazil":    "D9EAD3",
    "Colombia":  "FCE5CD",
    "Chile":     "D0E4F5",
    "Peru":      "FFF2CC",
    "Argentina": "EAD1DC",
    "USA":       "CFE2F3",
    "Canada":    "E6D0DE",
    "Mexico":    "D9D2E9",
    "Korea":     "F4CCCC",
}

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp / Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn",
    "Notes / Action",
]

thin   = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill  = PatternFill("solid", fgColor=COLOR_HEADER_BG)
hfont  = Font(bold=True, color=COLOR_HEADER_FONT, size=11)
halign = Alignment(horizontal="center", vertical="center", wrap_text=True)

for ci, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=ci, value=h)
    cell.fill = hfill; cell.font = hfont
    cell.alignment = halign; cell.border = border
ws.row_dimensions[1].height = 36

for ri, lead in enumerate(leads, 2):
    row_data = [
        lead["no"], lead["region"], lead["country"], lead["city"],
        lead["company_en"], lead["company_local"],
        lead["contact_name"], lead["title"],
        lead["email"], lead["phone_whatsapp"],
        lead["website"], lead["business"],
        lead["facebook"], lead["instagram"], lead["linkedin"],
        "",
    ]
    rfill = PatternFill("solid", fgColor=COUNTRY_COLORS.get(lead["country"], "FFFFFF"))
    for ci, val in enumerate(row_data, 1):
        cell = ws.cell(row=ri, column=ci, value=val)
        cell.fill = rfill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[ri].height = 26

col_widths = [5, 14, 12, 25, 28, 28, 18, 14, 32, 32, 34, 50, 32, 30, 42, 20]
for ci, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.freeze_panes = "A2"

# ─── Summary sheet ────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2["A1"] = "LED Display Leads — Americas + Korea  (v5)"
ws2["A1"].font = Font(bold=True, size=14, color=COLOR_HEADER_BG)

ws2["A3"] = "Region"; ws2["B3"] = "Country"; ws2["C3"] = "# Companies"
for cell in ws2["A3:C3"][0]:
    cell.fill = hfill; cell.font = hfont
    cell.alignment = Alignment(horizontal="center")

country_count = Counter(lead["country"] for lead in leads)

summary = [
    ("South America", "Brazil",    country_count["Brazil"]),
    ("South America", "Colombia",  country_count["Colombia"]),
    ("South America", "Chile",     country_count["Chile"]),
    ("South America", "Peru",      country_count["Peru"]),
    ("South America", "Argentina", country_count["Argentina"]),
    ("North America", "USA",       country_count["USA"]),
    ("North America", "Canada",    country_count["Canada"]),
    ("North America", "Mexico",    country_count["Mexico"]),
    ("Asia",          "Korea",     country_count["Korea"]),
]
for i, (region, country, cnt) in enumerate(summary, 4):
    ws2.cell(row=i, column=1, value=region)
    ws2.cell(row=i, column=2, value=country)
    ws2.cell(row=i, column=3, value=cnt)
    f = PatternFill("solid", fgColor=COUNTRY_COLORS.get(country, "FFFFFF"))
    for c in range(1, 4):
        ws2.cell(row=i, column=c).fill = f

total_row = len(summary) + 4
ws2.cell(row=total_row, column=1, value="TOTAL")
ws2.cell(row=total_row, column=3, value=len(leads))
for c in range(1, 4):
    ws2.cell(row=total_row, column=c).font = Font(bold=True)
for col in ["A","B","C"]:
    ws2.column_dimensions[col].width = 20

# ─── 开发建议 sheet ────────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3["A1"] = "B2B开发建议 — LED显示屏海外客户"
ws3["A1"].font = Font(bold=True, size=13, color=COLOR_HEADER_BG)

tips = [
    ("巴西 (Brazil)",
     "首选WhatsApp联系；提及Absen/Leyard已在巴西有分支，证明市场成熟度。"
     "针对Crialed/LedWave等租赁公司，主推P3.9/P4.8租赁屏。"),
    ("哥伦比亚 (Colombia)",
     "Digital Signage Colombia是三星官方经销商，可切入竞品替换。"
     "Medellín科技创业氛围浓厚，Tekus/MachineTronics是潜力客户。"),
    ("智利 (Chile)",
     "Santiago集中大量系统集成商；PLED/Tecnogroup/VideoWall体量较大。"
     "提供西班牙语技术支持文档有助于建立信任。"),
    ("美国 (USA)",
     "Fusion LED/Aurora LED主攻B2B批发；SNA Displays/PixelFLEX做高端项目。"
     "佛罗里达(Miami/Doral)是LatAm出口中转站，LED Nation USA值得重点跟进。"),
    ("加拿大 (Canada)",
     "Christie Digital是行业标杆；Leader LED/NUMMAX等中小经销商更易切入。"
     "BC省(温哥华)华人商圈活跃，ClearLED/Adtronics响应快。"),
    ("韩国 (Korea)",
     "韩国客户重视技术规格与认证(KC/CE)，首封邮件附产品规格书。"
     "通过Kompass/Naver搜索补充联系人；釜山/大邱聚集中小型系统集成商。"
     "Samik Electronics(1969年)和MLS Korea是当地权威标杆品牌。"),
    ("墨西哥 (Mexico)",
     "Monterrey(RGB Tronics/iLED)和CDMX(DMX TEC/Medios Mexico)是主要市场。"
     "活动/租赁市场大，Showco/Luft Screen是重点目标。"),
    ("秘鲁 (Peru)",
     "Lima集中绝大部分需求；EXCTECLED是当地最知名分销商。"
     "价格竞争激烈，强调性价比和售后响应速度。"),
    ("阿根廷 (Argentina)",
     "Dinalight/Grupo UNO LED是本地制造商，可能成为OEM客户。"
     "阿根廷外汇管制影响进口，优先推动样品单/试用合作。"),
]

for i, (country, tip) in enumerate(tips, 3):
    ws3.cell(row=i, column=1, value=country).font = Font(bold=True)
    ws3.cell(row=i, column=2, value=tip)
    ws3.row_dimensions[i].height = 45

ws3.column_dimensions["A"].width = 20
ws3.column_dimensions["B"].width = 90
ws3["B2"] = "开发要点"
ws3["B2"].font = Font(bold=True)
ws3["A2"] = "市场"
ws3["A2"].font = Font(bold=True)
for cell in ws3["A2:B2"][0]:
    cell.fill = hfill; cell.font = Font(bold=True, color=COLOR_HEADER_FONT)
    cell.alignment = Alignment(horizontal="center")

# ─── Save ──────────────────────────────────────────────────────────────────────
out = r"C:\Users\Administrator\ai-topic-generator\output\leads\LED_Display_Leads_v5.xlsx"
wb.save(out)
print(f"Saved: {out}")
print(f"Total records: {len(leads)}")
print("Country breakdown:")
for region, country, cnt in summary:
    print(f"  {country:<12} {cnt}")
print(f"  {'TOTAL':<12} {len(leads)}")
