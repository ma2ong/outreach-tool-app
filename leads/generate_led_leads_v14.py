"""LED Display Leads — Americas + Korea — v14
Extends v13 (413 records) with 7 new entries = 420 total.
Sources: Instagram opencli + Facebook opencli + Jina website fetch
New: Colombia +3, Chile +1, USA +2, Korea +1
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v13 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12", "v13"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (414 – 420) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Colombia new (3) ====================
    {
        "no":414,"country":"Colombia","region":"Americas",
        "company_en":"LED Master Colombia","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+57 300-204-4085",
        "website":"",
        "business":"LED screen sales and installation; Colombia; Instagram @led_master_colombia (464 followers); phone confirmed from bio",
        "facebook":"","instagram":"led_master_colombia","linkedin":"",
    },
    {
        "no":415,"country":"Colombia","region":"Americas",
        "company_en":"360GROUP Colombia","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Full-service event production company; LED screens, stages, lighting, sound; transforms spaces into unforgettable events; Instagram @360groupcolombia (24,274 followers, Verified, 959 posts)",
        "facebook":"","instagram":"360groupcolombia","linkedin":"",
    },
    {
        "no":416,"country":"Colombia","region":"Americas",
        "company_en":"Eventos Kaos","company_local":"",
        "city":"Medellín",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"AV event production; professional audio, lighting, LED panels, stages, structures, effects; Medellín Colombia; Instagram @eventoskaoscolombia (6,347 followers, 490 posts)",
        "facebook":"","instagram":"eventoskaoscolombia","linkedin":"",
    },
    # ==================== Chile new (1) ====================
    {
        "no":417,"country":"Chile","region":"Americas",
        "company_en":"LED Chile Pantallas","company_local":"",
        "city":"Chile",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen sales and rental (arriendo); commercial and event LED displays; Instagram @ledchilepantallas (215 followers)",
        "facebook":"","instagram":"ledchilepantallas","linkedin":"",
    },
    # ==================== USA new (2) ====================
    {
        "no":418,"country":"USA","region":"Americas",
        "company_en":"LED Factory USA","company_local":"",
        "city":"Houston, TX",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Premier LED retailer and supplier; LED panels, signs, modules; Houston Texas; Instagram @ledfactoryusa (510 followers, 89 posts)",
        "facebook":"","instagram":"ledfactoryusa","linkedin":"",
    },
    {
        "no":419,"country":"USA","region":"Americas",
        "company_en":"Brightlink AV","company_local":"",
        "city":"USA / Canada",
        "contact_name":"","title":"",
        "email":"sales@brightlinkav.com","phone_whatsapp":"+1 855-449-4733",
        "website":"https://brightlinkav.com",
        "business":"LED & LCD video wall supplier and AV distribution; indoor fine-pitch LED walls, outdoor LED, interactive flat panels; North America; email + phone confirmed from website",
        "facebook":"brightlinkav","instagram":"","linkedin":"",
    },
    # ==================== Korea new (1) ====================
    {
        "no":420,"country":"Korea","region":"Asia",
        "company_en":"Sign Code","company_local":"사인코드",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"010-5180-0775",
        "website":"",
        "business":"LED sign + signboard specialist; LED간판, 입간판, 채널간판, 실내간판; 3,000+ installations; 10 years experience; space/brand design approach; Instagram @sign.co.design (4,329 followers, Verified, 249 posts); phone confirmed from bio",
        "facebook":"","instagram":"sign.co.design","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v14 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-07  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v14"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","사인코드(010-5180-0775, 3000+装案, Verified)做LED招牌, 빛가람시스템(bgr@bitgaram.biz, 1660-1645)釜山海云台。현대퓨처넷大型集成商。LSI정보기술大田(042-627-8400)。이지미디어(010-8780-1003)教堂市场。"),
    ("USA","Brightlink(sales@brightlinkav.com, 1-855-449-4733, brightlinkav.com)全北美覆盖。LED Factory USA(Houston)做零售供应。FastInstall(425-766-3232)Seattle安装。Eagle LED USA(info@eagleled.us)分销商。"),
    ("Brazil","AMS Locação(André Moreira +55 18 99749-5003)。Azul Light(+55 44 3031-5492, 22K粉)最大目标。MW LED(mwled.com.br)。Infinity(2,873粉)Rio Verde。Locação Pelotas(2,132粉)。"),
    ("Canada","Inno-Leader有深圳背景，沟通顺畅。AVT/Illusion Universe/Digital Edge体量较大适合经销商。Brightlink也覆盖加拿大。"),
    ("Chile","CGS Chile(+56 2 2418-1225)、arriendopantallasledchile是重点。LED Chile Pantallas新增小型租赁商。"),
    ("Argentina","Tucumán(+54 386-551-0502)、Si-Led Buenos Aires等在库。经济不稳优先推性价比型号。"),
    ("Colombia","360GROUP(24K粉, Verified)、Eventos Kaos(6K粉, Medellín)是高价值目标。LED Master Colombia(+57 300-204-4085)有电话。Visionled(Fredy Ruiz +57 310-492-3133)重点跟进。"),
    ("Peru","Lima核心，WhatsApp是唯一有效渠道。LED Maker Peru(ledmaker.com.pe)已在库。"),
    ("Mexico","PUNTO LED(Zapopan)制造分销。MV Producciones GDL(Guadalajara)租赁。Pantallas LED Monterrey在库。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v14.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
