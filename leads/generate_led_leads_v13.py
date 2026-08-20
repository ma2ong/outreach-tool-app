"""LED Display Leads — Americas + Korea — v13
Extends v12 (403 records) with 10 new entries = 413 total.
Sources: Instagram opencli search
New: Korea +3, Brazil +7
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v12 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (404 – 413) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Korea new (3) ====================
    {
        "no":404,"country":"Korea","region":"Asia",
        "company_en":"Bitgaram System","company_local":"빛가람시스템",
        "city":"Busan (Haeundae-gu)",
        "contact_name":"","title":"",
        "email":"bgr@bitgaram.biz","phone_whatsapp":"1660-1645",
        "website":"",
        "business":"Digital signage + LED display specialist; CSAP/GS/KC certified; patent holder; 2024 Busan leading company award; Instagram @bitgaramsystem (26 followers); phone + email confirmed from bio",
        "facebook":"","instagram":"bitgaramsystem","linkedin":"",
    },
    {
        "no":405,"country":"Korea","region":"Asia",
        "company_en":"Hyundai FutureNet Signage","company_local":"현대퓨처넷 사이니지 사업부",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Signage division of Hyundai FutureNet (large Korean IT company); digital signage, DID, indoor/outdoor LED display solutions for corporate and retail; Instagram @hyundaifuturenet_signage (215 followers)",
        "facebook":"","instagram":"hyundaifuturenet_signage","linkedin":"",
    },
    {
        "no":406,"country":"Korea","region":"Asia",
        "company_en":"Harmony Company","company_local":"하모니컴퍼니",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"010-4867-5113",
        "website":"",
        "business":"Stage lighting design & event rental agency; LED display rental for events, concerts, live halls; Instagram @_harmony.company (270 followers); phone confirmed from bio",
        "facebook":"","instagram":"_harmony.company","linkedin":"",
    },
    # ==================== Brazil new (7) ====================
    {
        "no":407,"country":"Brazil","region":"Americas",
        "company_en":"BH Evento","company_local":"",
        "city":"Belo Horizonte, MG",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + lighting + stage + bands + DJ rental; Belo Horizonte; Instagram @bh.evento (555 followers)",
        "facebook":"","instagram":"bh.evento","linkedin":"",
    },
    {
        "no":408,"country":"Brazil","region":"Americas",
        "company_en":"Martins Som Luz","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + lighting + stage full-service event rental; 23 years experience; Instagram @martinssomluz (1,070 followers)",
        "facebook":"","instagram":"martinssomluz","linkedin":"",
    },
    {
        "no":409,"country":"Brazil","region":"Americas",
        "company_en":"Infinity Produções e Eventos","company_local":"",
        "city":"Rio Verde, GO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + DJ + band + stage + ground rental; wedding, corporate, agribusiness events; Rio Verde Goiás; Instagram @infinityproducoeseventos (2,873 followers)",
        "facebook":"","instagram":"infinityproducoeseventos","linkedin":"",
    },
    {
        "no":410,"country":"Brazil","region":"Americas",
        "company_en":"Locação Eventos Pelotas","company_local":"",
        "city":"Pelotas, RS",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + light + projection rental; event production services; Pelotas Rio Grande do Sul; Instagram @locacaoeventos (2,132 followers, 1,060 posts)",
        "facebook":"","instagram":"locacaoeventos","linkedin":"",
    },
    {
        "no":411,"country":"Brazil","region":"Americas",
        "company_en":"Haro Eventos","company_local":"",
        "city":"Palmas, TO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + lighting + stage rental; DJ; Palmas Tocantins; Instagram @haro.eventos (968 followers)",
        "facebook":"","instagram":"haro.eventos","linkedin":"",
    },
    {
        "no":412,"country":"Brazil","region":"Americas",
        "company_en":"GDX Brasil","company_local":"",
        "city":"Palmas, TO",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel + sound + lighting equipment rental; event production; podcast studio rental; training room rental; Palmas Tocantins; Instagram @gdx.brasil (1,709 followers)",
        "facebook":"","instagram":"gdx.brasil","linkedin":"",
    },
    {
        "no":413,"country":"Brazil","region":"Americas",
        "company_en":"AMS Locação Som Luz","company_local":"",
        "city":"Brazil",
        "contact_name":"André Moreira","title":"",
        "email":"","phone_whatsapp":"+55 18 99749-5003",
        "website":"",
        "business":"LED panel + sound + lighting rental for events; contact André Moreira; Instagram @ams.someluz (424 followers); phone confirmed from bio",
        "facebook":"","instagram":"ams.someluz","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v13 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-07  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v13"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","빛가람시스템(bgr@bitgaram.biz, 1660-1645)釜山海云台，CSAP认证。现代퓨처넷 사이니지사업부是现代集团旗下，适合大型项目。하모니컴퍼니(010-4867-5113)做活动LED租赁。LSI정보기술大田(042-627-8400, lsi8400@lsit.kr)有直接联系方式。이지미디어(010-8780-1003)专攻教堂市场。"),
    ("USA","FastInstall(425-766-3232, fastinstall.us)做视频墙安装，Eagle LED USA(info@eagleled.us)是分销商。重点推视频墙和数字标牌产品。"),
    ("Brazil","AMS Locação(André Moreira +55 18 99749-5003)有直接联系人。Azul Light(+55 44 3031-5492, 22K粉丝)是最大目标。MW LED(mwled.com.br)、RC LED、Infinity(2,873粉)、Locação Pelotas(2,132粉)活跃在不同州。"),
    ("Canada","Inno-Leader有深圳背景，沟通顺畅。AVT/Illusion Universe体量较大，适合做经销商。"),
    ("Chile","CGS Chile(+56 2 2418-1225 / WA +56 9 9532-2717)已联系。"),
    ("Argentina","Tucumán(+54 386-551-0502)、Si-Led Buenos Aires等在库。经济不稳，优先推性价比型号。"),
    ("Colombia","Visionled(Fredy Ruiz +57 310-492-3133)、Alquiler Pantallas LED Colombia(Bogotá)是重点。用WhatsApp+西班牙语。"),
    ("Peru","Lima核心，LED Maker Peru已在库。可跟进alquiler_de_pantallas_led_peru。"),
    ("Mexico","PUNTO LED(Zapopan Jalisco)制造分销大屏。MV Producciones GDL做Guadalajara租赁。Monterrey/CDMX继续开发。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v13.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
