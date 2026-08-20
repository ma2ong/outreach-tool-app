"""
LED Display Leads Generator v44 — USA batch (13 companies, nos 663-675)
Seattle WA x2, Minneapolis MN x2, Tampa FL, San Diego CA,
Salt Lake City UT x2, Cleveland OH x3, Charlotte NC, Buffalo NY
"""
import json, ast, re, os
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font

base = Path(__file__).parent

_v4_src = open(base / "generate_led_leads_v4.py", encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
for _vname in ("v5","v6","v7","v8","v9","v10","v11","v12","v13","v14","v15","v16",
               "v17","v18","v19","v20","v21","v22","v23","v24","v25","v26","v27",
               "v28","v29","v30","v31","v32","v33","v34","v35","v36","v37","v38",
               "v39","v40","v41","v42","v43"):
    _path = base / f"generate_led_leads_{_vname}.py"
    if not _path.exists():
        continue
    _src = open(_path, encoding="utf-8").read()
    _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

new_entries = [
    {
        "no": 663,
        "country": "USA",
        "region": "Americas",
        "company_en": "Northwest Video Wall",
        "company_local": "",
        "city": "Seattle, WA",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "nwvideowall.com",
        "business": "LED video wall rental, sales & installation; 1.9mm-4.8mm fine-pitch panels; 4K processing; Seattle area. Website nwvideowall.com",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 664,
        "country": "USA",
        "region": "Americas",
        "company_en": "LightSmiths",
        "company_local": "",
        "city": "Seattle, WA",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "lightsmiths.com",
        "business": "Premium LED video walls (YesTech MG9 2.9mm panels) + custom lighting/audio for events; Seattle area. Instagram @lightsmiths.seattle (8k+ followers). Website lightsmiths.com",
        "facebook": "",
        "instagram": "lightsmiths.seattle",
        "linkedin": "",
    },
    {
        "no": 665,
        "country": "USA",
        "region": "Americas",
        "company_en": "UltimateX Displays",
        "company_local": "",
        "city": "Minneapolis, MN",
        "contact_name": "",
        "title": "",
        "email": "info@ultimatexdisplays.com",
        "phone_whatsapp": "612-666-9705",
        "website": "ultimatexdisplays.com",
        "business": "Mobile MAX XL LED screen (15'x8', rotating); 10-min setup; festivals, corporate, outdoor events; Minneapolis MN. Website ultimatexdisplays.com",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 666,
        "country": "USA",
        "region": "Americas",
        "company_en": "Jagen Events",
        "company_local": "",
        "city": "Minneapolis, MN",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "jagenevents.com",
        "business": "Leading LED screen rentals in Minneapolis; trailer and modular LED walls for concerts and corporate events. Website jagenevents.com",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 667,
        "country": "USA",
        "region": "Americas",
        "company_en": "ERG247 Event Resource Group",
        "company_local": "",
        "city": "Tampa, FL",
        "contact_name": "",
        "title": "",
        "email": "hello@erg247.com",
        "phone_whatsapp": "800-941-8852",
        "website": "erg247.com",
        "business": "20+ years LED expertise; mobile LED/jumbotron rentals across Florida; Tampa, Orlando, Jacksonville, Miami coverage. Instagram @erg247. Website erg247.com",
        "facebook": "",
        "instagram": "erg247",
        "linkedin": "",
    },
    {
        "no": 668,
        "country": "USA",
        "region": "Americas",
        "company_en": "Control Entertainment",
        "company_local": "",
        "city": "San Diego, CA",
        "contact_name": "Alisha",
        "title": "",
        "email": "alisha@control-entertainment.com",
        "phone_whatsapp": "",
        "website": "control-entertainment.com",
        "business": "LED video wall and jumbotron rentals + full AV event services; established 2005; San Diego CA. Instagram @control_entertainment. Website control-entertainment.com",
        "facebook": "",
        "instagram": "control_entertainment",
        "linkedin": "",
    },
    {
        "no": 669,
        "country": "USA",
        "region": "Americas",
        "company_en": "Take One Audiovisual",
        "company_local": "",
        "city": "Salt Lake City, UT",
        "contact_name": "",
        "title": "",
        "email": "info@takeoneav.com",
        "phone_whatsapp": "",
        "website": "takeoneav.com",
        "business": "LED wall rentals + on-site tech support (IMAG, content playback); serves Utah statewide; Salt Lake City UT. Instagram @takeoneav. Website takeoneav.com",
        "facebook": "",
        "instagram": "takeoneav",
        "linkedin": "",
    },
    {
        "no": 670,
        "country": "USA",
        "region": "Americas",
        "company_en": "Mountain AV Events",
        "company_local": "",
        "city": "Salt Lake City, UT",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "mountainav-events.com",
        "business": "Micro LED video walls + audio + full event production; SLC, Provo, Park City coverage. Instagram @mountain_av. Website mountainav-events.com",
        "facebook": "",
        "instagram": "mountain_av",
        "linkedin": "",
    },
    {
        "no": 671,
        "country": "USA",
        "region": "Americas",
        "company_en": "Ohio LED Wall",
        "company_local": "",
        "city": "Cleveland, OH",
        "contact_name": "",
        "title": "",
        "email": "info@ohioledwall.com",
        "phone_whatsapp": "",
        "website": "ohioledwall.com",
        "business": "Specialist LED video wall rental, sales & installation; Cleveland, Akron, Canton OH. Website ohioledwall.com",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 672,
        "country": "USA",
        "region": "Americas",
        "company_en": "Great Lakes Audio Visual",
        "company_local": "",
        "city": "Milan, OH",
        "contact_name": "",
        "title": "",
        "email": "info@greatlakesaudiovisual.com",
        "phone_whatsapp": "",
        "website": "greatlakesaudiovisual.com",
        "business": "Full-service AV production + LED wall rentals; 200+ brand partners; Cleveland/Milan OH; serves nationwide. Instagram @greatlakesav. Website greatlakesaudiovisual.com",
        "facebook": "",
        "instagram": "greatlakesav",
        "linkedin": "",
    },
    {
        "no": 673,
        "country": "USA",
        "region": "Americas",
        "company_en": "NPi Audio Visual Solutions",
        "company_local": "",
        "city": "Cleveland, OH",
        "contact_name": "",
        "title": "",
        "email": "info@npiav.com",
        "phone_whatsapp": "",
        "website": "npiav.com",
        "business": "AV rentals + staging + fine-pitch LED video wall; established 1992; Cleveland OH; national coverage. Instagram @npiaudiovisualsolutions. Website npiav.com",
        "facebook": "",
        "instagram": "npiaudiovisualsolutions",
        "linkedin": "",
    },
    {
        "no": 674,
        "country": "USA",
        "region": "Americas",
        "company_en": "AV Pro Sound and Video",
        "company_local": "",
        "city": "Monroe, NC",
        "contact_name": "",
        "title": "",
        "email": "info@avprosoundandvideo.com",
        "phone_whatsapp": "",
        "website": "avprosoundandvideo.com",
        "business": "LED video wall + AV production for events; Monroe NC (Charlotte metro area). Instagram @avproductionsnc. Website avprosoundandvideo.com",
        "facebook": "",
        "instagram": "avproductionsnc",
        "linkedin": "",
    },
    {
        "no": 675,
        "country": "USA",
        "region": "Americas",
        "company_en": "Buffalo Audio Visual",
        "company_local": "",
        "city": "Buffalo, NY",
        "contact_name": "",
        "title": "",
        "email": "info@buffaloaudiovisual.com",
        "phone_whatsapp": "",
        "website": "buffaloaudiovisual.com",
        "business": "LED screen + AV rentals; Buffalo & Western New York; 10+ years in business. Website buffaloaudiovisual.com",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

leads.extend(new_entries)

print(f"Total leads: {len(leads)}")
usa = [l for l in leads if l.get("country") == "USA"]
print(f"USA leads: {len(usa)}")
print(f"New entries: {len(new_entries)} (nos {new_entries[0]['no']}–{new_entries[-1]['no']})")

COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads v44"

headers = ["No","Country","Region","Company EN","Company Local","City",
           "Contact Name","Title","Email","Phone/WhatsApp",
           "Website","Business","Facebook","Instagram","LinkedIn"]
ws.append(headers)
for cell in ws[1]:
    cell.fill = PatternFill("solid", fgColor="334455")
    cell.font = Font(bold=True, color="FFFFFF")

for lead in leads:
    row = [
        lead.get("no"), lead.get("country"), lead.get("region"),
        lead.get("company_en"), lead.get("company_local"), lead.get("city"),
        lead.get("contact_name"), lead.get("title"), lead.get("email"),
        lead.get("phone_whatsapp"), lead.get("website"), lead.get("business"),
        lead.get("facebook"), lead.get("instagram"), lead.get("linkedin"),
    ]
    ws.append(row)
    color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=color)

ws.column_dimensions["A"].width = 6
ws.column_dimensions["D"].width = 30
ws.column_dimensions["F"].width = 25
ws.column_dimensions["I"].width = 32
ws.column_dimensions["J"].width = 20
ws.column_dimensions["K"].width = 28
ws.column_dimensions["L"].width = 60

out = base / "LED_Display_Leads_v44.xlsx"
wb.save(out)
print(f"Saved: {out}")
