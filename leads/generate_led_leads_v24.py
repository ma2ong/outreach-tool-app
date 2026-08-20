"""
LED Display Leads v24 — Colombia (Bogota) + Mexico (Leon, Aguascalientes, Guadalajara, CDMX)
1 Colombia + 4 Mexico entries, nos 542-546
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 24)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 542,
        "country": "Colombia",
        "region": "Americas",
        "company_en": "Ofimax Colombia",
        "company_local": "Ofimax",
        "city": "Bogota",
        "contact_name": "",
        "title": "",
        "email": "ventas@ofimax.org",
        "phone_whatsapp": "+57 310 2881682",
        "website": "https://ofimax.org",
        "business": "LED display screen sales and installation; full-color video LED for events, advertising, stadiums, roads; 25+ years experience; Av. Calle 19 No. 10-06, Bogota; Twitter @ofimax",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 543,
        "country": "Mexico",
        "region": "Americas",
        "company_en": "Visual Stage Mexico",
        "company_local": "VISUAL STAGE",
        "city": "Guadalajara",
        "contact_name": "",
        "title": "",
        "email": "info@visualstage.com.mx",
        "phone_whatsapp": "+52 331 5431089",
        "website": "https://visualstage.com.mx",
        "business": "LED screen sales and rental for events; Guadalajara, Jalisco Mexico; national coverage",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 544,
        "country": "Mexico",
        "region": "Americas",
        "company_en": "TLmedia LED Aguascalientes",
        "company_local": "TLmedia Pantallas y Espectaculares",
        "city": "Aguascalientes",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+52 449 9110164",
        "website": "https://tlmedia.mx",
        "business": "LED advertising screens and billboards rental; Aguascalientes Mexico; OOH advertising spaces; Instagram @tlmedia.ags",
        "facebook": "",
        "instagram": "tlmedia.ags",
        "linkedin": "",
    },
    {
        "no": 545,
        "country": "Mexico",
        "region": "Americas",
        "company_en": "Multicorp LED Leon",
        "company_local": "Multicorp",
        "city": "Leon, Guanajuato",
        "contact_name": "Victor",
        "title": "",
        "email": "victormulticorp@gmail.com",
        "phone_whatsapp": "+52 477 7552000",
        "website": "https://multicorp.com.mx",
        "business": "LED screen design, installation, maintenance; 200+ clients, 70 active projects; 3-year warranty; Leon de los Aldama, Guanajuato Mexico; Instagram @multicorp.leon",
        "facebook": "",
        "instagram": "multicorp.leon",
        "linkedin": "",
    },
    {
        "no": 546,
        "country": "Mexico",
        "region": "Americas",
        "company_en": "JunglaMix Producciones CDMX",
        "company_local": "Producciones JunglaMix",
        "city": "CDMX",
        "contact_name": "",
        "title": "",
        "email": "info@junglamix.com",
        "phone_whatsapp": "+52 55 37229161",
        "website": "https://junglamix.com",
        "business": "LED screen rental for events; P2.9/P3.9/P5 pitches; indoor and outdoor; corporate events, conferences, social events; Mexico City (CDMX)",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v24"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v24.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
