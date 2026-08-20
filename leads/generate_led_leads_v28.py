"""
LED Display Leads v28 — Argentina (Neuquén, Buenos Aires x2)
3 Argentina entries, nos 559-561
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 28)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 559,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "DV TECH Neuquen",
        "company_local": "DV TECH Neuquén",
        "city": "Neuquén",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+54 9 299 5041259",
        "website": "",
        "business": "AV and LED screen rental for events in Neuquén Patagonia — up to 60m LED screens, projectors, simultaneous translation, sound and lighting; San Luis 583 Neuquén 8300; also WA 2996231582; Facebook dvtechneuquen",
        "facebook": "dvtechneuquen",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 560,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Group Effect Eventos Buenos Aires",
        "company_local": "Group Effect",
        "city": "Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+54 9 11 5812-9097",
        "website": "https://groupeffect.com.ar",
        "business": "LED P2 curved indoor and outdoor screen rental for corporate events, social events, conferences and outdoor shows; 25+ years experience; LED panels 6000 nit outdoor, 7680Hz refresh; Buenos Aires; Instagram @groupeffecteventos",
        "facebook": "groupeffect.efectosespeciales",
        "instagram": "groupeffecteventos",
        "linkedin": "",
    },
    {
        "no": 561,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Penta Eventos Buenos Aires",
        "company_local": "PentaEVENTOS PRODUCCIONES",
        "city": "Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "info@penta-eventos.com.ar",
        "phone_whatsapp": "+54 9 11 3949-1872",
        "website": "https://www.penta-eventos.com.ar",
        "business": "LED screen rental for fiestas, quinceañeras and corporate events; also full-service event production (artists, shows, entertainment); 15+ years in Buenos Aires events industry; penta-eventos.com.ar",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v28"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v28.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
