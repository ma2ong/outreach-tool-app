"""LED Display Leads — Americas + Korea — v7
Extends v6 (331 records) with 27 new entries = 358 total.
Sources: Google Search + WebFetch/Jina + WebSearch
New: Korea +8, USA +4, Brazil +6, Canada +3, Colombia +2, Mexico +4
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
_leads_src = re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1)
leads = ast.literal_eval(_leads_src)

# ── Load v5 new_entries ────────────────────────────────────────────────────────
_v5_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v5.py")
_v5_src  = open(_v5_path, encoding="utf-8").read()
_v5_new_src = re.search(r"^new_entries = (\[.+?^\])", _v5_src, re.M | re.S).group(1)
leads.extend(ast.literal_eval(_v5_new_src))

# ── Load v6 new_entries ────────────────────────────────────────────────────────
_v6_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v6.py")
_v6_src  = open(_v6_path, encoding="utf-8").read()
_v6_new_src = re.search(r"^new_entries = (\[.+?^\])", _v6_src, re.M | re.S).group(1)
leads.extend(ast.literal_eval(_v6_new_src))

# ── New entries (332 – 358) ────────────────────────────────────────────────────
new_entries = [
    # ==================== KOREA new (8) ====================
    {
        "no":332,"country":"Korea","region":"Asia",
        "company_en":"LKS Media","company_local":"LKS미디어",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"dcafe@naver.com","phone_whatsapp":"+82 2-6278-7890",
        "website":"",
        "business":"LED 전광판 제조·설치; indoor/outdoor LED display systems",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":333,"country":"Korea","region":"Asia",
        "company_en":"Kored","company_local":"코레드",
        "city":"Ansan, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 31-478-5330",
        "website":"",
        "business":"LED 전광판 제조·공급; outdoor LED signage, custom display panels",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":334,"country":"Korea","region":"Asia",
        "company_en":"Union LED","company_local":"유니온엘이디",
        "city":"Osan, Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"http://unionled.asia",
        "business":"LED display manufacturer; full-color outdoor/indoor panels, public-sector signage",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":335,"country":"Korea","region":"Asia",
        "company_en":"PST (Daegu)","company_local":"피에스티",
        "city":"Daegu",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 53-255-0770",
        "website":"",
        "business":"LED 전광판 설치·시공; outdoor advertising LED displays, Daegu region",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":336,"country":"Korea","region":"Asia",
        "company_en":"EONE OMS","company_local":"이원오엠에스",
        "city":"Seoul (Geumcheon-gu)",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 70-4211-0510",
        "website":"http://eoneoms.com",
        "business":"LED display systems integrator; outdoor/indoor LED panels, maintenance services",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":337,"country":"Korea","region":"Asia",
        "company_en":"Shinelight","company_local":"샤인라이트",
        "city":"Suwon, Gyeonggi-do",
        "contact_name":"김현진","title":"대표",
        "email":"led@shinelight.kr","phone_whatsapp":"+82 1666-1548",
        "website":"http://shinelight.kr",
        "business":"LED 전광판 전문; public/commercial outdoor signage, church LED, window LED; nationwide branches",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":338,"country":"Korea","region":"Asia",
        "company_en":"Postech Networks","company_local":"포스텍네트웍스",
        "city":"Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"http://postechnetworks.co.kr",
        "business":"LED 전광판 전문기업; public-sector (municipality, sports, traffic, disaster), 24h A/S; 4K LED, 3D screens",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":339,"country":"Korea","region":"Asia",
        "company_en":"Korea Netvision","company_local":"한국네트비젼",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display systems; outdoor/indoor LED signage, video walls — Instagram: @koreanetvision_iknv",
        "facebook":"","instagram":"koreanetvision_iknv","linkedin":"",
    },
    # ==================== USA new (4) ====================
    {
        "no":340,"country":"USA","region":"Americas",
        "company_en":"American LED Group","company_local":"",
        "city":"Miami, FL",
        "contact_name":"","title":"",
        "email":"info@americanledisplays.com","phone_whatsapp":"+1 305-420-5631",
        "website":"",
        "business":"LED display distributor/integrator; outdoor billboards, indoor screens, video walls — WhatsApp confirmed via wa.me link",
        "facebook":"","instagram":"americanledgroup","linkedin":"",
    },
    {
        "no":341,"country":"USA","region":"Americas",
        "company_en":"TechLed","company_local":"",
        "city":"Irving, TX",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+1 888-398-5891",
        "website":"https://techledwall.com",
        "business":"LED video wall specialist; indoor/outdoor LED panels, turnkey installation, nationwide USA",
        "facebook":"","instagram":"techledwall","linkedin":"",
    },
    {
        "no":342,"country":"USA","region":"Americas",
        "company_en":"Screen-LED USA","company_local":"",
        "city":"Tampa, FL",
        "contact_name":"","title":"",
        "email":"sales@screen-led.us","phone_whatsapp":"",
        "website":"https://screen-led.us",
        "business":"Mobile LED screen trailers; SimpLED/MobiLED/ContainerLED product lines; 463+ screens sold in 38 countries",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":343,"country":"USA","region":"Americas",
        "company_en":"RDL LED","company_local":"",
        "city":"California",
        "contact_name":"","title":"",
        "email":"red.logics.llc@gmail.com","phone_whatsapp":"+1 909-680-0141",
        "website":"",
        "business":"LED display solutions; indoor/outdoor LED screens, rental & permanent installs",
        "facebook":"","instagram":"rdl_led","linkedin":"",
    },
    # ==================== BRAZIL new (6) ====================
    {
        "no":344,"country":"Brazil","region":"Americas",
        "company_en":"Brasil Som Luz","company_local":"Brasil Som Luz",
        "city":"Santa Catarina, SC",
        "contact_name":"","title":"",
        "email":"contato@brasilsomluz.com.br","phone_whatsapp":"+55 49 9994-1829",
        "website":"https://brasilsomluz.com.br",
        "business":"LED panels, AV equipment, sound systems; WhatsApp: +55 49 9994-1829 (confirmed on website)",
        "facebook":"brasilsomluz","instagram":"brasilsomluz","linkedin":"",
    },
    {
        "no":345,"country":"Brazil","region":"Americas",
        "company_en":"Loja do Painel de LED","company_local":"Loja do Painel de LED",
        "city":"São Paulo, SP",
        "contact_name":"","title":"",
        "email":"contato@lojadopaineldeled.com.br","phone_whatsapp":"+55 11 91151-6975",
        "website":"https://lojadopaineldeled.com.br",
        "business":"LED panel specialist retailer; indoor/outdoor LED displays, pixel pitch 2–10mm, nationwide delivery",
        "facebook":"","instagram":"lojadopaineldeled","linkedin":"",
    },
    {
        "no":346,"country":"Brazil","region":"Americas",
        "company_en":"OneLight Brasil","company_local":"OneLight Brasil",
        "city":"João Pessoa, PB",
        "contact_name":"","title":"",
        "email":"contato@onelightbrasil.com.br","phone_whatsapp":"+55 83 99613-5133",
        "website":"https://onelightbrasil.com.br",
        "business":"LED display solutions; outdoor/indoor LED screens, digital signage, Northeast Brazil",
        "facebook":"","instagram":"onelightbrasil","linkedin":"",
    },
    {
        "no":347,"country":"Brazil","region":"Americas",
        "company_en":"Planet Iluminação","company_local":"Planet Iluminação",
        "city":"São Paulo, SP",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+55 11 93001-7804",
        "website":"",
        "business":"LED lighting & display solutions; WhatsApp: +55 11 93001-7804, Tel: +55 11 3223-7275",
        "facebook":"","instagram":"planetiluminacao","linkedin":"",
    },
    {
        "no":348,"country":"Brazil","region":"Americas",
        "company_en":"R2 Luz","company_local":"R2 Luz",
        "city":"Belo Horizonte, MG",
        "contact_name":"","title":"",
        "email":"atendimento@r2luz.com.br","phone_whatsapp":"+55 31 97153-9688",
        "website":"https://r2luz.com.br",
        "business":"LED lighting & display solutions; WhatsApp: +55 31 97153-9688, Tel: +55 31 3201-2526; Minas Gerais region",
        "facebook":"","instagram":"r2luz","linkedin":"",
    },
    {
        "no":349,"country":"Brazil","region":"Americas",
        "company_en":"LED Brasil RS","company_local":"LED Brasil",
        "city":"Novo Hamburgo, RS",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+55 51 9689-8669",
        "website":"",
        "business":"LED display supplier; outdoor/indoor panels; WhatsApp: +55 51 9689-8669; Rio Grande do Sul region",
        "facebook":"","instagram":"ledbrasil_rs","linkedin":"",
    },
    # ==================== CANADA new (3) ====================
    {
        "no":350,"country":"Canada","region":"Americas",
        "company_en":"AVT Canada","company_local":"AVT.ca",
        "city":"Mississauga, ON",
        "contact_name":"","title":"",
        "email":"info@avt.ca","phone_whatsapp":"+1 866-282-0099",
        "website":"https://avt.ca",
        "business":"AV & LED display integrator; video walls, LED signage, digital displays for corporate/events",
        "facebook":"","instagram":"avt.ca","linkedin":"",
    },
    {
        "no":351,"country":"Canada","region":"Americas",
        "company_en":"Illusion Universe","company_local":"",
        "city":"Toronto, ON",
        "contact_name":"","title":"",
        "email":"sales@illusionuniverse.ca","phone_whatsapp":"+1 647-212-2609",
        "website":"https://illusionuniverse.ca",
        "business":"LED display & AV solutions; indoor/outdoor LED, video walls, event displays; Toronto area",
        "facebook":"","instagram":"illusionuniverse","linkedin":"",
    },
    {
        "no":352,"country":"Canada","region":"Americas",
        "company_en":"Inno-Leader LED","company_local":"",
        "city":"Mississauga, ON / Richmond, BC",
        "contact_name":"","title":"",
        "email":"sales@innoleader.ca","phone_whatsapp":"+1 647-395-8882",
        "website":"http://innoleader.ca",
        "business":"LED display manufacturer/distributor (Shenzhen-backed); indoor, outdoor, rental & staging, creative LED; offices in Mississauga & Richmond",
        "facebook":"","instagram":"innoleaderled","linkedin":"",
    },
    # ==================== COLOMBIA new (2) ====================
    {
        "no":353,"country":"Colombia","region":"Americas",
        "company_en":"Pantallas LED Bogota","company_local":"Pantallas LED Bogotá",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"pantallasled@gmail.com","phone_whatsapp":"",
        "website":"https://pantallasledbogota.com",
        "business":"LED screen rental & sales; event/concert LED displays, indoor/outdoor, Bogotá & nationwide",
        "facebook":"","instagram":"pantallasledbogota","linkedin":"",
    },
    {
        "no":354,"country":"Colombia","region":"Americas",
        "company_en":"Pantallas LED Colombia","company_local":"",
        "city":"Bogotá / Cali / Medellín",
        "contact_name":"","title":"",
        "email":"info@pantallasledcolombia.com","phone_whatsapp":"",
        "website":"https://pantallasledcolombia.com",
        "business":"LED display solutions; commercial, academic, corporate; fixed indoor/outdoor, digital signage; nationwide coverage",
        "facebook":"","instagram":"pantallasledcolombia","linkedin":"",
    },
    # ==================== MEXICO new (4) ====================
    {
        "no":355,"country":"Mexico","region":"Americas",
        "company_en":"DMX Technologies","company_local":"DMX Tecnologías",
        "city":"Ciudad de México (CDMX)",
        "contact_name":"","title":"",
        "email":"ventas@pantallasled.mx","phone_whatsapp":"+52 55 3316-9827",
        "website":"https://pantallasled.mx",
        "business":"Mexican LED screen manufacturer; large-format indoor/outdoor LED, advertising & stadium screens; 20+ years",
        "facebook":"","instagram":"pantallasled.mx","linkedin":"",
    },
    {
        "no":356,"country":"Mexico","region":"Americas",
        "company_en":"HPMLED","company_local":"HPMLED",
        "city":"Naucalpan, Estado de México",
        "contact_name":"","title":"",
        "email":"cotiza@hpmled.com","phone_whatsapp":"+52 81 1158-0000",
        "website":"https://hpmled.com",
        "business":"LED display solutions; indoor/outdoor video walls, digital signage, event displays; quotation service",
        "facebook":"","instagram":"hpmled","linkedin":"",
    },
    {
        "no":357,"country":"Mexico","region":"Americas",
        "company_en":"Mundo Videowall","company_local":"El Mundo Del Videowall",
        "city":"Estado de México",
        "contact_name":"","title":"",
        "email":"info@videowall.com.mx","phone_whatsapp":"+52 55 7583-8168",
        "website":"https://videowall.com.mx",
        "business":"Videowall & LED display integrator; 15+ years; corporate, events, signage; certified AV specialists",
        "facebook":"","instagram":"mundovideowall","linkedin":"",
    },
    {
        "no":358,"country":"Mexico","region":"Americas",
        "company_en":"SAP LED","company_local":"SAP LED",
        "city":"San Luis Potosí, SLP",
        "contact_name":"","title":"",
        "email":"contacto@sapled.mx","phone_whatsapp":"+52 444-210-0824",
        "website":"https://sapled.mx",
        "business":"LED display manufacturer/integrator; outdoor/indoor LED, custom configurations; San Luis Potosí & nationwide",
        "facebook":"","instagram":"sapled_mx","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

# ── Sheet 1: All Markets ───────────────────────────────────────────────────────
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn",
    "Notes/Action",
]

header_font   = Font(bold=True, color="FFFFFF", size=10)
header_fill   = PatternFill("solid", fgColor="2F5496")
header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font    = header_font
    cell.fill    = header_fill
    cell.alignment = header_align
    cell.border  = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""),
        lead.get("region",""),
        lead.get("country",""),
        lead.get("city",""),
        lead.get("company_en",""),
        lead.get("company_local",""),
        lead.get("contact_name",""),
        lead.get("title",""),
        lead.get("email",""),
        lead.get("phone_whatsapp",""),
        wa_flag,
        kk_flag,
        lead.get("website",""),
        lead.get("business",""),
        lead.get("facebook",""),
        lead.get("instagram",""),
        lead.get("linkedin",""),
        "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill      = row_fill
        cell.border    = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font      = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14

title_font  = Font(bold=True, size=12, color="2F5496")
ws2["A1"]   = "LED Display Leads — v7 Summary"
ws2["A1"].font = title_font
ws2.merge_cells("A1:B1")

ws2["A3"] = "Country"
ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font   = Font(bold=True, color="FFFFFF")
    cell.fill   = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea", "USA", "Brazil", "Canada", "Chile", "Argentina", "Colombia", "Peru", "Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    ws2.cell(row=r, column=2, value=counts.get(country, 0)).border = thin_border
    ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    r += 1

ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
ws2.cell(row=r, column=2, value=len(leads)).font = Font(bold=True)
ws2.cell(row=r, column=2).border = thin_border
ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")

ws2["A" + str(r+2)] = f"Last updated: 2026-05-06  |  Total: {len(leads)} records"
ws2["A" + str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60

ws3["A1"] = "市场开发建议 — v7"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")

advice = [
    ("Korea", "优先Kakao联系，邮件备用。公共事业/政府项目多，重点推系统集成方案。已确认Kakao的公司直接发消息开场。"),
    ("USA", "优先WhatsApp/邮件。重点推大屏广告、活动租赁和体育场馆方向。强调交期+美标认证。"),
    ("Brazil", "多数公司留有WhatsApp，直接发WA最有效。优先东南部SP/RJ市场，其次南部RS市场。葡语问候拉近距离。"),
    ("Canada", "邮件为主，电话跟进。重点推零售/企业室内屏和政府招标。AVT/Illusion Universe体量较大，可做经销商谈。"),
    ("Chile", "邮件+WhatsApp。圣地亚哥为核心，活动/广告市场优先。"),
    ("Argentina", "通货膨胀环境，买家更关注性价比。优先推标准现货型号。"),
    ("Colombia", "Bogotá市场最大，Medellín其次。活动/租赁方向需求活跃。西班牙语开发信效果更好。"),
    ("Peru", "Lima为核心市场，零售广告需求为主。WhatsApp是主流沟通工具。"),
    ("Mexico", "CDMX市场最大，Monterrey/Guadalajara其次。重点推户外广告屏和活动租赁屏。当地有工厂竞争，强调品质差异化。"),
]

ws3["A3"] = "国家"
ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font   = Font(bold=True, color="FFFFFF")
    cell.fill   = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border

for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v7.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
