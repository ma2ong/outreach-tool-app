"""
LED Display Leads v27 — Argentina (Buenos Aires x2, Bahía Blanca, Mar del Plata)
4 Argentina entries, nos 555-558
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 27)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 555,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Sidisound Argentina",
        "company_local": "Sidisound",
        "city": "Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "info@sidisound.com.ar",
        "phone_whatsapp": "+54 9 11 5581 8801",
        "website": "https://www.sidisound.com.ar",
        "business": "LED screen rental for corporate and social events, Buenos Aires; large-format LED (200m+ for corporate events like Pfizer); indoor and outdoor curved/straight LED panels; Instagram @sidisound_eventos_con_exito",
        "facebook": "",
        "instagram": "sidisound_eventos_con_exito",
        "linkedin": "",
    },
    {
        "no": 556,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Ave Fenix LEDs Bahia Blanca",
        "company_local": "AVE FÉNIX LEDS",
        "city": "Bahía Blanca, Buenos Aires Province",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+54 291 412-1109",
        "website": "https://avefenixleds.com.ar",
        "business": "LED screen provider and installer — totem, indoor/outdoor LED panels; turnkey installation with structure, config, support; also operates DOOH ad circuit covering 85% of Bahía Blanca population; Bahía Blanca Argentina",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 557,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Jupiter LED Argentina",
        "company_local": "Júpiter LED",
        "city": "Buenos Aires",
        "contact_name": "",
        "title": "",
        "email": "contact@jupiterled.com.ar",
        "phone_whatsapp": "+54 9 11 3075-7688",
        "website": "http://www.jupiterled.com.ar",
        "business": "LED screen rental and permanent installation for large events and stadiums; Pantallas Led Rental (concerts, exhibitions) and Fix (permanent indoor/outdoor); Brandsen 2059 Buenos Aires; Instagram @jupiterled",
        "facebook": "",
        "instagram": "jupiterled",
        "linkedin": "",
    },
    {
        "no": 558,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "Camilo Rental Mar del Plata",
        "company_local": "Camilorental Audiovisuales",
        "city": "Mar del Plata, Buenos Aires Province",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+54 9 223 5837161",
        "website": "https://www.camilorental.com",
        "business": "AV equipment rental for events — LED screens, projectors, sound, lighting; 40+ years serving Mar del Plata and region; professional pre-event equipment testing included; Mar del Plata Argentina",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v27"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v27.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
