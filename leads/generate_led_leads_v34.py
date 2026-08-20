"""
LED Display Leads v34 — USA (14 new entries, nos 597-610)
AV For You, Fire Up Creative, Showtime LED, KC Event Company, TSV Sound & Vision,
JAWS AVL, DPC Event Services, Outdoor LED Rentals, Elite Multimedia,
Fairfield Pro AV, FireFly AV Design, Mathes Event Productions,
Picture This Production Services, North State Audio Visual
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 34)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 597,
        "country": "USA",
        "region": "Americas",
        "company_en": "AV For You",
        "company_local": "",
        "city": "Crystal, MN",
        "contact_name": "",
        "title": "",
        "email": "rentals@avforyou.com",
        "phone_whatsapp": "952-500-8839",
        "website": "https://avforyou.com",
        "business": "Full-service AV rental company in Minneapolis area; outdoor mobile LED video wall trailers for events. Instagram @av.for.you",
        "facebook": "",
        "instagram": "av.for.you",
        "linkedin": "",
    },
    {
        "no": 598,
        "country": "USA",
        "region": "Americas",
        "company_en": "Fire Up Creative",
        "company_local": "",
        "city": "Andover, MN",
        "contact_name": "Dave",
        "title": "",
        "email": "dave@fireupvideo.com",
        "phone_whatsapp": "612-759-1012",
        "website": "https://fireupcreative.com",
        "business": "Modular LED video wall rental/production; mobile outdoor LED trailer units for Midwest events. Instagram @fireupcreative",
        "facebook": "",
        "instagram": "fireupcreative",
        "linkedin": "",
    },
    {
        "no": 599,
        "country": "USA",
        "region": "Americas",
        "company_en": "Showtime LED",
        "company_local": "",
        "city": "Kansas City, MO",
        "contact_name": "",
        "title": "",
        "email": "info@showtimeled.com",
        "phone_whatsapp": "(816) 506-9477",
        "website": "https://showtimeled.com",
        "business": "Mobile LED screen and video wall rental for watch parties, corporate events, outdoor activations in Kansas City. Instagram @showtime.led",
        "facebook": "",
        "instagram": "showtime.led",
        "linkedin": "",
    },
    {
        "no": 600,
        "country": "USA",
        "region": "Americas",
        "company_en": "KC Event Company",
        "company_local": "",
        "city": "Overland Park, KS",
        "contact_name": "",
        "title": "",
        "email": "info@kceventcompany.com",
        "phone_whatsapp": "314-550-9264",
        "website": "https://kceventcompany.com",
        "business": "LED video wall rental and full-service event production in Kansas City metro. Instagram @kceventcompany",
        "facebook": "",
        "instagram": "kceventcompany",
        "linkedin": "",
    },
    {
        "no": 601,
        "country": "USA",
        "region": "Americas",
        "company_en": "TSV Sound & Vision",
        "company_local": "",
        "city": "St. Louis, MO",
        "contact_name": "",
        "title": "",
        "email": "info@tsvstl.com",
        "phone_whatsapp": "314-658-9516",
        "website": "https://tsvsoundandvision.com",
        "business": "AV production and modular LED display rental/sales for indoor and outdoor St. Louis events. Instagram @tsvusa",
        "facebook": "",
        "instagram": "tsvusa",
        "linkedin": "",
    },
    {
        "no": 602,
        "country": "USA",
        "region": "Americas",
        "company_en": "JAWS AVL",
        "company_local": "",
        "city": "San Antonio, TX",
        "contact_name": "",
        "title": "",
        "email": "info@jawsaudio.com",
        "phone_whatsapp": "(210) 442-5847",
        "website": "https://jawsaudio.com",
        "business": "Full-service event AV and LED video wall rental/installation serving San Antonio, Houston, Dallas, and broader Texas. Instagram @jawsaudio",
        "facebook": "",
        "instagram": "jawsaudio",
        "linkedin": "",
    },
    {
        "no": 603,
        "country": "USA",
        "region": "Americas",
        "company_en": "DPC Event Services",
        "company_local": "",
        "city": "San Antonio, TX",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "210-479-5541",
        "website": "https://dpceventservices.com",
        "business": "LED video wall rental and corporate AV event production in San Antonio. Contact form only; reach via IG. Instagram @dpcevents",
        "facebook": "",
        "instagram": "dpcevents",
        "linkedin": "",
    },
    {
        "no": 604,
        "country": "USA",
        "region": "Americas",
        "company_en": "Outdoor LED Rentals",
        "company_local": "",
        "city": "Nashville, TN",
        "contact_name": "",
        "title": "",
        "email": "info@outdoorledrentals.com",
        "phone_whatsapp": "(615) 575-3945",
        "website": "https://outdoorledrentals.com",
        "business": "Outdoor HD LED screen rental with onsite technician; serves Nashville and South Florida markets. Instagram @outdoorledrentals",
        "facebook": "",
        "instagram": "outdoorledrentals",
        "linkedin": "",
    },
    {
        "no": 605,
        "country": "USA",
        "region": "Americas",
        "company_en": "Elite Multimedia",
        "company_local": "",
        "city": "Mount Juliet, TN",
        "contact_name": "",
        "title": "",
        "email": "rentals@elitemultimedia.com",
        "phone_whatsapp": "(615) 457-3540",
        "website": "https://elitemultimedia.com",
        "business": "Creative LED wall and screen rental for trade shows, conferences, and live events; serves Nashville and Indianapolis. Instagram @elitemultimedia",
        "facebook": "",
        "instagram": "elitemultimedia",
        "linkedin": "",
    },
    {
        "no": 606,
        "country": "USA",
        "region": "Americas",
        "company_en": "Fairfield Pro AV",
        "company_local": "",
        "city": "Charlotte, NC",
        "contact_name": "Sam Jr",
        "title": "",
        "email": "samjr@fairfieldpro.com",
        "phone_whatsapp": "(321) 437-5857",
        "website": "https://fairfieldproav.com",
        "business": "LED video wall rental and event AV production in Charlotte, NC. Instagram @fairfieldproav",
        "facebook": "",
        "instagram": "fairfieldproav",
        "linkedin": "",
    },
    {
        "no": 607,
        "country": "USA",
        "region": "Americas",
        "company_en": "FireFly AV Design",
        "company_local": "",
        "city": "Charlotte, NC",
        "contact_name": "Noah",
        "title": "",
        "email": "noah@fireflyavdesign.com",
        "phone_whatsapp": "",
        "website": "https://fireflyavdesign.com",
        "business": "LED wall rental and event camera services in Charlotte, NC. Instagram @fireflyclt",
        "facebook": "",
        "instagram": "fireflyclt",
        "linkedin": "",
    },
    {
        "no": 608,
        "country": "USA",
        "region": "Americas",
        "company_en": "Mathes Event Productions",
        "company_local": "",
        "city": "Chamblee, GA",
        "contact_name": "",
        "title": "",
        "email": "info@mathesevents.com",
        "phone_whatsapp": "(678) 310-7332",
        "website": "https://mathesevents.com",
        "business": "LED display rental and full event production; serves Charlotte, Nashville, and Southeast US from Atlanta base. Instagram @eventsmathes",
        "facebook": "",
        "instagram": "eventsmathes",
        "linkedin": "",
    },
    {
        "no": 609,
        "country": "USA",
        "region": "Americas",
        "company_en": "Picture This Production Services",
        "company_local": "",
        "city": "Portland, OR",
        "contact_name": "",
        "title": "",
        "email": "info@pixthis.com",
        "phone_whatsapp": "(503) 235-3456",
        "website": "https://pixthis.com",
        "business": "Camera rentals, AV production, and LED video wall rental; first AV company to offer video wall in Oregon. Instagram @picturethisportland",
        "facebook": "",
        "instagram": "picturethisportland",
        "linkedin": "",
    },
    {
        "no": 610,
        "country": "USA",
        "region": "Americas",
        "company_en": "North State Audio Visual",
        "company_local": "",
        "city": "Chico, CA",
        "contact_name": "",
        "title": "",
        "email": "info@northstateav.com",
        "phone_whatsapp": "1-833-879-6728",
        "website": "https://northstateav.com",
        "business": "Full-service AV event production with LED video wall rental; serves Sacramento, Bay Area, and Northern California. Instagram @northstateaudiovisual",
        "facebook": "",
        "instagram": "northstateaudiovisual",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    from pathlib import Path
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
        HAS_EXCEL = True
    except ImportError:
        HAS_EXCEL = False

    all_leads = load_all_leads() + new_entries
    print(f"Total leads: {len(all_leads)}")
    usa = [l for l in all_leads if l.get("country") == "USA"]
    print(f"USA: {len(usa)}")

    if not HAS_EXCEL:
        print("openpyxl not installed, skipping Excel export")
    else:
        COUNTRY_COLORS = {
            "Korea": "FFF2CC", "USA": "DDEEFF", "Brazil": "E2EFDA",
            "Canada": "FCE4D6", "Chile": "EAD1DC", "Argentina": "D9E1F2",
            "Colombia": "F4CCCC", "Peru": "FFE5B4", "Mexico": "D5E8D4",
        }
        wb = Workbook()
        ws = wb.active
        ws.title = "LED Leads"
        headers = ["No", "Country", "Region", "Company EN", "Company Local",
                   "City", "Contact", "Title", "Email", "Phone/WhatsApp",
                   "Website", "Business", "Facebook", "Instagram", "LinkedIn"]
        ws.append(headers)
        for h_cell in ws[1]:
            h_cell.font = Font(bold=True)
            h_cell.fill = PatternFill("solid", fgColor="D9D9D9")
        for lead in all_leads:
            color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
            row = [
                lead.get("no", ""), lead.get("country", ""), lead.get("region", ""),
                lead.get("company_en", ""), lead.get("company_local", ""),
                lead.get("city", ""), lead.get("contact_name", ""), lead.get("title", ""),
                lead.get("email", ""), lead.get("phone_whatsapp", ""),
                lead.get("website", ""), lead.get("business", ""),
                lead.get("facebook", ""), lead.get("instagram", ""), lead.get("linkedin", ""),
            ]
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
        out = Path(BASE) / "LED_Display_Leads_v34.xlsx"
        wb.save(out)
        print(f"Saved: {out}")
