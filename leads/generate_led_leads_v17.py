"""LED Display Leads — Americas + Korea — v17
Extends v16 (436 records) with 38 new entries = 474 total.
Sources: Instagram opencli + Facebook opencli + WebSearch
New: Chile +3, Colombia +5, USA +11, Brazil +12, Korea +2, Mexico +1, Argentina +1, Peru +2, Canada +1
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v16 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12", "v13", "v14", "v15", "v16"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (437 – 474) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Chile new (3) ====================
    {
        "no":437,"country":"Chile","region":"Americas",
        "company_en":"NordeLED","company_local":"",
        "city":"Viña del Mar",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Venta y arriendo pantallas LED; 3-year technical warranty; Viña del Mar service area Chile; Instagram @nordeled (1,130 followers, 47 posts)",
        "facebook":"","instagram":"nordeled","linkedin":"",
    },
    {
        "no":438,"country":"Chile","region":"Americas",
        "company_en":"LED Station","company_local":"",
        "city":"Chile",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED P2.5 screen rental for brands, products, events and projections; Chile; Instagram @ledstation.cl (78 followers, 7 posts); .cl domain confirmed",
        "facebook":"","instagram":"ledstation.cl","linkedin":"",
    },
    {
        "no":439,"country":"Chile","region":"Americas",
        "company_en":"Pantallas LED en Santiago","company_local":"",
        "city":"Santiago",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental (arriendo pantallas LED) for events; Santiago Chile; Facebook Pantallas Led En Santiago (1,661 followers)",
        "facebook":"Pantallas Led En Santiago","instagram":"","linkedin":"",
    },
    # ==================== Colombia new (5) ====================
    {
        "no":440,"country":"Colombia","region":"Americas",
        "company_en":"Pantallas LED Bogotá","company_local":"",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen importer; best prices in all Colombia; wholesale and retail; Bogotá; Instagram @pantallas_ledbogota (88 followers, 12 posts)",
        "facebook":"","instagram":"pantallas_ledbogota","linkedin":"",
    },
    {
        "no":441,"country":"Colombia","region":"Americas",
        "company_en":"LED Service Colombia","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Professional LED display screen maintenance nationwide; Colombia; Instagram @ledservicecolombia (136 followers, 5 posts)",
        "facebook":"","instagram":"ledservicecolombia","linkedin":"",
    },
    {
        "no":442,"country":"Colombia","region":"Americas",
        "company_en":"Fabulux","company_local":"ANDIVISION S.A.S.",
        "city":"Bogotá",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display leader in Colombia; partner of ANDIVISION S.A.S.; outdoor LED Bogotá projects; Instagram @fabulux_col (513 followers, 10 posts); Facebook @FabuluxLED (1,544 followers)",
        "facebook":"FabuluxLED","instagram":"fabulux_col","linkedin":"",
    },
    {
        "no":443,"country":"Colombia","region":"Americas",
        "company_en":"Aster LED Colombia","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display products and installations; ASTER LED COLOMBIA; Instagram @asterledco (94 followers, 53 posts)",
        "facebook":"","instagram":"asterledco","linkedin":"",
    },
    {
        "no":444,"country":"Colombia","region":"Americas",
        "company_en":"LED Pixel Colombia","company_local":"",
        "city":"Colombia",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Pixel LED installation; nationwide coverage; online quotes; a nivel nacional Colombia; Instagram @led_pixel_colombia (1,713 followers, 285 posts)",
        "facebook":"","instagram":"led_pixel_colombia","linkedin":"",
    },
    # ==================== USA new (11) ====================
    {
        "no":445,"country":"USA","region":"Americas",
        "company_en":"AV Rental 305","company_local":"",
        "city":"Miami, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED Video Wall rental; AV event production; Miami and South Florida full coverage; Instagram @avrental305 (3,277 followers, 91 posts)",
        "facebook":"","instagram":"avrental305","linkedin":"",
    },
    {
        "no":446,"country":"USA","region":"Americas",
        "company_en":"Event Smart Technology","company_local":"",
        "city":"Las Vegas, NV",
        "contact_name":"","title":"",
        "email":"info@eventstecnology.com","phone_whatsapp":"+1 (702)702-9792",
        "website":"",
        "business":"LED wall rental; AV events; covers East+West Coast and Midwest USA; Las Vegas; Instagram @event_smart_technology (1,010 followers, 191 posts); phone+email confirmed from bio",
        "facebook":"","instagram":"event_smart_technology","linkedin":"",
    },
    {
        "no":447,"country":"USA","region":"Americas",
        "company_en":"Show Boss AV","company_local":"",
        "city":"Arizona",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED video wall rental and installation; located in Arizona; preferred supplier covering all USA; Instagram @showbossav (1,053 followers, 248 posts)",
        "facebook":"","instagram":"showbossav","linkedin":"",
    },
    {
        "no":448,"country":"USA","region":"Americas",
        "company_en":"LED Productions Miami","company_local":"",
        "city":"Miami, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED video wall, floor, tunnel, cold flame, full stage production; Miami Florida; Instagram @ledproductionsmiami (4,434 followers, 633 posts)",
        "facebook":"","instagram":"ledproductionsmiami","linkedin":"",
    },
    {
        "no":449,"country":"USA","region":"Americas",
        "company_en":"LED Productions Orlando","company_local":"",
        "city":"Orlando, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED video wall, floor, cold flame, cloud effects production; Orlando Florida; Instagram @ledproductionsorlando (99 followers, 62 posts)",
        "facebook":"","instagram":"ledproductionsorlando","linkedin":"",
    },
    {
        "no":450,"country":"USA","region":"Americas",
        "company_en":"LED Tech Miami","company_local":"",
        "city":"Miami, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Professional lighting and LED screen company; Miami best LED solutions Mon-Fri 8am-5pm; Verified account; Instagram @ledtechmiami (22,892 followers, 747 posts)",
        "facebook":"","instagram":"ledtechmiami","linkedin":"",
    },
    {
        "no":451,"country":"USA","region":"Americas",
        "company_en":"LED Miami Rental Signs","company_local":"",
        "city":"Miami, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"305-328-9557",
        "website":"",
        "business":"LED rental signs; Miami FL 33144; Instagram @ledmiamirentalsigns (244 followers, 1 post); phone confirmed from bio",
        "facebook":"","instagram":"ledmiamirentalsigns","linkedin":"",
    },
    {
        "no":452,"country":"USA","region":"Americas",
        "company_en":"LED Miami Signs","company_local":"",
        "city":"Miami, FL / New York, NY",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"305-328-9557",
        "website":"",
        "business":"LED digital signage direct factory supply; New York and Miami dual locations; Instagram @ledmiamisigns (7,861 followers, 122 posts); phone confirmed from bio",
        "facebook":"","instagram":"ledmiamisigns","linkedin":"",
    },
    {
        "no":453,"country":"USA","region":"Americas",
        "company_en":"Screen LED US","company_local":"",
        "city":"Tampa, FL",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Mobile LED screen solutions; Poland factory + Tampa Florida warehouse; indoor/outdoor LED; Instagram @screenled.us (3,257 followers, 118 posts)",
        "facebook":"","instagram":"screenled.us","linkedin":"",
    },
    {
        "no":454,"country":"USA","region":"Americas",
        "company_en":"American LED Screens","company_local":"",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"High performance indoor and outdoor LED screens designed for the U.S. market; advertising and events; Instagram @americanledscreens (250 followers, 20 posts)",
        "facebook":"","instagram":"americanledscreens","linkedin":"",
    },
    {
        "no":455,"country":"USA","region":"Americas",
        "company_en":"Los Angeles LED Screens","company_local":"",
        "city":"Los Angeles, CA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Los Angeles #1 video wall and LED screen company; Instagram @losangelesledscreens (61 followers, 109 posts)",
        "facebook":"","instagram":"losangelesledscreens","linkedin":"",
    },
    # ==================== Brazil new (12) ====================
    {
        "no":456,"country":"Brazil","region":"Americas",
        "company_en":"Lumens LED","company_local":"",
        "city":"Rio de Janeiro, RJ",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Full-service event structure company; LED panel rental; 1,000+ events completed; Rio de Janeiro; Instagram @lumensledoficial (3,517 followers, 876 posts)",
        "facebook":"","instagram":"lumensledoficial","linkedin":"",
    },
    {
        "no":457,"country":"Brazil","region":"Americas",
        "company_en":"LED One","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel sales and rental; custom projects for events, residential, corporate; Brazil; Instagram @led__one (13,000 followers, 456 posts)",
        "facebook":"","instagram":"led__one","linkedin":"",
    },
    {
        "no":458,"country":"Brazil","region":"Americas",
        "company_en":"LED Bras Ourinhos","company_local":"",
        "city":"Ourinhos, SP",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel operator in Ourinhos region; 5x2m LED panel outdoor advertising; Ourinhos São Paulo Brazil; Instagram @ledbras_ourinhos",
        "facebook":"","instagram":"ledbras_ourinhos","linkedin":"",
    },
    {
        "no":459,"country":"Brazil","region":"Americas",
        "company_en":"LED Bragança","company_local":"",
        "city":"Bragança",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED advertising screen on main street of Bragança; brand and corporate advertising; Brazil; Instagram @ledbraganca (281 followers, 91 posts)",
        "facebook":"","instagram":"ledbraganca","linkedin":"",
    },
    {
        "no":460,"country":"Brazil","region":"Americas",
        "company_en":"The LED","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Brazil's largest digital visual communication company specializing in LED technology; Instagram @theledoficial (50,000 followers, 1,904 posts)",
        "facebook":"","instagram":"theledoficial","linkedin":"",
    },
    {
        "no":461,"country":"Brazil","region":"Americas",
        "company_en":"InLED Brasil","company_local":"",
        "city":"Brasília, DF",
        "contact_name":"","title":"",
        "email":"contato@inledbrasil.com","phone_whatsapp":"(61)99205-7388",
        "website":"",
        "business":"Digital LED panels; visual communication innovation; Brasília DF; Instagram @inledbrasil (1,602 followers, 353 posts); phone+email confirmed from bio",
        "facebook":"","instagram":"inledbrasil","linkedin":"",
    },
    {
        "no":462,"country":"Brazil","region":"Americas",
        "company_en":"LED Eart","company_local":"",
        "city":"Ceará, CE",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"(85)9.9432.4544",
        "website":"",
        "business":"LED panel rental and sales across Ceará state; event screens; Brazil; Instagram @ledeartoficial (12,000 followers, 473 posts); phone confirmed from bio",
        "facebook":"","instagram":"ledeartoficial","linkedin":"",
    },
    {
        "no":463,"country":"Brazil","region":"Americas",
        "company_en":"CIA do LED","company_local":"",
        "city":"Rio Grande do Sul / Santa Catarina",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"(51)992314663",
        "website":"",
        "business":"LED panel rental + pyrotechnic effects; serves RS and SC states; Brazil; Instagram @ciadoledd (4,447 followers, 320 posts); phone confirmed from bio",
        "facebook":"","instagram":"ciadoledd","linkedin":"",
    },
    {
        "no":464,"country":"Brazil","region":"Americas",
        "company_en":"LED Seven Audiovisual","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Audio, lighting and LED panel specialist; 1,000+ successful events; Brazil; Instagram @ledsevenaudiovisual (4,015 followers, 603 posts)",
        "facebook":"","instagram":"ledsevenaudiovisual","linkedin":"",
    },
    {
        "no":465,"country":"Brazil","region":"Americas",
        "company_en":"Projeta Produções","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"83 98864-8221",
        "website":"",
        "business":"Events structure specialist; LED panels, screens, projection, audio, lighting, stage; Brazil; Instagram @projetaproducoes (6,741 followers, 699 posts); phone confirmed from bio",
        "facebook":"","instagram":"projetaproducoes","linkedin":"",
    },
    {
        "no":466,"country":"Brazil","region":"Americas",
        "company_en":"Mídia InLED","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Brazil/USA/Europe large integrated AV+LED panel company; rental and sales; WhatsApp available; Instagram @midiainledreal (15,000 followers, 118 posts)",
        "facebook":"","instagram":"midiainledreal","linkedin":"",
    },
    {
        "no":467,"country":"Brazil","region":"Americas",
        "company_en":"Capital Mídias Rio","company_local":"",
        "city":"Rio de Janeiro, RJ",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"300+ outdoor advertising + LED media spaces in Rio de Janeiro; brand-audience connection; Instagram @capitalmidiasrio (1,024 followers, 192 posts)",
        "facebook":"","instagram":"capitalmidiasrio","linkedin":"",
    },
    # ==================== Korea new (2) ====================
    {
        "no":468,"country":"Korea","region":"Asia",
        "company_en":"STN Media","company_local":"에스티엔 미디어",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"02-2088-1662",
        "website":"",
        "business":"Video hardware design+construction; software content planning+operation+maintenance; projection mapping, media facade, digital signage; Seoul Korea; Instagram @stn_media (1,365 followers, 210 posts); phone confirmed from bio",
        "facebook":"","instagram":"stn_media","linkedin":"",
    },
    {
        "no":469,"country":"Korea","region":"Asia",
        "company_en":"RCH Display","company_local":"전광판 DISPLAY RCH GLOW",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED전광판 + space creation and transformation; RCH GLOW LED solutions for interior/exterior spaces; Korea; Instagram @rch_display (1,640 followers, 320 posts)",
        "facebook":"","instagram":"rch_display","linkedin":"",
    },
    # ==================== Mexico new (1) ====================
    {
        "no":470,"country":"Mexico","region":"Americas",
        "company_en":"RGB Media","company_local":"",
        "city":"Mexico",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen manufacturing and installation; México y Latinoamérica; 1,000+ projects completed; Instagram @rgb.media.led (1,433 followers, 126 posts)",
        "facebook":"","instagram":"rgb.media.led","linkedin":"",
    },
    # ==================== Argentina new (1) ====================
    {
        "no":471,"country":"Argentina","region":"Americas",
        "company_en":"Prina","company_local":"Prina | Inteligencia Visual",
        "city":"Argentina",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Technology marketing expert; LED panels for brand visual development, shows and advertising spaces; Argentina; Verified; Instagram @prina.argentina (16,435 followers, 313 posts)",
        "facebook":"","instagram":"prina.argentina","linkedin":"",
    },
    # ==================== Peru new (2) ====================
    {
        "no":472,"country":"Peru","region":"Americas",
        "company_en":"PubliLED SAC","company_local":"",
        "city":"Arequipa / Tacna",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"919 474 350",
        "website":"",
        "business":"Outdoor advertising company; LED screens in Arequipa (AQP) and Tacna Peru; Instagram @publiled.aqp (47 followers, 67 posts); phone confirmed from bio",
        "facebook":"","instagram":"publiled.aqp","linkedin":"",
    },
    {
        "no":473,"country":"Peru","region":"Americas",
        "company_en":"Ledex Digital Outdoor","company_local":"",
        "city":"Lima",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Peruvian outdoor advertising company with latest digital LED elements; strategic locations in Lima; Instagram @ledexperu (2,008 followers, 145 posts)",
        "facebook":"","instagram":"ledexperu","linkedin":"",
    },
    # ==================== Canada new (1) ====================
    {
        "no":474,"country":"Canada","region":"Americas",
        "company_en":"Genoptic Smart Displays","company_local":"",
        "city":"Canada",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Borderless modular LED digital smart signage solutions; Canada; Instagram @genopticsmartdisplays (1,221 followers, 59 posts)",
        "facebook":"","instagram":"genopticsmartdisplays","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v17 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-08  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v17"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","케이시스(1600-6187, Novastar合作, 1700+公共装案)重点推。위디에스(02-6953-4419)透明屏/视频墙。디스플레이비투비(032-363-4770)三星商业伙伴。MUXWAVE(02-716-8850)透明LED专门。사인코드(010-5180-0775, Verified)3000+装案。네오사인Sign Pro(1644-2546)仁川机场项目。STN Media(02-2088-1662)数字标牌+媒体幕墙首尔。RCH Display(1640粉)LED전광판空间设计。"),
    ("USA","Brightlink(sales@brightlinkav.com, 855-449-4733)全北美。LED Productions Miami(4434粉)活动全套。LED Tech Miami(22K粉, Verified)迈阿密照明+LED屏。LED Miami Signs(7861粉, 305-328-9557)工厂直供。Event Smart Technology(+1 702-702-9792, info@eventstecnology.com)拉斯维加斯。AV Rental 305(3277粉)南佛罗里达。Screen LED US(3257粉, Tampa)移动LED屏。"),
    ("Brazil","InLED Brasil(contato@inledbrasil.com, 61-99205-7388, 巴西利亚)重点。CIA do LED(51-992314663, RS+SC州)。LED Eart(85-9.9432.4544, 锡阿拉州)。Projeta(83-98864-8221)。The LED(50K粉)巴西最大。LED One(13K粉)、Lumens LED(3517粉, Rio)。Mídia InLED(15K粉)巴西/美/欧三地。One Light(26K粉, Verified, São Paulo)。SonoraLED(2278粉)婚礼市场。"),
    ("Canada","Inno-Leader有深圳背景。AVT/Illusion Universe/Digital Edge适合经销商。Genoptic Smart Displays(1221粉)模块化LED数字标牌。Brightlink也覆盖加拿大。"),
    ("Chile","CGS Chile(+56 2 2418-1225)重点。NordeLED(1130粉, Viña del Mar, venta+arriendo)新增。Pantallas LED en Santiago(FB 1661粉, Santiago, 租赁)新增。AgsLed(FB, Concepción, 959粉)。"),
    ("Argentina","Prina(16435粉, Verified, 科技营销+LED展示)重点新增。TNX(6484粉, Verified)买卖全套。LED VJ(+54 11 5745-4880, Buenos Aires)活动租赁。Eventos LED(@eventosled.arg)。"),
    ("Colombia","Visionled(Fredy Ruiz +57 310-492-3133)重点。LED Pixel Colombia(1713粉)全国安装。Fabulux+FabuluxLED(513粉IG+1544粉FB, Bogotá, ANDIVISION合作)。360GROUP(24K粉)、Eventos Kaos(6K粉, Medellín)。Azkar Eventos(Bogotá, 世界杯活动大屏)。"),
    ("Peru","Ledex Digital Outdoor(2008粉, Lima)户外LED广告重点。PubliLED SAC(919 474 350, Arequipa+Tacna)有电话重点开发。LED Maker Peru(ledmaker.com.pe)已在库。WhatsApp是有效渠道。"),
    ("Mexico","RGB Media(1433粉, 1000+项目, 全美洲LED制造安装)重点。Pantallas LED MTY(Monterrey, 400粉)。PUNTO LED(Zapopan)制造分销。MV Producciones GDL(Guadalajara)租赁。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v17.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
