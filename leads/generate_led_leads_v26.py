"""
LED Display Leads v26 — Argentina (Córdoba) + Chile (Santiago x2, Santiago+Viña)
1 Argentina + 3 Chile entries, nos 551-554
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 26)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 551,
        "country": "Argentina",
        "region": "Americas",
        "company_en": "TEHATRO Córdoba",
        "company_local": "TEHATRO Eventos Integrales",
        "city": "General Paz, Córdoba",
        "contact_name": "",
        "title": "",
        "email": "info@tehatro.com",
        "phone_whatsapp": "+54 351 2359983",
        "website": "https://tehatro.com",
        "business": "LED screen rental for corporate and social events in Córdoba; 35+ years experience; LED lighting, sound, special effects; Lima 763, General Paz, Córdoba Argentina; Instagram @tehatroeventos",
        "facebook": "",
        "instagram": "tehatroeventos",
        "linkedin": "",
    },
    {
        "no": 552,
        "country": "Chile",
        "region": "Americas",
        "company_en": "B7 Arriendos Chile",
        "company_local": "B7 Arriendos",
        "city": "Santiago",
        "contact_name": "",
        "title": "",
        "email": "contacto@b7arriendos.cl",
        "phone_whatsapp": "+56 9 6394 5128",
        "website": "https://www.b7arriendos.cl",
        "business": "AV equipment rental including LED screens for events; lighting, sound, LED panels; Chile; Instagram @b7arriendos",
        "facebook": "",
        "instagram": "b7arriendos",
        "linkedin": "",
    },
    {
        "no": 553,
        "country": "Chile",
        "region": "Americas",
        "company_en": "CABEG SPA Chile",
        "company_local": "CABEG SPA",
        "city": "Santiago",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+56 9 92291654",
        "website": "https://www.cabeg.cl",
        "business": "AV equipment rental — LED screens (50\" Full HD), projectors, sound; presupuesto rápido por WhatsApp; Santiago Chile",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 554,
        "country": "Chile",
        "region": "Americas",
        "company_en": "VIPEVEN Chile",
        "company_local": "VIPEVEN",
        "city": "Santiago / Viña del Mar / Rancagua",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+56 9 7188 3566",
        "website": "https://www.vipeven.cl",
        "business": "LED screen and interactive screen rental for events; covers Santiago, Rancagua, Viña del Mar, Valparaíso; Instagram @productora_vipeven",
        "facebook": "",
        "instagram": "productora_vipeven",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v26"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v26.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
