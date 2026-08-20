"""LED Display Leads — Americas + Korea — v9
Extends v8 (372 records) with 7 new entries = 379 total.
Sources: eagerled.com Korea top-30 list + Instagram opencli + WebFetch/Jina
New: Korea +7
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 / v6 / v7 / v8 new_entries ────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (373 – 379) ────────────────────────────────────────────────────
new_entries = [
    # ==================== KOREA new (7) ====================
    {
        "no":373,"country":"Korea","region":"Asia",
        "company_en":"Illis Company","company_local":"일리스",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2-971-1215",
        "website":"https://illis.co.kr",
        "business":"LED display integrator; indoor/outdoor LED walls, media facade installations; LG BestShop, showrooms, galleries; full design-to-install service",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":374,"country":"Korea","region":"Asia",
        "company_en":"AVTEAM","company_local":"에이브이팀",
        "city":"Changwon, Gyeongnam",
        "contact_name":"","title":"",
        "email":"avteam2016@naver.com","phone_whatsapp":"+82 1544-0476",
        "website":"https://avteam.co.kr",
        "business":"LED display + AV total solutions; showroom, events, outdoor signage; KakaoTalk: pf.kakao.com/_MLqmG (confirmed); Naver SmartStore",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":375,"country":"Korea","region":"Asia",
        "company_en":"Ava Vision","company_local":"아바비젼",
        "city":"Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 31-443-3311",
        "website":"http://avavision.co.kr",
        "business":"LED display manufacturer & integrator; indoor/outdoor LED panels, signage; Coupang + Naver SmartStore presence",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":376,"country":"Korea","region":"Asia",
        "company_en":"Lzone","company_local":"엘존",
        "city":"Daegu",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 53-625-5966",
        "website":"https://lzone.co.kr",
        "business":"LED display specialist; outdoor advertising screens, indoor video walls; Daegu/Gyeongbuk region focus",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":377,"country":"Korea","region":"Asia",
        "company_en":"CS Sheet","company_local":"씨에스시트",
        "city":"Gyeonggi-do",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 31-935-7500",
        "website":"https://www.cssheet.co.kr",
        "business":"LED display & signage solutions; flexible LED, transparent LED, custom LED panels; manufacturing + install",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":378,"country":"Korea","region":"Asia",
        "company_en":"Coses","company_local":"코시스",
        "city":"Seoul",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"+82 2-470-2202",
        "website":"https://www.coses.co.kr",
        "business":"LED display systems; indoor/outdoor LED panels, digital signage, information display boards; nationwide install",
        "facebook":"","instagram":"","linkedin":"",
    },
    {
        "no":379,"country":"Korea","region":"Asia",
        "company_en":"Hansae Media","company_local":"한쌔미디어",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Media facade + media wall + LED display + hologram; landscape lighting, media architecture, Resolume Arena — Instagram: @hansae_media",
        "facebook":"","instagram":"hansae_media","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v9 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-06  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v9"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","韩国现有110条，仍是最大缺口。AVTEAM有Kakao官方账号，可直接Kakao联系。Lzone专注大邱/庆北地区，CS Sheet有柔性LED产品线需求可能高。"),
    ("USA","邮件+WhatsApp为主。重点推大屏广告、体育场馆、教堂市场。XcolorLED/TechLed适合做分销商开发。"),
    ("Brazil","LedWave有全国网点+美国奥兰多，是高潜力经销商目标。LED Expert有ISO认证，适合企业级采购。"),
    ("Canada","Inno-Leader有深圳背景，沟通顺畅。AVT/Illusion Universe体量较大，适合做经销商。"),
    ("Chile","Liu King做批发分销，适合报大批量价格。Santiago活动/广告市场优先。"),
    ("Argentina","经济不稳，优先推性价比型号，避免高端定制。Buenos Aires活动租赁市场活跃。"),
    ("Colombia","Medellín+Bogotá双核市场。活动/娱乐场所LED需求强劲。用WhatsApp+西班牙语开发信。"),
    ("Peru","Lima核心，WhatsApp是唯一有效渠道。从展会/活动rental切入，再推销售。"),
    ("Mexico","CDMX+Guadalajara+Monterrey三大市场。Fénix Evolution专注活动租赁，推rental产品线。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v9.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
