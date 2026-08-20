"""
generate_led_leads_v50.py
New entries: nos 715-722 (8 USA leads)
Cities: Chicopee MA / Boston MA / Lowell MA / Lafayette LA /
        Hartford CT / Wilmington DE / Harrisburg PA / Tampa FL
Note: All companies verified — LED display rental, sales or integration.
"""
import ast, re, os

_base = os.path.dirname(__file__)

def _load(path):
    return open(path, encoding="utf-8").read()

_v4_src = _load(os.path.join(_base, "generate_led_leads_v4.py"))
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

for _vname in ("v5","v6","v7","v8","v9","v10","v11","v12","v13","v14","v15","v16",
               "v17","v18","v19","v20","v21","v22","v23","v24","v25","v26","v27",
               "v28","v29","v30","v31","v32","v33","v34","v35","v36","v37","v38",
               "v39","v40","v41","v42","v43","v44","v45","v46","v47","v48","v49"):
    _path = os.path.join(_base, f"generate_led_leads_{_vname}.py")
    if not os.path.exists(_path):
        continue
    _src = _load(_path)
    _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

new_entries = [
    {
        "no": 715,
        "country": "USA",
        "region": "Americas",
        "company_en": "Zasco Productions",
        "company_local": "",
        "city": "Chicopee, MA",
        "contact_name": "",
        "title": "",
        "email": "info@zasco.com",
        "phone_whatsapp": "413-534-6677",
        "website": "zasco.com",
        "business": "Massachusetts full-service production company 25+ years; LED video wall (jumbotron) rentals & production for commencements, concerts, corporate & sports events across New England; also operates bigvideoscreen.com. Instagram @zascoproduction.",
        "facebook": "ZascoProductions",
        "instagram": "zascoproduction",
        "linkedin": "",
    },
    {
        "no": 716,
        "country": "USA",
        "region": "Americas",
        "company_en": "Boston Video Wall",
        "company_local": "",
        "city": "Boston, MA",
        "contact_name": "",
        "title": "",
        "email": "hello@videowallboston.com",
        "phone_whatsapp": "",
        "website": "videowallboston.com",
        "business": "Boston LED screen rental; conference, curved, mobile LED trailers, billboard vans, kiosk displays and LED posters for corporate events, trade shows, concerts & experiential marketing across Greater Boston.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 717,
        "country": "USA",
        "region": "Americas",
        "company_en": "Audio East (AE Event Systems)",
        "company_local": "",
        "city": "Lowell, MA",
        "contact_name": "",
        "title": "",
        "email": "sales@audioeast.com",
        "phone_whatsapp": "978-937-3944",
        "website": "audioeast.com",
        "business": "New England concert & event production company 20+ years; LED video wall rental for EDM festivals, corporate events, commencements & brand activations across Greater Boston, Springfield, NY and beyond.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 718,
        "country": "USA",
        "region": "Americas",
        "company_en": "Go Media LLC",
        "company_local": "",
        "city": "Lafayette, LA",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "337-984-9126",
        "website": "gomediallc.com",
        "business": "Louisiana AV systems integrator; direct-view LED video walls, outdoor LED, digital signage & conference room AV for corporate, government, church & sports venues across Louisiana, Dallas, Houston & Denver. Facebook: gomedialafayette.",
        "facebook": "gomedialafayette",
        "instagram": "",
        "linkedin": "gomediala",
    },
    {
        "no": 719,
        "country": "USA",
        "region": "Americas",
        "company_en": "Hartford Rents",
        "company_local": "",
        "city": "Hartford, CT",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "888-520-5667",
        "website": "hartfordrents.com",
        "business": "AV rental company serving Connecticut and Rhode Island; LED video wall tile and LED poster rentals for corporate events, exhibitions, conferences and digital signage.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 720,
        "country": "USA",
        "region": "Americas",
        "company_en": "Video WallTronics",
        "company_local": "",
        "city": "Wilmington, DE",
        "contact_name": "Mitch Kaplan",
        "title": "President",
        "email": "mitch@rentbigscreens.com",
        "phone_whatsapp": "302-328-4511",
        "website": "videowalltronics.com",
        "business": "One of the USA's first video wall rental & sales companies since 1988; LED video wall rental, staging, sales & installation for Mid-Atlantic and Northeast; features digiLED, Absen, Leyard panels. Also: rentbigscreens.com.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 721,
        "country": "USA",
        "region": "Americas",
        "company_en": "JP Lilley & Son",
        "company_local": "",
        "city": "Harrisburg, PA",
        "contact_name": "",
        "title": "",
        "email": "john.kurtinecz@jplilley.com",
        "phone_whatsapp": "717-238-8123",
        "website": "jplilley.com",
        "business": "Pennsylvania's largest AV company since 1928; LED video wall rental + sales + system integration for corporate events, hotel conferences, convention centers and commencements; in-house AV provider to major venues. Facebook: jplilleyav.",
        "facebook": "jplilleyav",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 722,
        "country": "USA",
        "region": "Americas",
        "company_en": "Tampa Video Wall Rental",
        "company_local": "",
        "city": "Tampa, FL",
        "contact_name": "",
        "title": "",
        "email": "TampaVideoWallRental@gmail.com",
        "phone_whatsapp": "813-857-2508",
        "website": "tampavideowallrental.myshopify.com",
        "business": "Tampa-based LED video wall rental for indoor and outdoor events across Florida.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

leads.extend(new_entries)

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from pathlib import Path

    base = Path(_base)
    usa = [l for l in leads if l.get("country") == "USA"]
    print(f"v50: {len(new_entries)} new entries, nos {new_entries[0]['no']}-{new_entries[-1]['no']}")
    print(f"Total: {len(leads)}, USA: {len(usa)}")
    for e in new_entries:
        ig = f"@{e['instagram']}" if e.get("instagram") else "no IG"
        print(f"  {e['no']} {e['company_en']} ({e['city']}) {ig} {e.get('email','') or 'no email'}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = [
        "No", "Country", "Region", "Company (EN)", "Company (Local)",
        "City", "Contact Name", "Title", "Email", "Phone/WhatsApp",
        "Website", "Business", "Facebook", "Instagram", "LinkedIn",
    ]
    field_keys = [
        "no", "country", "region", "company_en", "company_local",
        "city", "contact_name", "title", "email", "phone_whatsapp",
        "website", "business", "facebook", "instagram", "linkedin",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    COUNTRY_COLORS = {
        "Korea": "FFF2CC", "USA": "DDEEFF", "Brazil": "E2EFDA",
        "Canada": "FCE4D6", "Chile": "EAD1DC", "Argentina": "D9E1F2",
        "Colombia": "F4CCCC", "Peru": "FFE5B4", "Mexico": "D5E8D4",
    }

    for row_idx, lead in enumerate(leads, 2):
        country = lead.get("country", "")
        color = COUNTRY_COLORS.get(country, "F2F2F2")
        fill = PatternFill("solid", fgColor=color)
        for col_idx, key in enumerate(field_keys, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=lead.get(key, ""))
            cell.fill = fill

    col_widths = [6, 10, 10, 28, 20, 18, 18, 16, 32, 18, 28, 60, 25, 25, 25]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    out_path = base / "LED_Display_Leads_v50.xlsx"
    wb.save(out_path)
    print(f"Saved {out_path}")
