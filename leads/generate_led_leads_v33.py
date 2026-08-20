"""
LED Display Leads v33 — USA (9 new entries, nos 588-596)
Special FX Rentals, AB AV Rentals, Atlanta Pro AV, Technical Elements,
Rayne Events, Promosa, Colorado Live Events, Centric Events, Las Vegas LED Rentals
All confirmed new (dedup passed). LED wall rental, AV production, event companies.
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 33)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 588,
        "country": "USA",
        "region": "Americas",
        "company_en": "Special FX Rentals",
        "company_local": "",
        "city": "Langhorne, PA / Las Vegas, NV",
        "contact_name": "",
        "title": "",
        "email": "Info@SpecialFXRentals.com",
        "phone_whatsapp": "(800) 685-2187",
        "website": "https://specialfxrentals.com",
        "business": "Multi-city LED video wall rental + full AV production. Warehouses in Nashville, Dallas, Chicago, NYC, Tampa, Las Vegas. Large-scale event LED inventory. Instagram @specialfxrentals",
        "facebook": "",
        "instagram": "specialfxrentals",
        "linkedin": "",
    },
    {
        "no": 589,
        "country": "USA",
        "region": "Americas",
        "company_en": "AB AV Rentals",
        "company_local": "",
        "city": "Houston, TX",
        "contact_name": "Allen",
        "title": "",
        "email": "Allen@abavrentals.com",
        "phone_whatsapp": "281-235-7444",
        "website": "https://abavrentals.com",
        "business": "Houston-based AV rental specializing in LED video walls for corporate events and trade shows. Instagram @abavrentals",
        "facebook": "",
        "instagram": "abavrentals",
        "linkedin": "",
    },
    {
        "no": 590,
        "country": "USA",
        "region": "Americas",
        "company_en": "Atlanta Pro AV",
        "company_local": "",
        "city": "Atlanta, GA",
        "contact_name": "",
        "title": "",
        "email": "info@atlantaproav.com",
        "phone_whatsapp": "404-835-2230",
        "website": "https://atlantaproav.com",
        "business": "Full-service AV production company in Atlanta specializing in live/virtual event production with LED wall inventory. Instagram @atl_proav",
        "facebook": "",
        "instagram": "atl_proav",
        "linkedin": "",
    },
    {
        "no": 591,
        "country": "USA",
        "region": "Americas",
        "company_en": "Technical Elements",
        "company_local": "",
        "city": "Woodstock, GA",
        "contact_name": "",
        "title": "",
        "email": "info@teatlanta.com",
        "phone_whatsapp": "678-303-4401",
        "website": "https://teatlanta.com",
        "business": "Atlanta metro event production & rental — LED wall rentals for corporate events, concerts, product launches. Self-described Atlanta's #1 event production company.",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 592,
        "country": "USA",
        "region": "Americas",
        "company_en": "Rayne Events",
        "company_local": "",
        "city": "Washington, DC",
        "contact_name": "",
        "title": "",
        "email": "info@rayneevents.com",
        "phone_whatsapp": "(202) 695-3325",
        "website": "https://rayneevents.com",
        "business": "Full AV production company since 2010, LED video wall rental serving DC, Dallas, Nashville, Chicago, Atlanta. Multi-city high-value contact. Instagram @rayne_events",
        "facebook": "",
        "instagram": "rayne_events",
        "linkedin": "",
    },
    {
        "no": 593,
        "country": "USA",
        "region": "Americas",
        "company_en": "Promosa",
        "company_local": "",
        "city": "Kent, WA",
        "contact_name": "Stephen Dilts",
        "title": "US Sales",
        "email": "info@promosa.com",
        "phone_whatsapp": "1-888-610-6710",
        "website": "https://promosa.com",
        "business": "Seattle-area (Kent, WA) AV production with LED video wall rental for events, concerts, conferences. Also serves film/television market. Instagram @promosamgmt",
        "facebook": "",
        "instagram": "promosamgmt",
        "linkedin": "",
    },
    {
        "no": 594,
        "country": "USA",
        "region": "Americas",
        "company_en": "Colorado Live Events",
        "company_local": "",
        "city": "Centennial, CO",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "303-515-6400",
        "website": "https://coloradoliveevents.com",
        "business": "Denver-area LED video wall rental + full AV services for conferences and events. Serves Colorado Convention Center. Contact form only. Instagram @coloradoliveevents",
        "facebook": "",
        "instagram": "coloradoliveevents",
        "linkedin": "",
    },
    {
        "no": 595,
        "country": "USA",
        "region": "Americas",
        "company_en": "Centric Events",
        "company_local": "",
        "city": "Phoenix, AZ",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "602-932-6393",
        "website": "https://centric.events",
        "business": "Phoenix/Scottsdale LED video wall rental specialist — serves 16 AZ cities, trade show displays, keynote backdrops, outdoor setups. Active Instagram @centricevents",
        "facebook": "",
        "instagram": "centricevents",
        "linkedin": "",
    },
    {
        "no": 596,
        "country": "USA",
        "region": "Americas",
        "company_en": "Las Vegas LED Rentals",
        "company_local": "",
        "city": "Las Vegas, NV",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "725-699-1115",
        "website": "https://vegasledrentals.com",
        "business": "Dedicated Las Vegas LED screen rental for weddings, corporate functions, trade shows. Active Instagram @vegasledrentals",
        "facebook": "",
        "instagram": "vegasledrentals",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import json
    from pathlib import Path
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        HAS_EXCEL = True
    except ImportError:
        HAS_EXCEL = False

    all_leads = load_all_leads() + new_entries
    print(f"Total leads: {len(all_leads)}")
    usa = [l for l in all_leads if l.get("country") == "USA"]
    print(f"USA: {len(usa)}")

    if not HAS_EXCEL:
        print("openpyxl not installed, skipping Excel export")
    else:
        COUNTRY_COLORS = {
            "Korea":     "FFF2CC",
            "USA":       "DDEEFF",
            "Brazil":    "E2EFDA",
            "Canada":    "FCE4D6",
            "Chile":     "EAD1DC",
            "Argentina": "D9E1F2",
            "Colombia":  "F4CCCC",
            "Peru":      "FFE5B4",
            "Mexico":    "D5E8D4",
        }
        wb = Workbook()
        ws = wb.active
        ws.title = "LED Leads"
        headers = ["No", "Country", "Region", "Company EN", "Company Local",
                   "City", "Contact", "Title", "Email", "Phone/WhatsApp",
                   "Website", "Business", "Facebook", "Instagram", "LinkedIn"]
        ws.append(headers)
        for h_cell in ws[1]:
            h_cell.font = Font(bold=True)
            h_cell.fill = PatternFill("solid", fgColor="D9D9D9")
        for lead in all_leads:
            color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
            row = [
                lead.get("no", ""), lead.get("country", ""), lead.get("region", ""),
                lead.get("company_en", ""), lead.get("company_local", ""),
                lead.get("city", ""), lead.get("contact_name", ""), lead.get("title", ""),
                lead.get("email", ""), lead.get("phone_whatsapp", ""),
                lead.get("website", ""), lead.get("business", ""),
                lead.get("facebook", ""), lead.get("instagram", ""), lead.get("linkedin", ""),
            ]
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
        out = Path(BASE) / "LED_Display_Leads_v33.xlsx"
        wb.save(out)
        print(f"Saved: {out}")
