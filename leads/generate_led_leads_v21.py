"""
LED Display Leads v21 — Brazil (Northeast + Amazon cities)
New cities: Recife, Manaus, Fortaleza, Belém — 5 new entries, nos 526-530
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 21)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 526,
        "country": "Brazil",
        "region": "Americas",
        "company_en": "Teloes Recife",
        "company_local": "Telões Recife",
        "city": "Recife, PE",
        "contact_name": "",
        "title": "",
        "email": "contato@teloesrecife.com.br",
        "phone_whatsapp": "+55 81 3242-9099",
        "website": "https://teloesrecife.com.br",
        "business": "LED panel and AV equipment rental for corporate and social events; serves all Northeast Brazil; since 2002; Instagram @teloesrecife (4,953 followers)",
        "facebook": "",
        "instagram": "teloesrecife",
        "linkedin": "",
    },
    {
        "no": 527,
        "country": "Brazil",
        "region": "Americas",
        "company_en": "Manaus Midia LED",
        "company_local": "Manaus Midias leds e totens",
        "city": "Manaus, AM",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+55 92 99185-4822",
        "website": "",
        "business": "LED panels and digital totems at strategic fixed locations in Manaus; OOH advertising media; Instagram @manausmidia",
        "facebook": "",
        "instagram": "manausmidia",
        "linkedin": "",
    },
    {
        "no": 528,
        "country": "Brazil",
        "region": "Americas",
        "company_en": "The Voice Paineis de LED Manaus",
        "company_local": "The Voice - Painéis de LED",
        "city": "Manaus, AM",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+55 92 98419-4574",
        "website": "https://thevoice.com.br",
        "business": "Outdoor LED billboard advertising company; 12 panels (40m² and 18m²) at major avenues in Manaus; since 2009; Instagram @thevoice.manaus, Facebook thevoicemidias",
        "facebook": "thevoicemidias",
        "instagram": "thevoice.manaus",
        "linkedin": "",
    },
    {
        "no": 529,
        "country": "Brazil",
        "region": "Americas",
        "company_en": "Leddoor Locacoes Fortaleza",
        "company_local": "Leddoor Locações",
        "city": "Fortaleza, CE",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "",
        "website": "https://leddoor.com.br",
        "business": "LED panel rental for events; P5/P10/P25 panels; complete AV equipment rental; Rua Armando Monteiro 310, Vila União, Fortaleza; Facebook leddoordigital",
        "facebook": "leddoordigital",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 530,
        "country": "Brazil",
        "region": "Americas",
        "company_en": "Led Show Belem",
        "company_local": "Led Show Belém",
        "city": "Belém, PA",
        "contact_name": "",
        "title": "",
        "email": "",
        "phone_whatsapp": "+55 91 99257-2166",
        "website": "",
        "business": "LED screen, lighting, sound, and structure services for events; Belém, Pará; Instagram @led.show (2,799 followers)",
        "facebook": "",
        "instagram": "led.show",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from collections import Counter

    COUNTRY_COLORS = {
        "Korea":     "FFF2CC",
        "USA":       "DDEEFF",
        "Brazil":    "E2EFDA",
        "Canada":    "FCE4D6",
        "Chile":     "EAD1DC",
        "Argentina": "D9E1F2",
        "Colombia":  "F4CCCC",
        "Peru":      "FFE5B4",
        "Mexico":    "D5E8D4",
    }

    leads = load_all_leads()
    leads.extend(new_entries)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LED Leads v21"

    headers = ["No","Country","Region","Company EN","Company Local","City",
               "Contact","Title","Email","Phone/WA","Website",
               "Business","Facebook","Instagram","LinkedIn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(bold=True, color="FFFFFF")

    for lead in leads:
        row = [
            lead.get("no",""), lead.get("country",""), lead.get("region",""),
            lead.get("company_en",""), lead.get("company_local",""), lead.get("city",""),
            lead.get("contact_name",""), lead.get("title",""), lead.get("email",""),
            lead.get("phone_whatsapp",""), lead.get("website",""),
            lead.get("business",""), lead.get("facebook",""),
            lead.get("instagram",""), lead.get("linkedin",""),
        ]
        ws.append(row)
        color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor=color)

    out = os.path.join(BASE, "LED_Display_Leads_v21.xlsx")
    wb.save(out)
    print(f"Saved: {out}")

    counter = Counter(l.get("country") for l in leads)
    print(f"Total: {len(leads)} leads")
    for c, n in sorted(counter.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")
