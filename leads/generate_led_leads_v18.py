"""LED Display Leads — Americas + Korea — v18
Extends v17 (474 records) with 15 new entries = 489 total.
Sources: Instagram opencli
New: Colombia +1, Peru +1, Argentina +2, Mexico +1, USA +2, Canada +1, Brazil +4, Korea +3
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import ast, re, os

# ── Load v4 leads ──────────────────────────────────────────────────────────────
_v4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generate_led_leads_v4.py")
_v4_src  = open(_v4_path, encoding="utf-8").read()
leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))

# ── Load v5 – v17 new_entries ─────────────────────────────────────────────────
for _vname in ("v5", "v6", "v7", "v8", "v9", "v10", "v11", "v12", "v13", "v14", "v15", "v16", "v17"):
    _path = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"generate_led_leads_{_vname}.py")
    _src  = open(_path, encoding="utf-8").read()
    _m    = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
    if _m:
        leads.extend(ast.literal_eval(_m.group(1)))

# ── New entries (475 – 489) ────────────────────────────────────────────────────
new_entries = [
    # ==================== Colombia new (1) ====================
    {
        "no":475,"country":"Colombia","region":"Americas",
        "company_en":"Pantallas LED Cali","company_local":"",
        "city":"Cali",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen rental for events in Cali, Colombia; Instagram @pantallasledcali (23 followers, 20 posts); city confirmed from bio",
        "facebook":"","instagram":"pantallasledcali","linkedin":"",
    },
    # ==================== Peru new (1) ====================
    {
        "no":476,"country":"Peru","region":"Americas",
        "company_en":"ZYC LED Peru","company_local":"",
        "city":"Lima",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screens to create unique client experiences and achieve business objectives; Peru; Instagram @zycled.pe (65 followers, 28 posts); .pe domain confirmed",
        "facebook":"","instagram":"zycled.pe","linkedin":"",
    },
    # ==================== Argentina new (2) ====================
    {
        "no":477,"country":"Argentina","region":"Americas",
        "company_en":"Golden Rocket","company_local":"",
        "city":"Argentina",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen, totem and LED tunnel rental for events; high visual impact technology; Argentina (Argentine Spanish voseo confirmed); Instagram @goldenrocket_ar (615 followers, 151 posts)",
        "facebook":"","instagram":"goldenrocket_ar","linkedin":"",
    },
    {
        "no":478,"country":"Argentina","region":"Americas",
        "company_en":"GC Pantallas LED","company_local":"",
        "city":"Tucumán",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED advertising screen manufacturing and installation; from Tucumán to all of Argentina; best after-sales service; Instagram @gcpantallasled (4,263 followers, 196 posts)",
        "facebook":"","instagram":"gcpantallasled","linkedin":"",
    },
    # ==================== Mexico new (1) ====================
    {
        "no":479,"country":"Mexico","region":"Americas",
        "company_en":"Pantallas LED Rino Monterrey","company_local":"",
        "city":"Monterrey",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Premium LED screen rental 3.9 pitch for weddings, quinceañeras, expos, conferences, events; Monterrey, Mexico; Instagram @pantallasledrinomty (855 followers, 89 posts)",
        "facebook":"","instagram":"pantallasledrinomty","linkedin":"",
    },
    # ==================== USA new (2) ====================
    {
        "no":480,"country":"USA","region":"Americas",
        "company_en":"LED Screen Company LLC","company_local":"",
        "city":"USA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen company LLC; serves USA and Latin America (#pantallaledestadosunidos #ledscrenlatinoamerica); USA (LLC registration confirmed); Instagram @ledscreencompanyllc (89 followers, 28 posts)",
        "facebook":"","instagram":"ledscreencompanyllc","linkedin":"",
    },
    {
        "no":481,"country":"USA","region":"Americas",
        "company_en":"Led Video Wall","company_local":"",
        "city":"Rocklin, CA",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED video walls in a wide variety of sizes and resolutions; since 2005; ships worldwide; Rocklin, CA USA; Instagram @ledvideo_wall (470 followers, 24 posts); location confirmed from bio",
        "facebook":"","instagram":"ledvideo_wall","linkedin":"",
    },
    # ==================== Canada new (1) ====================
    {
        "no":482,"country":"Canada","region":"Americas",
        "company_en":"Toronto LED Video Wall Rental","company_local":"",
        "city":"Toronto",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Cutting-edge LED video wall technology with flexible rental options for unforgettable event displays; Toronto, Canada; Instagram @torontoledvideowallrental (645 followers, 10 posts)",
        "facebook":"","instagram":"torontoledvideowallrental","linkedin":"",
    },
    # ==================== Brazil new (4) ====================
    {
        "no":483,"country":"Brazil","region":"Americas",
        "company_en":"LED ABC","company_local":"",
        "city":"ABC Region, São Paulo",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED panel rental and sales for facades, commercial, corporate, fairs and events; interactivity specialist; ABC region São Paulo, Brazil; Instagram @ledabcpaineldeled (3,090 followers, 198 posts)",
        "facebook":"","instagram":"ledabcpaineldeled","linkedin":"",
    },
    {
        "no":484,"country":"Brazil","region":"Americas",
        "company_en":"ALMO RAEL","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"20 years of market leadership; customized LED screens Indoor & Outdoor; Brazil; Instagram @almorael (15,624 followers, 408 posts)",
        "facebook":"","instagram":"almorael","linkedin":"",
    },
    {
        "no":485,"country":"Brazil","region":"Americas",
        "company_en":"PLL Painéis LED","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED screen manufacturing for residential and commercial; indoor and outdoor custom installation projects; 🇧🇷 Brazil confirmed; Instagram @pllpaineldeled (542 followers, 88 posts)",
        "facebook":"","instagram":"pllpaineldeled","linkedin":"",
    },
    {
        "no":486,"country":"Brazil","region":"Americas",
        "company_en":"Ledfy","company_local":"",
        "city":"Brazil",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"OOH digital advertising LED screens; high visibility panels; measurable results; Brazil (Portuguese); Instagram @ledfyoficial (1,135 followers, 37 posts)",
        "facebook":"","instagram":"ledfyoficial","linkedin":"",
    },
    # ==================== Korea new (3) ====================
    {
        "no":487,"country":"Korea","region":"Asia",
        "company_en":"DS Lighting","company_local":"대성특수조명",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Special lighting company; laser show, laser mapping, LED display, stage lighting for concerts and clubs; Korea; Instagram @dslighting (1,131 followers, 303 posts)",
        "facebook":"","instagram":"dslighting","linkedin":"",
    },
    {
        "no":488,"country":"Korea","region":"Asia",
        "company_en":"WS International","company_local":"",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"LED display installation specialist; 전광판, media wall, custom space-fit solutions; Korea; Instagram @ws.international.official (81 followers, 67 posts)",
        "facebook":"","instagram":"ws.international.official","linkedin":"",
    },
    {
        "no":489,"country":"Korea","region":"Asia",
        "company_en":"Ksys Digital Signage","company_local":"케이시스",
        "city":"Korea",
        "contact_name":"","title":"",
        "email":"","phone_whatsapp":"",
        "website":"",
        "business":"Digital signage group; LED display, video wall, signage, LED banner, LED banner stand; Korea; Instagram @ksys_signage (449 followers, 63 posts)",
        "facebook":"","instagram":"ksys_signage","linkedin":"",
    },
]

leads.extend(new_entries)

# ── Country colour map ─────────────────────────────────────────────────────────
COUNTRY_COLORS = {
    "Korea":     "FFF2CC",
    "USA":       "DDEEFF",
    "Brazil":    "E2EFDA",
    "Canada":    "FCE4D6",
    "Chile":     "EAD1DC",
    "Argentina": "D9E1F2",
    "Colombia":  "F4CCCC",
    "Peru":      "FFE5B4",
    "Mexico":    "D5E8D4",
}

# ── Build Excel ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LED Leads - All Markets"

headers = [
    "No.", "Region", "Country", "City",
    "Company (English)", "Company (Local Language)",
    "Contact Person", "Title",
    "Email", "Phone / WhatsApp", "WhatsApp", "Kakao",
    "Website", "Main Business",
    "Facebook", "Instagram", "LinkedIn", "Notes/Action",
]

header_font  = Font(bold=True, color="FFFFFF", size=10)
header_fill  = PatternFill("solid", fgColor="2F5496")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

for col_idx, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font; cell.fill = header_fill
    cell.alignment = header_align; cell.border = thin_border

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
col_widths = [5, 10, 12, 18, 28, 24, 16, 14, 28, 18, 14, 14, 28, 36, 18, 18, 20, 22]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

for row_idx, lead in enumerate(leads, start=2):
    wa_flag = "✅" if "wa.me" in lead.get("business","").lower() or "whatsapp" in lead.get("business","").lower() else ""
    kk_flag = "✅" if "kakao" in lead.get("business","").lower() or "pf.kakao" in lead.get("business","").lower() else ""
    row = [
        lead.get("no",""), lead.get("region",""), lead.get("country",""), lead.get("city",""),
        lead.get("company_en",""), lead.get("company_local",""),
        lead.get("contact_name",""), lead.get("title",""),
        lead.get("email",""), lead.get("phone_whatsapp",""), wa_flag, kk_flag,
        lead.get("website",""), lead.get("business",""),
        lead.get("facebook",""), lead.get("instagram",""), lead.get("linkedin",""), "",
    ]
    fill_color = COUNTRY_COLORS.get(lead.get("country",""), "FFFFFF")
    row_fill   = PatternFill("solid", fgColor=fill_color)
    for col_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = row_fill; cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9)

ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet 2: Summary ───────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 14
ws2["A1"] = "LED Display Leads — v18 Summary"
ws2["A1"].font = Font(bold=True, size=12, color="2F5496")
ws2.merge_cells("A1:B1")
ws2["A3"] = "Country"; ws2["B3"] = "Company Count"
for cell in ws2["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

counts = Counter(lead["country"] for lead in leads)
countries_order = ["Korea","USA","Brazil","Canada","Chile","Argentina","Colombia","Peru","Mexico"]
r = 4
for country in countries_order:
    ws2.cell(row=r, column=1, value=country).border = thin_border
    c = ws2.cell(row=r, column=2, value=counts.get(country, 0))
    c.border = thin_border; c.alignment = Alignment(horizontal="center")
    r += 1
ws2.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws2.cell(row=r, column=1).border = thin_border
c2 = ws2.cell(row=r, column=2, value=len(leads))
c2.font = Font(bold=True); c2.border = thin_border
c2.alignment = Alignment(horizontal="center")
ws2["A"+str(r+2)] = f"Last updated: 2026-05-12  |  Total: {len(leads)} records"
ws2["A"+str(r+2)].font = Font(italic=True, color="666666")

# ── Sheet 3: 开发建议 ───────────────────────────────────────────────────────────
ws3 = wb.create_sheet("开发建议")
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 60
ws3["A1"] = "市场开发建议 — v18"
ws3["A1"].font = Font(bold=True, size=12, color="2F5496")
ws3.merge_cells("A1:B1")
advice = [
    ("Korea","케이시스(1600-6187, Novastar合作)重点推。위디에스(02-6953-4419)透明屏/视频墙。디스플레이비투비(032-363-4770)三星商业伙伴。MUXWAVE(02-716-8850)透明LED专门。사인코드(010-5180-0775)3000+装案。STN Media(02-2088-1662)数字标牌+媒体幕墙。Ksys Digital Signage(449粉)LED/视频墙/标牌组合。DS Lighting(1131粉)激光+LED舞台。WS International LED安装专门。"),
    ("USA","Brightlink(sales@brightlinkav.com)全北美。LED Productions Miami(4434粉)活动全套。LED Tech Miami(22K粉)迈阿密。Event Smart Technology(+1 702-702-9792)拉斯维加斯。AV Rental 305(3277粉)南佛罗里达。Led Video Wall Rocklin CA(470粉, since 2005)全球发货。LED Screen Company LLC(89粉)美国+拉美。"),
    ("Brazil","InLED Brasil(contato@inledbrasil.com, 61-99205-7388)重点。CIA do LED(51-992314663)。LED Eart(85-9.9432.4544)。The LED(50K粉)巴西最大。Mídia InLED(15K粉)巴西/美/欧。ALMO RAEL(15624粉, 20年市场)新增重点。LED ABC(3090粉, ABC SP)新增。PLL Painéis LED(542粉)住宅+商业。Ledfy(1135粉)OOH户外数字。"),
    ("Canada","Genoptic Smart Displays(1221粉)模块化LED标牌。Toronto LED Video Wall Rental(645粉)多伦多活动租赁新增。Inno-Leader深圳背景。Brightlink也覆盖加拿大。"),
    ("Chile","CGS Chile(+56 2 2418-1225)重点。NordeLED(1130粉, Viña del Mar)。Pantallas LED en Santiago(FB 1661粉)租赁。AgsLed(FB, Concepción)。"),
    ("Argentina","Prina(16435粉, Verified)重点。GC Pantallas LED(4263粉, Tucumán到全国)新增重点。Golden Rocket(615粉)LED/Totem/Tunnel租赁新增。TNX(6484粉, Verified)买卖全套。LED VJ(+54 11 5745-4880)活动租赁。"),
    ("Colombia","Visionled(Fredy Ruiz +57 310-492-3133)重点。LED Pixel Colombia(1713粉)全国安装。Fabulux+FabuluxLED(Bogotá, ANDIVISION合作)。360GROUP(24K粉)。Pantallas LED Cali(Cali, 活动租赁)新增。"),
    ("Peru","Ledex Digital Outdoor(2008粉, Lima)户外LED广告重点。PubliLED SAC(919 474 350, Arequipa+Tacna)有电话重点。ZYC LED Peru(65粉, .pe域名)新增。WhatsApp是有效渠道。"),
    ("Mexico","RGB Media(1433粉, 1000+项目)重点。Pantallas LED Rino Monterrey(855粉, 3.9pitch活动租赁)新增。Pantallas LED MTY(400粉)。PUNTO LED(Zapopan)制造分销。MV Producciones GDL(Guadalajara)。"),
]
ws3["A3"] = "国家"; ws3["B3"] = "开发策略"
for cell in ws3["3:3"]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2F5496")
    cell.border = thin_border
for i, (country, tip) in enumerate(advice, start=4):
    ws3.cell(row=i, column=1, value=country).border = thin_border
    ws3.cell(row=i, column=2, value=tip).border = thin_border
    ws3.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "LED_Display_Leads_v18.xlsx")
wb.save(out_path)
print(f"Saved: {out_path} ({len(leads)} records)")
