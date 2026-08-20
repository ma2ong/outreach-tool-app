"""
LED Display Leads v32 — USA (7 new entries, nos 581-587)
Imagine Media Group, Soflo Studio, Kings Rentals, Stellar XP,
R90 Lighting, Studio46 Media, Lumina Event Lighting
All confirmed new (dedup passed). LED wall rental, AV production, event companies.
"""
import ast, re, os, sys
sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(__file__)

def load_all_leads():
    _v4_src = open(os.path.join(BASE, "generate_led_leads_v4.py"), encoding="utf-8").read()
    leads = ast.literal_eval(re.search(r"^leads = (\[.+?^\])", _v4_src, re.M | re.S).group(1))
    for _vname in [f"v{i}" for i in range(5, 32)]:
        _path = os.path.join(BASE, f"generate_led_leads_{_vname}.py")
        if not os.path.exists(_path): continue
        _src = open(_path, encoding="utf-8").read()
        _m = re.search(r"^new_entries = (\[.+?^\])", _src, re.M | re.S)
        if _m: leads.extend(ast.literal_eval(_m.group(1)))
    return leads

new_entries = [
    {
        "no": 581,
        "country": "USA",
        "region": "Americas",
        "company_en": "Imagine Media Group",
        "company_local": "",
        "city": "San Diego, CA",
        "contact_name": "",
        "title": "",
        "email": "info@imaginemediagroup.com",
        "phone_whatsapp": "",
        "website": "imaginemediagroup.com",
        "business": "LED video wall rental, AV systems, sales and installation for Southern California corporate events and trade shows. Email confirmed from website",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 582,
        "country": "USA",
        "region": "Americas",
        "company_en": "Soflo Studio",
        "company_local": "",
        "city": "Sunrise (Miami/Fort Lauderdale), FL",
        "contact_name": "",
        "title": "",
        "email": "info@soflostudio.com",
        "phone_whatsapp": "",
        "website": "soflostudio.com",
        "business": "South Florida leading AV and event production company. XR LED wall, full AV rental (audio, video, lighting) for Miami and Fort Lauderdale. Instagram and TikTok active. Email confirmed from website",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 583,
        "country": "USA",
        "region": "Americas",
        "company_en": "Kings Rentals",
        "company_local": "",
        "city": "Miami, FL",
        "contact_name": "",
        "title": "",
        "email": "kingsrental@hotmail.com",
        "phone_whatsapp": "",
        "website": "kings-rental.com",
        "business": "LED video wall screen rental Miami and Broward. Lighting equipment rental for events and production. Instagram @kingsrental. Email confirmed from website",
        "facebook": "",
        "instagram": "kingsrental",
        "linkedin": "",
    },
    {
        "no": 584,
        "country": "USA",
        "region": "Americas",
        "company_en": "Stellar XP",
        "company_local": "",
        "city": "Santa Rosa, CA",
        "contact_name": "",
        "title": "",
        "email": "info@stellarxp.com",
        "phone_whatsapp": "",
        "website": "stellarxp.com",
        "business": "LED screen rental Northern California. 17ft x 10ft 6500-nit mobile LED displays for outdoor events, concerts and festivals. Instagram @stellarxp. Email confirmed from website",
        "facebook": "",
        "instagram": "stellarxp",
        "linkedin": "",
    },
    {
        "no": 585,
        "country": "USA",
        "region": "Americas",
        "company_en": "R90 Lighting",
        "company_local": "",
        "city": "Seattle, WA",
        "contact_name": "Joe",
        "title": "",
        "email": "joe@r90lighting.com",
        "phone_whatsapp": "",
        "website": "r90lighting.com",
        "business": "LED video wall, concert and event lighting, touring-grade equipment rental in Seattle area. 10+ years experience. Instagram @r90.lighting (1,660 followers). Email confirmed from website",
        "facebook": "",
        "instagram": "r90.lighting",
        "linkedin": "",
    },
    {
        "no": 586,
        "country": "USA",
        "region": "Americas",
        "company_en": "Studio46 Media",
        "company_local": "",
        "city": "Lexington, KY",
        "contact_name": "",
        "title": "",
        "email": "TEAM@STUDIO46MEDIA.COM",
        "phone_whatsapp": "",
        "website": "studio46media.com",
        "business": "LED video wall rental, video production, live and virtual event services in Kentucky and surrounding region. 3mm-10mm pixel pitch options. Email confirmed from website",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
    {
        "no": 587,
        "country": "USA",
        "region": "Americas",
        "company_en": "Lumina Event Lighting",
        "company_local": "",
        "city": "Los Angeles, CA",
        "contact_name": "",
        "title": "",
        "email": "rentals@luminaeventlighting.com",
        "phone_whatsapp": "",
        "website": "luminaeventlighting.com",
        "business": "LED video wall rental, event lighting and AV production in Los Angeles. High-end event and entertainment clients. Email confirmed from website",
        "facebook": "",
        "instagram": "",
        "linkedin": "",
    },
]

if __name__ == "__main__":
    import json
    from pathlib import Path
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        HAS_EXCEL = True
    except ImportError:
        HAS_EXCEL = False

    all_leads = load_all_leads() + new_entries
    print(f"Total leads: {len(all_leads)}")
    usa = [l for l in all_leads if l.get("country") == "USA"]
    print(f"USA: {len(usa)}")

    if not HAS_EXCEL:
        print("openpyxl not installed, skipping Excel export")
    else:
        COUNTRY_COLORS = {
            "Korea":     "FFF2CC",
            "USA":       "DDEEFF",
            "Brazil":    "E2EFDA",
            "Canada":    "FCE4D6",
            "Chile":     "EAD1DC",
            "Argentina": "D9E1F2",
            "Colombia":  "F4CCCC",
            "Peru":      "FFE5B4",
            "Mexico":    "D5E8D4",
        }
        wb = Workbook()
        ws = wb.active
        ws.title = "LED Leads"
        headers = ["No", "Country", "Region", "Company EN", "Company Local",
                   "City", "Contact", "Title", "Email", "Phone/WhatsApp",
                   "Website", "Business", "Facebook", "Instagram", "LinkedIn"]
        ws.append(headers)
        for h_cell in ws[1]:
            h_cell.font = Font(bold=True)
            h_cell.fill = PatternFill("solid", fgColor="D9D9D9")
        for lead in all_leads:
            color = COUNTRY_COLORS.get(lead.get("country", ""), "FFFFFF")
            row = [
                lead.get("no", ""), lead.get("country", ""), lead.get("region", ""),
                lead.get("company_en", ""), lead.get("company_local", ""),
                lead.get("city", ""), lead.get("contact_name", ""), lead.get("title", ""),
                lead.get("email", ""), lead.get("phone_whatsapp", ""),
                lead.get("website", ""), lead.get("business", ""),
                lead.get("facebook", ""), lead.get("instagram", ""), lead.get("linkedin", ""),
            ]
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)
        col_widths = [6, 12, 10, 28, 20, 22, 15, 15, 30, 20, 25, 60, 20, 22, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        out = Path(BASE) / "LED_Display_Leads_v32.xlsx"
        wb.save(out)
        print(f"Saved: {out}")
